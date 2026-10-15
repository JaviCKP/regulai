"""Report regressions. Provider responses are simulated at the network boundary."""
import uuid
import base64
from copy import deepcopy
from datetime import datetime, timezone
import html
import io
import json
from pathlib import Path
from unittest.mock import patch
import zipfile

import httpx
from django.test import TestCase, SimpleTestCase, override_settings
from front.agents import jobs
from front.agents import document_jobs, research, report_contract
from front.agents.data_views import tool_result
from front.agents.request_contract import interpret, allowed_tools
from front.agents.service import AgentDenied, infer
from front.models import AgentExecution, AgentArtifact
from front.test_agent_chat import ChatFixture

FIXTURES = Path(__file__).resolve().parents[2] / 'evals/reports/task55-agent-documents-20261013'


def content_fixture(name='report'):
    return json.loads((FIXTURES / (name + '-fields.json')).read_text(encoding='utf-8'))['narrative']


def bundle_fixture(name='report'):
    return base64.b64encode((FIXTURES/name/'candidate.zip').read_bytes()).decode()


@override_settings(REGULAY_AGENT_ENABLED=True)
class ReportRequestTests(ChatFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.agent.tools += ['create_mock_excel', 'build_document', 'edit_document', 'export_pdf',
                             'build_presentation', 'edit_presentation']
        self.agent.save()

    def test_five_page_report_without_upload_never_grants_fictional_excel(self):
        text = 'Crea un informe de consultoría de 5 páginas sobre el impacto de la computación cuántica en la ciberseguridad'
        for url in ('/chat/', '/chat/send/', '/api/chat/send/'):
            result = self.send(url=url, file=False, prompt=text, content=text)
            self.assertEqual(result.status_code, 202, result.content)
            run = AgentExecution.objects.get(pk=result.json()['agent']['id'])
            token, attempt = self.claimed(run)
            grant = jobs.bootstrap(token, attempt, {})
            self.assertIn('build_document', grant['tools'])
            self.assertNotIn('create_mock_excel', grant['tools'])
            self.assertEqual(grant['model_context']['request']['pages'], 5)
            jobs.cancel(self.user, run.pk)
        self.assertFalse(self.requests)

    def test_real_excel_without_source_is_actionable_and_creates_no_fiction(self):
        text = 'Crea un Excel con las ventas reales de mi empresa'
        response = self.send(file=False, prompt=text, content=text)
        self.assertEqual(response.status_code, 422, response.content)
        self.assertEqual(response.json()['error_type'], 'agent_source_required')
        self.assertFalse(AgentExecution.objects.filter(conversation__isnull=False).exists())

    def test_negated_mock_is_never_permission_to_invent(self):
        for text in ('Crea un Excel sin datos mockup', 'Crea un Excel sin datos de ejemplo',
                'Excel no ficticio', 'Excel without sample data'):
            response = self.send(file=False, prompt=text, content=text)
            self.assertEqual(response.status_code, 422, response.content)
        self.assertFalse(self.requests)

    def test_missing_research_and_unsupported_contract_are_precise_before_inference(self):
        for prompt, code in [
            ('Crea un informe Word con fuentes actuales', 'agent_research_disabled'),
            ('Crea un informe Word de 21 páginas', 'agent_page_limit'),
            ('Crea un informe Word en francés', 'agent_language_unsupported'),
            ('Crea Word y Excel', 'agent_multiple_formats_unsupported')]:
            response = self.send(file=False, prompt=prompt, content=prompt)
            self.assertEqual(response.status_code, 422, response.content)
            self.assertEqual(response.json()['error_type'], code)
        self.assertFalse(self.requests)
        self.assertFalse(AgentExecution.objects.filter(conversation__isnull=False).exists())


class NarrativeContractTests(SimpleTestCase):
    def test_revisions_preserve_styles_parts_and_other_slide_bytes(self):
        for name, base, extension in [('report-edit', 'report', 'docx'),
                ('report-edit2', 'report-edit', 'docx'), ('deck-edit', 'deck', 'pptx')]:
            original = (FIXTURES / base / ('final.' + extension)).read_bytes()
            _, proof, _ = report_contract.verify_bundle(bundle_fixture(name), content_fixture(name),
                format=extension, request={'formats': [extension]}, base=original)
            self.assertTrue(proof['automated_passed'])
        before = (FIXTURES / 'deck/final.pptx').read_bytes()
        after = (FIXTURES / 'deck-edit/final.pptx').read_bytes()
        with zipfile.ZipFile(io.BytesIO(before)) as old, zipfile.ZipFile(io.BytesIO(after)) as new:
            self.assertEqual(old.read('ppt/slides/slide1.xml'), new.read('ppt/slides/slide1.xml'))
            self.assertEqual(old.read('ppt/slides/slide3.xml'), new.read('ppt/slides/slide3.xml'))
        # A plausible-looking edit with an unrelated style mutation is rejected.
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(after)) as new, zipfile.ZipFile(out, 'w') as changed:
            for info in new.infolist():
                data = new.read(info)
                if info.filename == 'ppt/slides/slide1.xml':
                    data = data.replace(b'F5F8F7', b'FFFFFF')
                changed.writestr(info, data)
        with self.assertRaisesMessage(ValueError, 'report_revision_style_changed'):
            report_contract.verify_revision(before, out.getvalue(), 'pptx')
        with self.assertRaisesMessage(ValueError, 'report_export_source_changed'):
            report_contract.verify_revision(before, after, 'pptx', export=True)

    def test_real_render_bundles_require_exact_final_pages_content_and_coherent_formats(self):
        for name, format, count in [('report', 'docx', 5), ('report', 'pdf', 5), ('deck', 'pptx', 3)]:
            content = content_fixture(name)
            _, proof, _ = report_contract.verify_bundle(bundle_fixture(name), content,
                format=format, request={'formats': [format], 'pages': count})
            self.assertEqual(proof['pages'], count)
            with self.assertRaisesMessage(ValueError, 'deliverable_page_count_mismatch'):
                report_contract.verify_bundle(bundle_fixture(name), content,
                    format=format, request={'formats': [format], 'pages': count + 1})
            changed = deepcopy(content)
            changed['pages'][0]['blocks'][0]['text'] += ' Contenido ausente del archivo.'
            with self.assertRaises(ValueError):
                report_contract.verify_bundle(bundle_fixture(name), changed,
                    format=format, request={'formats': [format]})

    def test_no_invented_figures_sources_estimates_or_implicit_examples(self):
        content = content_fixture()
        sources = {s['id']: s for s in json.loads((FIXTURES/'primary-sources.json').read_text(encoding='utf-8'))['sources']}
        args = {'title': content['title'], 'language': 'es', 'content': content['pages']}
        request = {'formats': ['docx'], 'pages': 5, 'language': 'es'}
        report_contract.normalize(args, request, sources)
        for kind, text, citations, rationale in [
                ('analysis', 'Presupuesto 250000 euros.', [], ''),
                ('fact', 'Un estudio demuestra un ahorro.', [], ''),
                ('fact', 'El presupuesto es 999999999 euros.', ['nist-pqc'], ''),
                ('fact', 'Contenido supuestamente respaldado.', ['fabricada'], ''),
                ('estimate', 'El presupuesto sería reducido.', [], 'Parece razonable'),
                ('example', 'La empresa ficticia compra un producto.', [], '')]:
            candidate = deepcopy(args)
            candidate['content'][0]['blocks'][0] = {'kind': kind, 'text': text,
                'sources': citations, 'justification': rationale}
            with self.subTest(kind=kind, text=text), self.assertRaises(ValueError):
                report_contract.normalize(candidate, request, sources)

    def test_pdf_annotation_action_and_cross_page_text_substitution_rejected(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, NumberObject
        source = FIXTURES/'report/final.pdf'
        writer = PdfWriter(clone_from=PdfReader(source))
        action = DictionaryObject({NameObject('/S'): NameObject('/JavaScript'), NameObject('/JS'): TextStringObject('synthetic-only')})
        annotation = DictionaryObject({NameObject('/Type'): NameObject('/Annot'), NameObject('/Subtype'): NameObject('/Link'),
            NameObject('/Rect'): ArrayObject([NumberObject(0)] * 4), NameObject('/AA'): DictionaryObject({NameObject('/E'): action})})
        writer.add_annotation(0, annotation)
        out = io.BytesIO()
        writer.write(out)
        with self.assertRaisesMessage(ValueError, 'report_pdf_annotations'):
            report_contract._verify_pdf(out.getvalue(), content_fixture(), 'docx')
        writer = PdfWriter(clone_from=PdfReader(source))
        writer.pages[0].merge_page(PdfReader(source).pages[0])
        out = io.BytesIO()
        writer.write(out)
        with self.assertRaisesMessage(ValueError, 'report_pdf_unexpected_content'):
            report_contract._verify_pdf(out.getvalue(), content_fixture(), 'docx')

    def test_narrative_presentation_can_be_transformed_to_word(self):
        request = interpret('Convierte la presentación a Word', base_format='pptx')
        # The noun describing the source must not silently cause a different output.
        self.assertEqual(request['formats'], ['docx'])
        self.assertEqual(allowed_tools(request, ['build_document', 'edit_document', 'export_pdf'],
            base=True, source=True, narrative=True), ['build_document'])
        unsupported = interpret('Convierte la presentación a Excel', base_format='pptx')
        self.assertEqual(allowed_tools(unsupported, ['rename_excel', 'build_excel'],
            base=True, source=True, narrative=True), [])
        self.assertEqual(interpret('Write a five-page report in English')['pages'], 5)


@override_settings(REGULAY_AGENT_ENABLED=True, REGULAY_AGENT_RESEARCH_ENABLED=True)
class ReportLifecycleTests(ChatFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.enterContext(patch('django.utils.timezone.now', return_value=datetime(2026, 10, 13, 12, tzinfo=timezone.utc)))
        self.agent.tools += ['create_mock_excel', 'build_document', 'edit_document', 'export_pdf',
            'build_presentation', 'edit_presentation', 'research_sources']
        self.agent.save()
        self.sources = {s['id']: s for s in json.loads((FIXTURES/'primary-sources.json').read_text(encoding='utf-8'))['sources']}
        self.company.ai_config['destinations']['internal'] += ['research:' + key for key in self.sources]
        self.company.save()
        self.enterContext(patch.object(research.socket, 'getaddrinfo', return_value=[(2, 1, 6, '', ('93.184.216.34', 443))]))
        original_fetch = research.fetch_source
        def fetch(entry, **kwargs):
            source = next(s for s in self.sources.values() if s['url'] == entry['url'])
            body = '<div id="' + entry['content_id'] + '">' + html.escape(source['excerpt']) + '</div>'
            return original_fetch(entry, transport=httpx.MockTransport(lambda request:
                httpx.Response(200, headers={'content-type': 'text/html; charset=utf-8'}, text=body)), **kwargs)
        self.enterContext(patch.object(research, 'fetch_source', side_effect=fetch))

    def issue_tool(self, token, attempt, name, args):
        call = {'id': name + '-' + uuid.uuid4().hex, 'name': name, 'args': args}
        self.reply_calls = [call]
        result = infer(token, attempt, {'call_id': str(uuid.uuid4()), 'messages': [{'role': 'user',
            'content': 'Cumple la petición y utiliza solo las fuentes autorizadas.'}], 'tools': [name]})
        self.assertEqual(result['status'], 'completed', result)
        return result['tool_calls'][0]

    def create(self, name='report', url='/chat/send/'):
        text = 'Crea un informe Word y PDF de cinco páginas sobre computación cuántica y ciberseguridad con fuentes' if name == 'report' else 'Crea una presentación de tres diapositivas sobre una transición controlada'
        response = self.send(url=url, file=False, prompt=text, content=text)
        self.assertEqual(response.status_code, 202, response.content)
        run = AgentExecution.objects.get(pk=response.json()['agent']['id'])
        token, attempt = self.claimed(run)
        if name == 'report':
            research_call = self.issue_tool(token, attempt, 'research_sources', {'ids': list(self.sources)})
            result = research.execute(token, attempt, {'call': research_call})
            view = tool_result(token, attempt, {'call': research_call,
                'raw': {'content': json.dumps({'reference': result['reference']})}})
            self.assertEqual(view['research_status'], 'completed', view)
        content = content_fixture(name)
        args = {'input_reference': '', 'title': content['title'], 'section': 'Resumen de la petición.',
            'language': 'es', 'content': content['pages']}
        args.update({'format': 'docx'} if name == 'report' else {'template': 'new'})
        call = self.issue_tool(token, attempt, 'build_document' if name == 'report' else 'build_presentation', args)
        fields = jobs.local_fields(token, attempt, call)
        self.assertEqual(fields['narrative']['pages'], content['pages'])
        payload = {'operation_id': str(uuid.uuid4()), 'call': call, 'document_bundle': bundle_fixture(name)}
        result = jobs.candidate(token, attempt, payload)
        artifact = AgentArtifact.objects.get(pk=result['reference'])
        self.assertTrue(artifact.verification['automated_passed'], artifact.verification)
        self.assertEqual(artifact.status, 'draft')
        self.assertEqual(jobs.candidate(token, attempt, payload)['reference'], str(artifact.pk))
        return run, token, attempt, call, artifact, payload

    def approve(self, run, token, attempt, artifact):
        # Simulates the authenticated user's approval; pixel inspection is an
        # independent recorded step, never attributed to this test assertion.
        document_jobs.review(self.user, artifact.pk, {'sha256': artifact.sha256,
            'page_sha256': artifact.verification['page_sha256'], 'checks': document_jobs.VISUAL_CHECKS})
        jobs.complete(token, attempt, {})

    def test_three_http_entries_deliver_five_page_word_and_pdf_with_no_fictitious_amounts(self):
        from docx import Document
        for url in ('/chat/', '/chat/send/', '/api/chat/send/'):
            run, token, attempt, call, artifact, _ = self.create(url=url)
            self.assertEqual(artifact.verification['pages'], 5)
            self.assertNotIn('total_cents', artifact.verification)
            self.assertEqual(self.client.get(f'/chat/agent/artifacts/{artifact.pk}/download/').status_code, 404)
            self.approve(run, token, attempt, artifact)
            word = self.client.get(f'/chat/agent/artifacts/{artifact.pk}/download/?format=docx')
            pdf = self.client.get(f'/chat/agent/artifacts/{artifact.pk}/download/?format=pdf')
            self.assertEqual(word.status_code, 200)
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(report_contract.sha256(pdf.content), artifact.verification['pdf_sha256'])
            doc = Document(io.BytesIO(word.content))
            self.assertFalse(doc.tables)
            self.assertNotRegex(' '.join(p.text for p in doc.paragraphs), r'[€$£]|Importe|Datos ficticios')
            snapshot = self.client.get(f'/chat/agent/runs/{run.pk}/').json()
            self.assertEqual(snapshot['receipt']['verification']['status'], 'partial')
            self.assertEqual(snapshot['artifacts'][0]['verification_checks']['sources']['status'], 'not_performed')

    def test_presentation_creation_revision_edit_scope_and_cancelled_work(self):
        run, token, attempt, call, artifact, _ = self.create('deck')
        self.approve(run, token, attempt, artifact)
        text = 'Cambia el título de la segunda diapositiva'
        response = self.send(url=f'/chat/{run.conversation_id}/send/', file=False, prompt=text, content=text,
            artifact_id=str(artifact.logical_id), base_revision=str(artifact.pk))
        self.assertEqual(response.status_code, 202, response.content)
        revision = AgentExecution.objects.get(pk=response.json()['agent']['id'])
        token, attempt = self.claimed(revision)
        call = self.issue_tool(token, attempt, 'edit_presentation', {'logical_id': str(artifact.logical_id),
            'base_revision': str(artifact.pk), 'slide': 2, 'field': 'title', 'value': 'Ensayo y recuperación'})
        fields = jobs.local_fields(token, attempt, call)['narrative']
        original = report_contract.source_content(artifact)
        self.assertEqual(fields['pages'][0], original['pages'][0])
        self.assertEqual(fields['pages'][2], original['pages'][2])
        self.assertEqual(fields['pages'][1]['blocks'], original['pages'][1]['blocks'])
        candidate = jobs.candidate(token, attempt, {'operation_id': str(uuid.uuid4()), 'call': call,
            'document_bundle': bundle_fixture('deck-edit')})
        changed = AgentArtifact.objects.get(pk=candidate['reference'])
        self.assertTrue(changed.verification['automated_passed'], changed.verification)
        self.approve(revision, token, attempt, changed)
        self.assertEqual(jobs.latest(run.conversation, artifact.logical_id), changed)
        with self.assertRaises(AgentDenied):
            jobs.local_fields(token, attempt, call)

    def test_two_document_edits_keep_references_and_base_content_and_cancel_next(self):
        run, token, attempt, _, artifact, _ = self.create()
        self.approve(run, token, attempt, artifact)
        first = artifact
        for name in ('report-edit', 'report-edit2'):
            content = content_fixture(name)
            text = 'Cambia el título del informe a ' + content['title']
            response = self.send(url=f'/chat/{run.conversation_id}/send/', file=False, prompt=text, content=text,
                artifact_id=str(artifact.logical_id), base_revision=str(artifact.pk))
            self.assertEqual(response.status_code, 202, response.content)
            revision = AgentExecution.objects.get(pk=response.json()['agent']['id'])
            token, attempt = self.claimed(revision)
            grant = jobs.bootstrap(token, attempt, {})
            self.assertTrue(grant['input']['path'].startswith('/workspace/input/'))
            self.assertEqual(grant['model_context']['columns'], [])
            call = self.issue_tool(token, attempt, 'edit_document', {'logical_id': str(artifact.logical_id),
                'base_revision': str(artifact.pk), 'field': 'title', 'value': content['title']})
            fields = jobs.local_fields(token, attempt, call)['narrative']
            self.assertEqual(fields['pages'], report_contract.source_content(first)['pages'])
            self.assertEqual(len(fields['references']), 4)
            result = jobs.candidate(token, attempt, {'operation_id': str(uuid.uuid4()), 'call': call,
                'document_bundle': bundle_fixture(name)})
            artifact = AgentArtifact.objects.get(pk=result['reference'])
            self.assertTrue(artifact.verification['automated_passed'], artifact.verification)
            self.approve(revision, token, attempt, artifact)
        stale = self.send(url=f'/chat/{run.conversation_id}/send/', file=False, prompt='Cambia el título', content='Cambia el título',
            artifact_id=str(first.logical_id), base_revision=str(first.pk))
        self.assertEqual(stale.status_code, 409)
        response = self.send(url=f'/chat/{run.conversation_id}/send/', file=False, prompt='Cambia el título', content='Cambia el título',
            artifact_id=str(artifact.logical_id), base_revision=str(artifact.pk))
        revision = AgentExecution.objects.get(pk=response.json()['agent']['id'])
        token, attempt = self.claimed(revision)
        jobs.cancel(self.user, revision.pk)
        with self.assertRaises(AgentDenied):
            jobs.bootstrap(token, attempt, {})
        self.assertEqual(jobs.latest(run.conversation, first.logical_id), artifact)

    def test_unusable_research_does_not_end_with_unexplained_success(self):
        prompt = 'Crea un informe con fuentes'
        response = self.send(file=False, prompt=prompt, content=prompt)
        run = AgentExecution.objects.get(pk=response.json()['agent']['id'])
        token, attempt = self.claimed(run)
        jobs.fail(token, attempt, {})
        snapshot = self.client.get(f'/chat/agent/runs/{run.pk}/').json()
        self.assertEqual(snapshot['run']['status'], 'failed')
        self.assertEqual(snapshot['result']['error_type'], 'agent_research_no_sources')
        self.assertFalse(snapshot['artifacts'])

    def test_wrong_format_unconsulted_references_failed_bytes_and_company_isolation(self):
        run, token, attempt, call, artifact, payload = self.create()
        title = 'Análisis de las implicaciones de la computación cuántica para la ciberseguridad de la organización'
        long_title = self.issue_tool(token, attempt, 'build_document', {**call['args'], 'title': title})
        self.assertEqual(jobs.local_fields(token, attempt, long_title)['title'], title)
        wrong = self.issue_tool(token, attempt, 'create_mock_excel', {'title': 'Inventado', 'rows': [
            {'item': str(n), 'amount_cents': 100} for n in range(3)]})
        with self.assertRaisesMessage(AgentDenied, 'tool_does_not_match_request'):
            jobs.local_fields(token, attempt, wrong)
        broken = jobs.candidate(token, attempt, {**payload, 'operation_id': str(uuid.uuid4()),
            'document_bundle': base64.b64encode(b'broken document').decode()})
        self.assertFalse(AgentArtifact.objects.get(pk=broken['reference']).verification['automated_passed'])
        with self.assertRaises(AgentDenied):
            self.approve(run, token, attempt, artifact)
        from django.contrib.auth.models import User
        from front.models import Profile, Company
        other = User.objects.create_user(username='other-company-viewer')
        Profile.objects.create(user=other, company=Company.objects.create(name='Other isolated test company'))
        self.client.force_login(other)
        self.assertEqual(self.client.get(f'/chat/agent/artifacts/{artifact.pk}/pages/1/').status_code, 404)
        self.assertEqual(self.client.get(f'/chat/agent/artifacts/{artifact.pk}/download/?format=pdf').status_code, 404)
