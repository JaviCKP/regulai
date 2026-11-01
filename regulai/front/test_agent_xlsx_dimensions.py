"""Actual XLSX bytes: missing/stale dimensions and pre-iteration resource bounds."""
import io
from pathlib import Path
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

from django.test import SimpleTestCase
from openpyxl import load_workbook

from front.agents.artifact_formats import inspect_file
from front.agents.professional_artifacts import _worksheet_extent


FIXTURE = Path(__file__).resolve().parents[2] / 'evals/corpora/p0-conversations-20261016/fixtures/inventory.xlsx'
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def changed_sheets(change):
    output = io.BytesIO()
    with zipfile.ZipFile(FIXTURE) as source, zipfile.ZipFile(output, 'w') as target:
        for name in source.namelist():
            raw = source.read(name)
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root = ET.fromstring(raw)
                change(root)
                raw = ET.tostring(root, encoding='utf-8')
            target.writestr(name, raw)
    return output.getvalue()


class XlsxDimensionTests(SimpleTestCase):
    def test_original_streaming_fixture_reads_every_sheet(self):
        book = load_workbook(FIXTURE, read_only=True)
        try:
            self.assertTrue(all(sheet.max_row is None and sheet.max_column is None for sheet in book))
        finally:
            book.close()
        proof, text, preview = inspect_file(FIXTURE.read_bytes(), 'xlsx')
        self.assertEqual(proof['coverage']['status'], 'complete')
        self.assertEqual({sheet['name']: (sheet['rows'], sheet['columns']) for sheet in preview['sheets']},
                         {'Productos': (4, 3), 'Movimientos': (6, 3)})
        self.assertIn('C | 0 | 3', text)
        self.assertIn('A | salida | 5', text)

    def test_stale_dimension_does_not_hide_later_cells(self):
        def stale(root):
            root.insert(0, ET.Element(NS + 'dimension', {'ref': 'A1:A1'}))
        _, text, preview = inspect_file(changed_sheets(stale), 'xlsx')
        self.assertIn('C | 0 | 3', text)
        self.assertIn('A | salida | 5', text)
        self.assertEqual(preview['sheet_count'], 2)

    def test_sparse_limit_is_checked_before_iterating(self):
        def sparse(root):
            data = root.find(NS + 'sheetData')
            data.clear()
            row = ET.SubElement(data, NS + 'row', {'r': '100001'})
            ET.SubElement(row, NS + 'c', {'r': 'A100001'})
        with patch('openpyxl.worksheet._read_only.ReadOnlyWorksheet.iter_rows',
                   side_effect=AssertionError('Must reject before iterating')):
            with self.assertRaisesRegex(ValueError, 'workbook_cell_limit'):
                inspect_file(changed_sheets(sparse), 'xlsx')

    def test_extent_boundary_and_implicit_coordinates(self):
        self.assertEqual(_worksheet_extent(f'<worksheet xmlns="{NS[1:-1]}"><sheetData>'
            '<row r="100000"><c r="A100000"/></row></sheetData></worksheet>'.encode()), (100000, 1))
        self.assertEqual(_worksheet_extent(f'<worksheet xmlns="{NS[1:-1]}"><sheetData>'
            '<row><c/><c/></row><row><c/></row></sheetData></worksheet>'.encode()), (2, 2))
        with self.assertRaisesRegex(ValueError, 'workbook_cell_limit'):
            _worksheet_extent(f'<worksheet xmlns="{NS[1:-1]}"><sheetData>'
                '<row r="50001"><c r="B50001"/></row></sheetData></worksheet>'.encode())
