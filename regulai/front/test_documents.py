"""Synthetic uploads exercise the public extraction interface without network calls."""
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, override_settings
import io
import zipfile

from docx import Document as WordDocument
from openpyxl import Workbook

from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from front.documents import extract_document


class DocumentTests(SimpleTestCase):
    def extract(self, name, content):
        return extract_document(SimpleUploadedFile(name, content))

    def test_text_keeps_late_evidence_and_hash_bound_fragments(self):
        content = ("ordinary text\n" * 900 + "RESTRICTED late evidence\n").encode()
        first = self.extract("notes.txt", content)
        second = self.extract("notes.txt", content)
        other = self.extract("notes.txt", b"Different content")
        self.assertEqual(first.status, "complete")
        self.assertIn("RESTRICTED late evidence", first.text)
        self.assertEqual(first.fragments[-1]["locator"], "line:901")
        self.assertEqual(first.fragments, second.fragments)
        self.assertNotEqual(first.fragments[0]["id"], other.fragments[0]["id"])
        self.assertEqual(first.units_read, 901)
        self.assertEqual(first.units_total, 901)
        metadata = first.metadata()
        self.assertEqual(metadata["digest"], first.digest)
        self.assertNotIn("text", metadata)
        self.assertNotIn("fragments", metadata)
        self.assertNotIn("RESTRICTED", str(metadata))

    def test_invalid_or_bounded_uploads_never_appear_complete(self):
        for name, content, status, warning in [
            ("legacy.xls", b"legacy", "unsupported", "unsupported_format"),
            ("image.png", b"\x89PNG\r\n\x1a\n", "unsupported", "unsupported_format"),
            ("fake.pdf", b"ordinary text", "unreadable", "signature_mismatch"),
            ("fake.txt", b"%PDF-1.4\n", "unreadable", "signature_mismatch"),
            ("binary.txt", b"abc\x00def", "unreadable", "invalid_text"),
            ("encoding.txt", b"\xff\xfe", "unreadable", "invalid_text"),
            ("large.txt", b"a" * (1024 * 1024 + 1), "unreadable", "file_size_limit"),
            ("long.md", b"a" * 200_001, "partial", "text_limit"),
        ]:
            with self.subTest(name=name):
                document = self.extract(name, content)
                self.assertEqual(document.status, status)
                self.assertIn(warning, document.warnings)
        for extension in ("md", "json", "py", "js"):
            with self.subTest(extension=extension):
                self.assertEqual(self.extract(f"source.{extension}", b"read as text").status, "complete")

    def pdf(self, texts, encrypted=False, image=False, compressed=False):
        writer = PdfWriter()
        for text in texts:
            writer.add_blank_page(width=200, height=200)
            page = writer.pages[-1]
            if text:
                font = DictionaryObject({
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                })
                page[NameObject("/Resources")] = DictionaryObject({
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                })
                stream = DecodedStreamObject()
                stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode())
                page[NameObject("/Contents")] = writer._add_object(stream.flate_encode() if compressed else stream)
            if image:
                bitmap = DecodedStreamObject()
                bitmap.set_data(b"\xff\xff\xff")
                bitmap.update({
                    NameObject("/Type"): NameObject("/XObject"),
                    NameObject("/Subtype"): NameObject("/Image"),
                    NameObject("/Width"): NumberObject(1),
                    NameObject("/Height"): NumberObject(1),
                    NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                    NameObject("/BitsPerComponent"): NumberObject(8),
                })
                page[NameObject("/Resources")][NameObject("/XObject")] = DictionaryObject({
                    NameObject("/Im1"): writer._add_object(bitmap),
                })
        if encrypted:
            writer.encrypt("synthetic-password")
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    def test_pdf_reports_pages_and_partial_or_unknown_coverage(self):
        readable = self.extract("report.pdf", self.pdf(["First evidence", "Late evidence"]))
        self.assertEqual(readable.status, "complete")
        self.assertEqual(readable.units_read, 2)
        self.assertEqual(readable.units_total, 2)
        self.assertEqual(readable.fragments[1]["locator"], "page:2")
        self.assertIn("Late evidence", readable.text)
        compressed = self.extract("compressed.pdf", self.pdf(["Compressed evidence"], compressed=True))
        self.assertEqual(compressed.status, "complete")
        self.assertIn("Compressed evidence", compressed.text)
        for content, status, warning in [
            (self.pdf(["Readable", None]), "partial", "page_without_text"),
            (self.pdf(["Hidden"], encrypted=True), "unreadable", "encrypted_document"),
            (b"%PDF-1.4\ncorrupt", "unreadable", "corrupt_document"),
            (self.pdf(["Visible"], image=True), "partial", "visual_coverage_unknown"),
            (self.pdf(["Page"] * 101), "partial", "page_limit"),
        ]:
            with self.subTest(warning=warning):
                document = self.extract("report.pdf", content)
                self.assertEqual(document.status, status)
                self.assertIn(warning, document.warnings)

    def test_docx_reads_body_tables_and_headers_with_paragraph_locators(self):
        word = WordDocument()
        word.add_paragraph("Body evidence")
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
        word.sections[0].header.paragraphs[0].text = "Header evidence"
        output = io.BytesIO()
        word.save(output)
        document = self.extract("letter.docx", output.getvalue())
        self.assertEqual(document.status, "complete")
        for evidence in ("Body evidence", "Table evidence", "Header evidence"):
            self.assertIn(evidence, document.text)
        self.assertEqual(document.fragments[0]["locator"], "word/document.xml/paragraph:1")
        self.assertTrue(any(fragment["locator"].startswith("word/header1.xml/") for fragment in document.fragments))
        for content, warning in [
            (b"PK\x03\x04bad archive", "corrupt_document"),
            (self.archive({"xl/workbook.xml": b"<book/>"}), "signature_mismatch"),
            (self.archive({"word/document.xml": b"x" * (33 * 1024 * 1024)}), "expansion_limit"),
            (self.archive({"word/document.xml": b"<!DOCTYPE root [<!ENTITY x 'secret'>]><root>&x;</root>"}), "unsafe_xml"),
        ]:
            with self.subTest(warning=warning):
                invalid = self.extract("invalid.docx", content)
                self.assertEqual(invalid.status, "unreadable")
                self.assertIn(warning, invalid.warnings)

    def archive(self, files, compression=zipfile.ZIP_DEFLATED):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression) as archive:
            for name, content in files.items():
                archive.writestr(name, content)
        return output.getvalue()

    def test_csv_preserves_quoted_cells_and_explicit_row_limit(self):
        document = self.extract("data.csv", b'name,note\nAda,"first\nsecond"\n')
        self.assertEqual(document.status, "complete")
        self.assertEqual(document.units_read, 2)
        self.assertEqual(document.units_total, 2)
        row = next(fragment for fragment in document.fragments if fragment["locator"] == "row:2/cells:A2-B2")
        self.assertIn("B2: first\nsecond", row["text"])
        limited = self.extract("rows.csv", b"value\n" * 10_001)
        self.assertEqual(limited.status, "partial")
        self.assertIn("unit_limit", limited.warnings)
        self.assertEqual(limited.units_read, 10_000)
        self.assertIsNone(limited.units_total)

    def test_xlsx_reads_hidden_sheets_and_keeps_cell_locations_without_executing_formulas(self):
        workbook = Workbook()
        workbook.active.title = "Evidence"
        workbook.active["A1"] = "Public text"
        workbook.active["B3"] = "Late evidence"
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "Hidden restricted evidence"
        output = io.BytesIO()
        workbook.save(output)
        document = self.extract("evidence.xlsx", output.getvalue())
        self.assertEqual(document.status, "complete")
        self.assertIn("Hidden restricted evidence", document.text)
        row = next(fragment for fragment in document.fragments if fragment["locator"] == "sheet:Evidence/row:3/cells:B3-B3")
        self.assertEqual(row["text"], "B3: Late evidence")
        hidden["A2"] = '=WEBSERVICE("https://example.invalid/never-called")'
        output = io.BytesIO()
        workbook.save(output)
        formula = self.extract("formula.xlsx", output.getvalue())
        self.assertEqual(formula.status, "partial")
        self.assertIn("formula_coverage_unknown", formula.warnings)
        self.assertIn("=WEBSERVICE", formula.text)

    @override_settings(DOCUMENT_PARSER_TIMEOUT=0.0001)
    def test_parser_time_limit_stops_work_and_returns_no_false_coverage(self):
        document = self.extract("timeout.pdf", self.pdf(["Readable document"]))
        self.assertEqual(document.status, "unreadable")
        self.assertEqual(document.warnings, ["parser_timeout"])
        self.assertEqual(document.units_read, 0)
        self.assertIsNone(document.units_total)
        self.assertEqual(document.text, "")

    def test_empty_documents_and_unsupported_office_content_are_explicit(self):
        for name, content, warning in [
            ("empty.txt", b"", "empty_document"),
            ("ratio.docx", self.archive({"word/document.xml": b"x" * 1_000_000}), "expansion_limit"),
            ("utf16.docx", self.archive({"word/document.xml": "<!DOCTYPE root [<!ENTITY x 'secret'>]><root>&x;</root>".encode("utf-16")}), "unsafe_xml"),
            ("fake.docx", self.archive({"word/document.xml": b"<ordinary>text</ordinary>"}), "signature_mismatch"),
            ("macro.docx", self.archive({"word/document.xml": b"<root/>", "word/vbaProject.bin": b"macro"}), "unsupported_active_content"),
            ("name.docx", self.archive({"word/document.xml": b"<root/>", "word/header" + "x" * 300 + ".xml": b"<root/>"}), "invalid_archive"),
        ]:
            with self.subTest(name=name):
                document = self.extract(name, content)
                self.assertEqual(document.status, "unreadable")
                self.assertIn(warning, document.warnings)

    def test_xlsx_false_dimensions_cannot_hide_late_cells_and_limits_remain_explicit(self):
        workbook = Workbook()
        workbook.active["A1"] = "First"
        workbook.active["B3"] = "Restricted late cell"
        output = io.BytesIO()
        workbook.save(output)
        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source:
            files = {name: source.read(name) for name in source.namelist()}
        sheet = "xl/worksheets/sheet1.xml"
        files[sheet] = files[sheet].replace(b'ref="A1:B3"', b'ref="A1:A1"')
        document = self.extract("false-dimensions.xlsx", self.archive(files))
        self.assertEqual(document.status, "complete")
        self.assertIn("Restricted late cell", document.text)
        for name, changed, warning in [
            (sheet, files[sheet].replace(b'r="B3"', b'r="IW3"'), "column_limit"),
            ("xl/workbook.xml", files["xl/workbook.xml"].replace(b'name="Sheet"', b'name="' + b'x' * 32 + b'"'), "sheet_name_limit"),
        ]:
            with self.subTest(warning=warning):
                invalid = dict(files, **{name: changed})
                bounded = self.extract("bounded.xlsx", self.archive(invalid))
                self.assertEqual(bounded.status, "unreadable")
                self.assertIn(warning, bounded.warnings)

    def test_document_digest_uses_original_bytes(self):
        document = self.extract("original.txt", b"abc")
        self.assertEqual(document.digest, "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad")
        self.assertEqual(document.file_size, 3)
        self.assertEqual(document.text, "abc")

    def test_office_xml_complexity_and_external_references_are_bounded(self):
        namespace = b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        for xml in (
            b"<w:document " + namespace + b">" + b"<w:p/>" * 100_001 + b"</w:document>",
            b"<w:document " + namespace + b">" + b"<w:body>" * 129 + b"</w:body>" * 129 + b"</w:document>",
        ):
            with self.subTest(size=len(xml)):
                document = self.extract("structure.docx", self.archive({"word/document.xml": xml}, zipfile.ZIP_STORED))
                self.assertEqual(document.status, "unreadable")
                self.assertIn("xml_structure_limit", document.warnings)
        word = WordDocument()
        word.add_paragraph("Visible text")
        word.part.relate_to(
            "https://example.invalid/never-called",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        output = io.BytesIO()
        word.save(output)
        external = self.extract("external.docx", output.getvalue())
        self.assertEqual(external.status, "complete")
        self.assertNotIn("external_content_unread", external.warnings)

    def test_docx_tabs_and_breaks_do_not_join_separate_words(self):
        word = WordDocument()
        word.add_paragraph("First\tSecond\nThird")
        output = io.BytesIO()
        word.save(output)
        self.assertEqual(self.extract("breaks.docx", output.getvalue()).text, "First\tSecond\nThird")
