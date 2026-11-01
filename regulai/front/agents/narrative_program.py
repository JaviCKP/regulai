"""Offline narrative composition. No network, model SDK, credentials or database."""
import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import zipfile


def block_text(block, language='es'):
    labels = {'es': {'analysis': 'Análisis', 'estimate': 'Estimación', 'example': 'Ejemplo ficticio',
                    'recommendation': 'Recomendación'},
              'en': {'analysis': 'Analysis', 'estimate': 'Estimate', 'example': 'Fictional example',
                    'recommendation': 'Recommendation'}}
    label = labels.get(language, labels['es']).get(block['kind'])
    text = (label + '. ' if label else '') + block['text']
    if block.get('justification'):
        text += (' Justificación: ' if language == 'es' else ' Rationale: ') + block['justification']
    if block.get('sources'):
        text += ' ' + ' '.join('[' + source + ']' for source in block['sources'])
    return text


def reference_text(source):
    return f"[{source['id']}] {source['title']} — {source['publisher']}. {source['url']} ({source['retrieved_at'][:10]})."


def paragraphs(content):
    result = [content['title']]
    for page in content['pages']:
        result.append(page['heading'])
        result += [block_text(block, content['language']) for block in page['blocks']]
    if content.get('limitations'):
        result += [('Alcance y límites' if content['language'] == 'es' else 'Scope and limitations'), *content['limitations']]
    if content['references']:
        result += [('Referencias consultadas' if content['language'] == 'es' else 'Sources consulted')]
        result += [reference_text(source) for source in content['references']]
    return result


def build_docx(content, output):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.05), Cm(1.85)
    section.left_margin = section.right_margin = Cm(2.1)
    section.header_distance = section.footer_distance = Cm(.85)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Liberation Sans', Pt(11)
    style.font.color.rgb = RGBColor.from_string('263C42')
    style.paragraph_format.line_spacing = 1.16
    style.paragraph_format.space_after = Pt(9)
    style.paragraph_format.widow_control = True
    for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12)]:
        target = doc.styles[name]
        target.font.name, target.font.size = 'Liberation Serif', Pt(size)
        target.font.color.rgb = RGBColor.from_string('163C43')
        target.paragraph_format.space_before = Pt(10)
        target.paragraph_format.space_after = Pt(10)
        target.paragraph_format.keep_with_next = True
    # Remove the starter template's decorative borders and theme font overrides.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for font in doc.styles.element.iter(qn('w:rFonts')):
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            font.attrib.pop(qn('w:' + key), None)
    doc.core_properties.author = 'Regulay'
    doc.core_properties.title = content['title']
    doc.core_properties.subject = 'Informe' if content['language'] == 'es' else 'Report'
    doc.core_properties.language = content['language']
    for element in list(doc.core_properties._element):
        if element.tag in (qn('dcterms:created'), qn('dcterms:modified')):
            doc.core_properties._element.remove(element)
    header = section.header.paragraphs[0]
    header.text = 'REGULAY'
    header.runs[0].font.name, header.runs[0].font.size = 'Liberation Sans', Pt(9)
    header.runs[0].bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Página ' if content['language'] == 'es' else 'Page ')
    for field, separator in [('PAGE', ' / '), ('NUMPAGES', '')]:
        element = OxmlElement('w:fldSimple')
        element.set(qn('w:instr'), field)
        footer._p.append(element)
        if separator:
            footer.add_run(separator)
    doc.add_paragraph(content['title'], 'Title')
    for index, page in enumerate(content['pages']):
        heading = doc.add_paragraph(page['heading'], 'Heading 1')
        if index:
            heading.paragraph_format.page_break_before = True
        for block in page['blocks']:
            doc.add_paragraph(block_text(block, content['language']))
    if content.get('limitations'):
        doc.add_paragraph('Alcance y límites' if content['language'] == 'es' else 'Scope and limitations', 'Heading 2')
        for limit in content['limitations']:
            doc.add_paragraph(limit)
    if content['references']:
        doc.add_paragraph('Referencias consultadas' if content['language'] == 'es' else 'Sources consulted', 'Heading 2')
        for source in content['references']:
            paragraph = doc.add_paragraph(reference_text(source))
            paragraph.paragraph_format.space_after = Pt(5)
            for run in paragraph.runs:
                run.font.size = Pt(8)
    doc.save(output)


def build_pptx(content, output):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.3333), Inches(7.5)
    for index, page in enumerate(content['pages']):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('F5F8F7')
        def box(x, y, w, h, text, size, color='193B45', bold=False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.name, paragraph.font.size = 'Liberation Sans', Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            return frame
        box(.6, .22, 12.1, .3, 'REGULAY', 11, bold=True)
        box(.6, .72, 12.1, 1.1, page['heading'], 28, bold=True)
        body = '\n\n'.join(block_text(b, content['language']) for b in page['blocks'])
        if index == len(content['pages']) - 1:
            if content.get('limitations'):
                body += '\n\n' + ('Alcance y límites' if content['language'] == 'es' else 'Scope and limitations') + '\n' + '\n'.join(content['limitations'])
            if content['references']:
                body += '\n\n' + ('Referencias consultadas' if content['language'] == 'es' else 'Sources consulted') + '\n' + '\n'.join(reference_text(s) for s in content['references'])
        # Short narrative slides need presentation-sized type; denser closing
        # notes retain room for the actual limitations and source references.
        size = 26 if len(body.split()) <= 65 else 20
        box(.65, 1.95, 12, 4.6, body.strip(), size)
        box(.65, 6.88, 11.4, .35, content['title'], 10)
        box(12, 6.88, .8, .35, str(index + 1), 10)
    presentation.core_properties.title, presentation.core_properties.author = content['title'], 'Regulay'
    presentation.save(output)


def edit(source, content, output, extension='docx'):
    """Change text nodes only; preserve styles, geometry and every other part."""
    from lxml import etree
    with zipfile.ZipFile(source) as archive:
        parts = {item.filename: archive.read(item) for item in archive.infolist()}
    if extension == 'docx':
        from docx import Document
        old = Document(source)
        values = paragraphs(content)
        if len(old.paragraphs) != len(values):
            raise ValueError('report_revision_structure_changed')
        tree = etree.fromstring(parts['word/document.xml'])
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        elements = tree.findall('w:body/w:p', ns)
        for element, before, after in zip(elements, old.paragraphs, values):
            if before.text != after:
                nodes = element.findall('.//w:t', ns)
                if not nodes:
                    raise ValueError('report_revision_text_missing')
                nodes[0].text = after
                for node in nodes[1:]:
                    node.text = ''
        parts['word/document.xml'] = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
    else:
        from pptx import Presentation
        old = Presentation(source)
        if len(old.slides) != len(content['pages']):
            raise ValueError('report_revision_structure_changed')
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
              'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
        for index, (slide, page) in enumerate(zip(old.slides, content['pages']), 1):
            part = f'ppt/slides/slide{index}.xml'
            tree = etree.fromstring(parts[part])
            shapes = tree.findall('.//p:sp', ns)
            changed = False
            # Only heading/body edits are supported for narrative decks.
            values = [(1, page['heading'])]
            if len(page['blocks']) == 1:
                before = slide.shapes[2].text.replace('\v', '\n')
                after = block_text(page['blocks'][0], content['language'])
                suffix = before.partition('\n\n')[2]
                values.append((2, after + ('\n\n' + suffix if suffix else '')))
            for shape_index, value in values:
                element = shapes[shape_index]
                existing = slide.shapes[shape_index].text.replace('\v', '\n')
                if existing == value:
                    continue
                changed = True
                lines = value.split('\n')
                text_nodes = element.findall('.//a:t', ns)
                existing_lines = existing.split('\n')
                if len(lines) != len(existing_lines) or len(text_nodes) != len([line for line in lines if line]):
                    raise ValueError('report_revision_text_structure')
                for node, value in zip(text_nodes, [line for line in lines if line]):
                    node.text = value
            if changed:
                parts[part] = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            target.writestr(item, parts[item.filename])


def render(output, extension='docx', *, converter=None, rasterizer=None):
    converter = converter or shutil.which('libreoffice') or shutil.which('soffice')
    rasterizer = rasterizer or shutil.which('pdftoppm')
    if not converter or not rasterizer:
        raise ValueError('document_converter_unavailable')
    target = output / ('final.' + extension)
    profile = (output / 'lo-profile').resolve().as_uri()
    export = 'writer_pdf_Export' if extension == 'docx' else 'impress_pdf_Export'
    result = subprocess.run([str(converter), '-env:UserInstallation=' + profile, '--headless',
        '--convert-to', 'pdf:' + export, '--outdir', str(output), str(target)], capture_output=True, timeout=90)
    pdf = output / 'final.pdf'
    if result.returncode or not pdf.is_file() or not pdf.stat().st_size:
        raise ValueError('document_converter_failed')
    subprocess.run([str(rasterizer), '-r', '100', '-png', str(pdf), str(output / 'render')],
                   capture_output=True, check=True, timeout=90)
    pages = sorted(output.glob('render-*.png'), key=lambda p: int(p.stem.rsplit('-', 1)[-1]))
    for number, page in enumerate(pages, 1):
        page.replace(output / f'page-{number}.png')


def seal(output, extension='docx'):
    from pypdf import PdfReader
    source, pdf = (output / ('final.' + extension)), output / 'final.pdf'
    count = len(PdfReader(pdf).pages)
    pages = [output / f'page-{n}.png' for n in range(1, count + 1)]
    sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {extension + '_sha256': sha(source), 'pdf_sha256': sha(pdf), 'pages': [sha(p) for p in pages]}
    (output / 'render.json').write_text(json.dumps(manifest), encoding='utf-8')
    with zipfile.ZipFile(output / 'candidate.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in [source, pdf, output / 'render.json', *pages]:
            archive.write(path, path.name)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('operation', choices=['build', 'edit', 'render', 'seal'])
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--fields', type=Path)
    parser.add_argument('--source', type=Path)
    parser.add_argument('--extension', choices=['docx', 'pptx'], default='docx')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.operation in ('build', 'edit'):
        content = json.loads(args.fields.read_text(encoding='utf-8'))['narrative']
        if args.operation == 'edit':
            edit(args.source, content, args.output / ('final.' + args.extension), args.extension)
        else:
            (build_docx if args.extension == 'docx' else build_pptx)(content, args.output / ('final.' + args.extension))
    elif args.operation == 'render':
        render(args.output, args.extension)
    else:
        seal(args.output, args.extension)


if __name__ == '__main__':
    main()
