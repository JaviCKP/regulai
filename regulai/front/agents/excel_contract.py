"""Trusted, bounded XLSX verification. No model, SDK, or document execution here."""
import csv
from decimal import Decimal
import hashlib
import io
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

MAX_BYTES = 2_000_000
MAX_ROWS = 1000
HEADERS = ['item', 'amount_cents']


def mock_source(rows):
    """Validate generated values before storing an immutable synthetic source."""
    if not isinstance(rows, list) or not 3 <= len(rows) <= MAX_ROWS:
        raise ValueError('mock_rows')
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, lineterminator='\n')
    writer.writerow(HEADERS)
    for row in rows:
        if (not isinstance(row, dict) or set(row) != set(HEADERS)
                or type(row['item']) is not str or type(row['amount_cents']) is not int):
            raise ValueError('mock_row')
        writer.writerow([row['item'], row['amount_cents']])
    data = stream.getvalue().encode('utf-8')
    input_rows(data, 'csv')
    return data


def mock_title(title):
    return valid_title('Datos ficticios · ' + valid_title(title))


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def workbook(data, *, values=False):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('excel_size')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (len(entries) > 100 or sum(x.file_size for x in entries) > 10_000_000
                or any(x.flag_bits & 1 for x in entries)):
            raise ValueError('excel_archive')
        for entry in entries:
            name = entry.filename.lower()
            if any(part in name for part in ('vbaproject', 'externallinks/', 'embeddings/', 'connections.xml')):
                raise ValueError('excel_active_content')
            if name.endswith('.rels'):
                relations = ET.fromstring(archive.read(entry))
                if any(row.attrib.get('TargetMode', '').lower() == 'external' for row in relations):
                    raise ValueError('excel_external_link')
    book = load_workbook(io.BytesIO(data), data_only=values, keep_links=False)
    if (len(book.worksheets) != 1 or book.active.max_row > MAX_ROWS + 4 or book.active.max_column > 3):
        raise ValueError('excel_shape')
    return book


def input_rows(data, kind):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('input_size')
    if kind == 'csv':
        try:
            table = list(csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), strict=True))
        except csv.Error as error:
            raise ValueError('input_csv') from error
    elif kind == 'xlsx':
        book = workbook(data)
        if any(c.data_type == 'f' for row in book.active for c in row):
            raise ValueError('input_formulas_unsupported')
        table = list(book.active.values)
        book.close()
    else:
        raise ValueError('input_format')
    if not table or list(table[0]) != HEADERS or not 1 <= len(table) - 1 <= MAX_ROWS:
        raise ValueError('input_columns')
    rows = []
    for row in table[1:]:
        if len(row) != 2 or not isinstance(row[0], str) or not 1 <= len(row[0]) <= 200:
            raise ValueError('input_item')
        value = row[1]
        if isinstance(value, bool) or not re.fullmatch(r'-?[0-9]{1,12}', str(value)):
            raise ValueError('input_cents')
        rows.append((row[0], int(value)))
    if abs(sum(v for _, v in rows)) > 999_999_999_999:
        raise ValueError('input_total')
    return rows


def valid_title(title):
    if not isinstance(title, str) or not title.strip() or len(title) > 120 or any(ord(c) < 32 for c in title):
        raise ValueError('excel_title')
    return title


def _style_signature(sheet):
    return [(c.coordinate, str(c.font), str(c.fill), str(c.border), str(c.alignment),
             c.number_format, str(c.protection)) for row in sheet for c in row]


def verify(data, source, kind, *, title, base=None):
    """Check final bytes against immutable input, independently of pandas/LO."""
    title = valid_title(title)
    rows = input_rows(source, kind)
    book, cached = workbook(data), workbook(data, values=True)
    sheet, values = book.active, cached.active
    if sheet.title != 'Resumen' or sheet['A1'].value != title or sheet['A1'].data_type == 'f':
        raise ValueError('excel_title_or_sheet')
    expected = {'A1': title, 'A3': 'Concepto', 'B3': 'Importe (céntimos)', 'C3': 'Importe (€)'}
    total = sum(amount for _, amount in rows)
    for n, (item, amount) in enumerate(rows, 4):
        expected.update({f'A{n}': item, f'B{n}': amount, f'C{n}': f'=B{n}/100'})
        if Decimal(str(values[f'C{n}'].value)) != Decimal(amount) / 100:
            raise ValueError('excel_cached_row')
    last = len(rows) + 4
    expected.update({f'A{last}': 'Total', f'B{last}': f'=SUM(B4:B{last-1})', f'C{last}': f'=B{last}/100'})
    formula_cells = {f'C{n}' for n in range(4, last + 1)} | {f'B{last}'}
    for row in sheet:
        for cell in row:
            if cell.value != expected.get(cell.coordinate):
                raise ValueError('excel_unexpected_cell')
            if cell.hyperlink or cell.data_type == 'e':
                raise ValueError('excel_active_or_error_cell')
            if (cell.data_type == 'f') != (cell.coordinate in formula_cells):
                raise ValueError('excel_formula_injection_or_missing')
    if any(sheet[key].value != value for key, value in expected.items()):
        raise ValueError('excel_missing_cell')
    if (values[f'B{last}'].value != total
            or Decimal(str(values[f'C{last}'].value)) != Decimal(total) / 100):
        raise ValueError('excel_cached_total')
    if sheet.freeze_panes != 'A4' or not sheet.auto_filter.ref or sheet.column_dimensions['A'].width < 20:
        raise ValueError('excel_layout')
    if base:
        previous = workbook(base)
        old = previous.active
        if ([(c.coordinate, c.value) for row in old for c in row if c.coordinate != 'A1'] !=
                [(c.coordinate, c.value) for row in sheet for c in row if c.coordinate != 'A1']
                or _style_signature(old) != _style_signature(sheet)
                or str(old.column_dimensions) != str(sheet.column_dimensions)
                or old.freeze_panes != sheet.freeze_panes or old.auto_filter.ref != sheet.auto_filter.ref
                or str(old.merged_cells) != str(sheet.merged_cells)):
            raise ValueError('excel_revision_changed_calculations_or_style')
        previous.close()
    preview = {'title': title, 'headers': ['Concepto', 'Céntimos', 'Euros'],
        'rows': [[item, amount, str(Decimal(amount) / 100)] for item, amount in rows[:20]],
        'total_cents': total, 'total_euros': str(Decimal(total) / 100),
        'rows_shown': min(20, len(rows)), 'rows_total': len(rows)}
    book.close()
    cached.close()
    return {'sha256': sha256(data), 'rows': len(rows), 'total_cents': total,
            'formulas': len(rows) + 2, 'checks': ['structure', 'formulas', 'cached_results',
            'independent_integer_total', 'layout', 'unchanged_revision' if base else 'new_workbook']}, preview
