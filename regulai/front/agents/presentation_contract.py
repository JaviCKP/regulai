"""Independent server-side inspection of final PPTX, embedded data and renders."""
import base64
from decimal import Decimal
import io
import json
from pathlib import Path
import re
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from pypdf.generic import ArrayObject

from .excel_contract import input_rows, sha256, workbook

TEMPLATE = Path(__file__).with_name('runtime')/'templates'/'corporate.pptx'
LIMIT = 20_000_000
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'


def valid_text(value, field):
    if (field not in ('title', 'body') or not isinstance(value, str) or not value.strip()
            or len(value) > (70 if field == 'title' else 340)
            or any(ord(c) < 32 for c in value) or any(s in value for s in ('{{', '}}'))):
        raise ValueError('presentation_text')
    return value


def source_rows(source, kind):
    rows = input_rows(source, kind)
    if not 1 <= len(rows) <= 6 or any(len(item) > 30 or ',' in item or '\\' in item or
            any(ord(c) < 32 for c in item) or '{{' in item or '}}' in item for item, _ in rows):
        raise ValueError('presentation_input_limit')
    return rows


def package(data):
    if not data or len(data) > LIMIT:
        raise ValueError('presentation_size')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (len(entries) > 160 or len({x.filename for x in entries}) != len(entries)
                or sum(x.file_size for x in entries) > LIMIT or any(x.flag_bits & 1 or
                '..' in x.filename.split('/') or x.filename.startswith('/') for x in entries)):
            raise ValueError('presentation_archive')
        parts = {x.filename: archive.read(x) for x in entries}
    for name, value in parts.items():
        if any(t in name.lower() for t in ('vbaproject', 'activex', 'oleobject')):
            raise ValueError('presentation_active_content')
        if name.endswith(('.xml', '.rels')):
            if b'<!DOCTYPE' in value or b'<!ENTITY' in value:
                raise ValueError('presentation_xml')
            tree = etree.fromstring(value, etree.XMLParser(resolve_entities=False, no_network=True))
            if name.endswith('.rels') and any(x.get('TargetMode', '').lower() == 'external' for x in tree):
                raise ValueError('presentation_external_relationship')
            if any(x.tag in (A+'hlinkClick', A+'hlinkMouseOver', P+'oleObj', P+'control', P+'timing') for x in tree.iter()):
                raise ValueError('presentation_active_content')
    if 'ppt/presentation.xml' not in parts:
        raise ValueError('presentation_missing')
    return parts


def fields(data):
    package(data)
    deck = Presentation(io.BytesIO(data))
    result = []
    for slide in deck.slides:
        item = {}
        for key in ('Title', 'Body'):
            matches = [s for s in slide.shapes if s.name == key and s.has_text_frame]
            if len(matches) != 1:
                raise ValueError('presentation_fields')
            item[key.lower()] = matches[0].text
        result.append(item)
    return result


def expected_fields(title, section):
    return [{'title': title, 'body': 'Datos económicos de España\nDesglose y revisión de importes'},
            {'title': 'Detalle de importes', 'body': 'Fuente: archivo aportado · Suma calculada en céntimos'},
            {'title': 'Distribución por concepto', 'body': section}]


def masked_slide(data, field):
    tree = etree.fromstring(data)
    for shape in tree.iter(P+'sp'):
        name = shape.find('./'+P+'nvSpPr/'+P+'cNvPr')
        if name is not None and name.get('name') == ('Title' if field == 'title' else 'Body'):
            for t in shape.iter(A+'t'):
                t.text = ''
    def semantic(node):
        return node.tag, sorted(node.attrib.items()), node.text or '', [semantic(n) for n in node]
    return semantic(tree)


def verify_pptx(raw, source, kind, expected, *, base=None, changed=None, template='new'):
    parts = package(raw)
    deck = Presentation(io.BytesIO(raw))
    if len(deck.slides) != 3 or len(deck.slide_masters) != 1 or fields(raw) != expected:
        raise ValueError('presentation_content')
    rows = source_rows(source, kind)
    tables, charts = [], []
    for slide in deck.slides:
        for s in slide.shapes:
            if s.left < 0 or s.top < 0 or s.left+s.width > deck.slide_width or s.top+s.height > deck.slide_height:
                raise ValueError('presentation_out_of_canvas')
            if s.has_text_frame:
                if '{{' in s.text or '}}' in s.text or '\ufffd' in s.text:
                    raise ValueError('presentation_template_residue')
                for p in s.text_frame.paragraphs:
                    for run in p.runs:
                        font = run.font.name or p.font.name
                        if font not in ('Liberation Sans', 'Liberation Serif'):
                            raise ValueError('presentation_font')
            if s.has_table:
                tables.append(s.table)
            if s.has_chart:
                charts.append(s.chart)
    money = lambda v: format(Decimal(v)/100, '.2f').replace('.', ',') + '\u00a0€'
    cover = {s.name: s.text for s in deck.slides[0].shapes if s.has_text_frame}
    if cover.get('Total') != money(sum(v for _, v in rows)) or cover.get('Count') != f'{len(rows)} conceptos':
        raise ValueError('presentation_cover_values')
    wanted = [['Concepto', 'Importe (€)']] + [[name, money(v)] for name, v in rows] + [['Total', money(sum(v for _, v in rows))]]
    if len(tables) != 1 or [[c.text for c in r.cells] for r in tables[0].rows] != wanted or len(charts) != 1:
        raise ValueError('presentation_native_table_chart')
    chart = charts[0]
    if (len(chart.series) != 1 or [c.label for c in chart.plots[0].categories] != [r[0] for r in rows]
            or [Decimal(str(v)) for v in chart.series[0].values] != [Decimal(v)/100 for _, v in rows]):
        raise ValueError('presentation_chart_values')
    labels = chart._chartSpace.findall('.//' + C + 'dLbl')
    if [''.join(n.itertext()) for label in labels for n in label.findall('.//' + A + 't')] != [money(v) for _, v in rows]:
        raise ValueError('presentation_chart_labels')
    # Verify the actual workbook reached by c:externalData, not merely a ZIP filename.
    blob = chart.part.chart_workbook.xlsx_part.blob
    book = workbook(blob)
    if list(book.active.values) != [('Concepto', 'Importe (€)'), *[(n, v/100) for n, v in rows]]:
        raise ValueError('presentation_embedded_data')
    if any(c.data_type == 'f' for row in book.active for c in row):
        raise ValueError('presentation_embedded_formula')
    book.close()
    if base:
        old = package(base)
        slide, field = changed
        name = f'ppt/slides/slide{slide}.xml'
        if (set(parts) != set(old) or any(parts[n] != old[n] for n in parts if n != name)
                or masked_slide(parts[name], field) != masked_slide(old[name], field)):
            raise ValueError('presentation_revision_changed_other_content')
    elif template == 'corporate':
        original = package(TEMPLATE.read_bytes())
        if any(parts.get(n) != v for n, v in original.items() if not n.startswith('ppt/slides/') and n != '[Content_Types].xml'):
            raise ValueError('presentation_template_changed')
    return rows


def verify_bundle(encoded, source, kind, expected, *, base=None, changed=None, template='new'):
    if len(encoded) > 2*LIMIT:
        raise ValueError('presentation_bundle_size')
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(encoded, validate=True))) as archive:
        entries = archive.infolist()
        names = {'final.pptx', 'final.pdf', 'render.json', 'page-1.png', 'page-2.png', 'page-3.png'}
        if len(entries) != 6 or {x.filename for x in entries} != names or sum(x.file_size for x in entries) > LIMIT:
            raise ValueError('presentation_bundle')
        bundle = {n: archive.read(n) for n in names}
    raw, pdf = bundle['final.pptx'], bundle['final.pdf']
    rows = verify_pptx(raw, source, kind, expected, base=base, changed=changed, template=template)
    reader = PdfReader(io.BytesIO(pdf), strict=True)
    if reader.is_encrypted or len(reader.pages) != 3:
        raise ValueError('presentation_pdf_pages')
    root = reader.trailer['/Root']
    action = root.get('/OpenAction')
    if action is not None:
        action = action.get_object()
        if not isinstance(action, ArrayObject) or len(action) != 5 or action[1] != '/XYZ':
            raise ValueError('presentation_pdf_action')
    if any(k in root for k in ('/AA', '/AcroForm')) or (root.get('/Names') and root['/Names'].get_object()):
        raise ValueError('presentation_pdf_action')
    for number, page in enumerate(reader.pages):
        text = re.sub(r'\s+', '', page.extract_text() or '')
        wanted = [*expected[number].values(), 'REGULAY']
        if number == 1:
            wanted += [n for n, _ in rows]
            wanted += [format(Decimal(v)/100, '.2f').replace('.', ',') for _, v in rows]
        if any(re.sub(r'\s+', '', v) not in text for v in wanted) or '\ufffd' in text or '/Annots' in page or '/AA' in page:
            raise ValueError('presentation_pdf_text')
        for font in page['/Resources'].get('/Font', {}).values():
            font = font.get_object()
            if '/DescendantFonts' in font:
                font = font['/DescendantFonts'][0].get_object()
            descriptor = font.get('/FontDescriptor')
            if not descriptor or not any(k in descriptor.get_object() for k in ('/FontFile', '/FontFile2', '/FontFile3')):
                raise ValueError('presentation_pdf_font')
    manifest = json.loads(bundle['render.json'])
    hashes = [sha256(bundle[f'page-{n}.png']) for n in range(1, 4)]
    if manifest != {'pptx_sha256': sha256(raw), 'pdf_sha256': sha256(pdf), 'pages': hashes}:
        raise ValueError('presentation_render_digest')
    for n in range(1, 4):
        with Image.open(io.BytesIO(bundle[f'page-{n}.png'])) as image:
            if image.format != 'PNG' or image.size != (1334, 750):
                raise ValueError('presentation_render_dimensions')
            image.verify()
    proof = {'sha256': sha256(raw), 'pdf_sha256': sha256(pdf), 'page_sha256': hashes, 'pages': 3,
        'rows': len(rows), 'total_cents': sum(v for _, v in rows), 'automated_passed': True, 'visual_review': 'required',
        'checks': ['native_text_table_chart', 'embedded_workbook', 'source_values', 'canvas_bounds',
                   'pdf_text_fonts', 'rendered_final_bytes', 'unchanged_revision' if base else template]}
    return raw, proof, json.dumps({n: base64.b64encode(v).decode() for n, v in bundle.items()})
