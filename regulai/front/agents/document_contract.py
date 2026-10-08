"""Bounded DOCX/PDF verification and private render bundles. No model access."""
import base64
from decimal import Decimal
import io
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from lxml import etree
from PIL import Image
from pypdf import PdfReader
from pypdf.generic import ArrayObject

from .excel_contract import input_rows, sha256, valid_title

MAX_DOCUMENT_BYTES = 20_000_000
MAX_PAGES = 40
TEMPLATE = Path(__file__).with_name('runtime') / 'templates' / 'corporate.docx'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def valid_section(value):
    if (not isinstance(value, str) or not value.strip() or len(value) > 1800
            or any(ord(c) < 32 for c in value) or '{{' in value or '}}' in value):
        raise ValueError('document_section')
    return value


def money(cents):
    return format(Decimal(cents) / 100, '.2f').replace('.', ',') + ' €'


def package(data):
    if not data or len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError('document_size')
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if (len(items) > 100 or len({x.filename for x in items}) != len(items)
                or sum(x.file_size for x in items) > MAX_DOCUMENT_BYTES
                or any(x.flag_bits & 1 or '..' in x.filename.split('/') for x in items)):
            raise ValueError('document_archive')
        parts = {x.filename: archive.read(x) for x in items}
    for name, content in parts.items():
        if any(key in name.lower() for key in ('vbaproject', 'embeddings/', 'activex/')):
            raise ValueError('document_active_content')
        if name.endswith(('.xml', '.rels')):
            if b'<!DOCTYPE' in content or b'<!ENTITY' in content:
                raise ValueError('document_xml')
            tree = etree.fromstring(content, etree.XMLParser(resolve_entities=False, no_network=True))
            if name.endswith('.rels') and any(x.get('TargetMode', '').lower() == 'external' for x in tree):
                raise ValueError('document_external_relationship')
            for node in tree.iter():
                if node.tag in (W + 'instrText', W + 'fldSimple'):
                    instruction = node.get(W + 'instr', node.text or '').strip()
                    if instruction not in ('PAGE', 'NUMPAGES'):
                        raise ValueError('document_field')
    if 'word/document.xml' not in parts:
        raise ValueError('document_missing_body')
    return parts


def fields(data):
    package(data)
    doc = Document(io.BytesIO(data))
    titles = [p for p in doc.paragraphs if p.style.name == 'Title']
    if len(titles) != 1:
        raise ValueError('document_title')
    paragraphs = doc.paragraphs
    index = next(i for i, p in enumerate(paragraphs) if p.text == 'Resumen ejecutivo')
    return {'title': titles[0].text, 'section': paragraphs[index + 1].text}


def _body_shape(data, *, clear_table=False, field=None):
    body = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
    paragraphs = list(body.iter(W + 'p'))
    for index, paragraph in enumerate(paragraphs):
        style = paragraph.find('./' + W + 'pPr/' + W + 'pStyle')
        is_title = style is not None and style.get(W + 'val') == 'Title'
        previous = ''.join(paragraphs[index - 1].itertext()) if index else ''
        is_section = previous == 'Resumen ejecutivo'
        if (field in (None, 'title') and is_title) or (field in (None, 'section') and is_section):
            for text in paragraph.iter(W + 't'):
                text.text = ''
    if clear_table:
        table = next(body.iter(W + 'tbl'))
        for row in list(table)[-1::-1]:
            if row.tag == W + 'tr':
                table.remove(row)
    return etree.tostring(body, method='c14n')


def verify_docx(data, source, kind, expected, *, base=None, changed=None):
    parts, template = package(data), package(TEMPLATE.read_bytes())
    if set(parts) != set(template):
        raise ValueError('document_parts')
    if any(parts[name] != template[name] for name in parts if name != 'word/document.xml'):
        raise ValueError('document_template_parts_changed')
    if fields(data) != expected:
        raise ValueError('document_fields')
    valid_title(expected['title'])
    valid_section(expected['section'])
    if _body_shape(parts['word/document.xml'], clear_table=True) != _body_shape(template['word/document.xml'], clear_table=True):
        raise ValueError('document_template_structure')
    doc = Document(io.BytesIO(data))
    rows = input_rows(source, kind)
    table = doc.tables[0]
    expected_table = [['Concepto', 'Importe (€)']] + [[item, money(cents)] for item, cents in rows]
    expected_table += [['Total', money(sum(cents for _, cents in rows))]]
    if [[c.text for c in row.cells] for row in table.rows] != expected_table:
        raise ValueError('document_table_values')
    if len(doc.sections) != 3 or len(doc.tables) != 1:
        raise ValueError('document_sections')
    xml_text = ''.join(etree.fromstring(parts['word/document.xml']).itertext())
    if '{{' in xml_text or '}}' in xml_text:
        raise ValueError('document_template_residue')
    if base:
        previous = package(base)
        if (set(previous) != set(parts) or any(previous[k] != v for k, v in parts.items() if k != 'word/document.xml')
                or _body_shape(previous['word/document.xml'], field=changed) != _body_shape(parts['word/document.xml'], field=changed)):
            raise ValueError('document_revision_changed_other_content')
    return rows


def verify_pdf(data, expected, rows):
    if not data or len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError('pdf_size')
    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted or not 3 <= len(reader.pages) <= MAX_PAGES:
        raise ValueError('pdf_pages')
    root = reader.trailer['/Root']
    opening = root.get('/OpenAction')
    if opening is not None:
        opening = opening.get_object()
        # Writer emits an initial page/zoom destination, not an executable action.
        if not isinstance(opening, ArrayObject) or len(opening) != 5 or opening[1] != '/XYZ':
            raise ValueError('pdf_active_content')
    names = root.get('/Names')
    if any(key in root for key in ('/AA', '/AcroForm')) or (names and names.get_object()):
        raise ValueError('pdf_active_content')
    texts = []
    for page in reader.pages:
        if '/AA' in page or '/Annots' in page:
            raise ValueError('pdf_annotations')
        text = page.extract_text() or ''
        if len(text.strip()) < 40 or 'REGULAY' not in text or 'Página' not in text:
            raise ValueError('pdf_blank_or_missing_header_footer')
        fonts = page['/Resources'].get('/Font', {})
        for font in fonts.values():
            font = font.get_object()
            if '/DescendantFonts' in font:
                font = font['/DescendantFonts'][0].get_object()
            descriptor = font.get('/FontDescriptor')
            if not descriptor or not any(k in descriptor.get_object() for k in ('/FontFile', '/FontFile2', '/FontFile3')):
                raise ValueError('pdf_font_not_embedded')
        texts.append(text)
    normalized = re.sub(r'\s+', '', '\n'.join(texts))
    required = [expected['title'], expected['section'], 'Resumen ejecutivo', 'Notas de revisión',
                money(sum(cents for _, cents in rows))]
    required += [value for item, cents in rows for value in (item, money(cents))]
    if any(re.sub(r'\s+', '', value) not in normalized for value in required) or '{{' in normalized or '\ufffd' in normalized:
        raise ValueError('pdf_missing_text_or_template_residue')
    return len(reader.pages)


def verify_bundle(encoded, source, kind, expected, *, format, base=None, changed=None):
    """Reopen final documents and PNGs; qualitative review remains a separate gate."""
    if format not in ('docx', 'pdf') or len(encoded) > MAX_DOCUMENT_BYTES * 2:
        raise ValueError('document_bundle_size')
    data = base64.b64decode(encoded, validate=True)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if len(items) > MAX_PAGES + 3 or sum(x.file_size for x in items) > MAX_DOCUMENT_BYTES:
            raise ValueError('document_bundle_size')
        if len({x.filename for x in items}) != len(items):
            raise ValueError('document_bundle_duplicate')
        parts = {x.filename: archive.read(x) for x in items}
    docx, pdf = parts['final.docx'], parts['final.pdf']
    rows = verify_docx(docx, source, kind, expected, base=base, changed=changed)
    count = verify_pdf(pdf, expected, rows)
    required = {'final.docx', 'final.pdf', 'render.json'} | {f'page-{n}.png' for n in range(1, count + 1)}
    if set(parts) != required:
        raise ValueError('document_render_pages')
    manifest = json.loads(parts['render.json'])
    if (manifest.get('docx_sha256') != sha256(docx) or manifest.get('pdf_sha256') != sha256(pdf)
            or manifest.get('pages') != [sha256(parts[f'page-{n}.png']) for n in range(1, count + 1)]):
        raise ValueError('document_render_digest')
    for n in range(1, count + 1):
        with Image.open(io.BytesIO(parts[f'page-{n}.png'])) as image:
            if image.format != 'PNG' or image.width != 827 or not 1168 <= image.height <= 1171:
                raise ValueError('document_render_dimensions')
            image.verify()
    raw = docx if format == 'docx' else pdf
    proof = {'sha256': sha256(raw), 'rows': len(rows), 'total_cents': sum(v for _, v in rows),
             'pages': count, 'docx_sha256': sha256(docx), 'pdf_sha256': sha256(pdf),
             'page_sha256': manifest['pages'], 'automated_passed': True, 'visual_review': 'required',
             'checks': ['template_parts', 'sections', 'table_values', 'pdf_text', 'embedded_fonts',
                        'rendered_final_bytes', 'unchanged_revision' if base else 'new_document']}
    return raw, proof, json.dumps({name: base64.b64encode(value).decode() for name, value in parts.items()})
