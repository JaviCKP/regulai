"""Narrative data contract and independent final-byte verification."""
import base64
from copy import deepcopy
import io
import json
import re
import zipfile

from docx import Document
from PIL import Image
from pypdf import PdfReader
from pypdf.generic import ArrayObject
from lxml import etree

from .document_contract import MAX_DOCUMENT_BYTES, package
from .excel_contract import sha256
from .narrative_program import paragraphs, block_text, reference_text


def text(value, maximum=5000):
    if (not isinstance(value, str) or not value.strip() or len(value) > maximum
            or any(ord(c) < 32 for c in value) or '{{' in value or '}}' in value
            or re.search(r'https?://|www\.', value, re.I)):
        raise ValueError('report_text_invalid')
    return value.strip()


def normalize(args, request, consulted):
    """Citations are server receipts, never URLs/titles invented by the model."""
    title = text(args['title'], 160)
    language = args.get('language') or request.get('language', 'es')
    if language not in ('es', 'en') or language != request.get('language', language):
        raise ValueError('report_language_mismatch')
    is_deck = request['formats'] == ['pptx']
    pages = args.get('content')
    if not isinstance(pages, list) or not 1 <= len(pages) <= 20:
        raise ValueError('report_pages_required')
    if request.get('pages') and len(pages) != request['pages']:
        raise ValueError('report_page_count_mismatch')
    result, used = [], set()
    for page in pages:
        if not isinstance(page, dict) or set(page) != {'heading', 'blocks'}:
            raise ValueError('report_page_shape')
        heading = text(page['heading'], 110)
        blocks = page['blocks']
        if not isinstance(blocks, list) or not 1 <= len(blocks) <= 8:
            raise ValueError('report_blocks_required')
        clean = []
        for block in blocks:
            if not isinstance(block, dict) or set(block) - {'kind', 'text', 'sources', 'justification'} or not {'kind', 'text'} <= set(block):
                raise ValueError('report_block_shape')
            kind, value = block['kind'], text(block['text'], 2000)
            if kind not in ('fact', 'analysis', 'recommendation', 'estimate', 'example'):
                raise ValueError('report_claim_kind')
            sources = block.get('sources', [])
            if (not isinstance(sources, list) or len(sources) > 8 or
                    any(not isinstance(s, str) or s not in consulted for s in sources)):
                raise ValueError('report_source_not_consulted')
            if kind == 'fact' and not sources:
                raise ValueError('report_fact_source_required')
            if kind == 'example' and not request.get('allow_mock'):
                raise ValueError('report_fiction_not_requested')
            justification = text(block['justification'], 500) if block.get('justification') else ''
            if kind == 'estimate' and (not justification or not sources):
                raise ValueError('report_estimate_basis_required')
            # Numerical observations need a source. Analysis is not a loophole
            # for fabricated budgets, percentages, dates or survey results.
            if re.search(r'\d|[€$£%]', value) and kind not in ('fact', 'estimate', 'example'):
                raise ValueError('report_numeric_evidence_required')
            if kind == 'fact':
                evidence = ' '.join(consulted[s].get('excerpt', '') for s in sources)
                for number in re.findall(r'\d+(?:[.,]\d+)*', value):
                    if not re.search(r'(?<!\d)' + re.escape(number) + r'(?!\d)', evidence):
                        raise ValueError('report_number_not_in_source')
            used.update(sources)
            clean.append({'kind': kind, 'text': value, 'sources': sources, 'justification': justification})
        words = sum(len(b['text'].split()) for b in clean)
        if not (12 if is_deck else 160) <= words <= (105 if is_deck else 440):
            raise ValueError('report_page_content_capacity')
        result.append({'heading': heading, 'blocks': clean})
    if request.get('research_required') and not used:
        raise ValueError('report_required_sources_unavailable')
    if len(used) > (3 if is_deck else 8):
        raise ValueError('report_reference_capacity')
    refs = [{k: consulted[s][k] for k in ('id', 'title', 'publisher', 'url', 'retrieved_at', 'sha256', 'excerpt')}
            for s in sorted(used)]
    limitations = [] if refs else [
        'No se han consultado fuentes externas. Este documento contiene análisis y recomendaciones sin verificación factual externa.'
        if language == 'es' else 'No external sources were consulted. This document contains analysis and recommendations without external factual verification.']
    return {'title': title, 'language': language, 'pages': result, 'references': refs, 'limitations': limitations}


def source_content(artifact):
    from .jobs import _decode
    if artifact.source.kind != 'narrative':
        raise ValueError('report_source_required')
    return json.loads(_decode(artifact.source))


def sources_for_artifact(artifact):
    from .research import consulted_sources
    sources = consulted_sources(artifact.execution)
    if artifact.verification.get('narrative'):
        # Saved only after server normalization, encrypted and hash checked.
        # Carry the actual consulted excerpt across multiple immutable edits.
        sources.update({s['id']: s for s in source_content(artifact)['references'] if 'excerpt' in s})
    return sources


def local_fields(run, call, resolve):
    from .jobs import _base
    from .research import consulted_sources
    request = run.job['request']
    if call.name in ('build_document', 'build_presentation'):
        args = deepcopy(call.args)
        args['title'] = resolve(args['title'])
        for page in args.get('content', []):
            page['heading'] = resolve(page['heading'])
            for block in page['blocks']:
                block['text'] = resolve(block['text'])
                if block.get('justification'):
                    block['justification'] = resolve(block['justification'])
        sources = consulted_sources(run)
        if run.job.get('base'):
            base = _base(run.conversation, run.job['logical_id'], run.job['base'])
            sources = sources_for_artifact(base) | sources
        content = normalize(args, request, sources)
        format = 'pptx' if call.name == 'build_presentation' else call.args['format']
    else:
        base = _base(run.conversation, run.job['logical_id'], run.job['base'])
        content = source_content(base)
        format = 'pdf' if call.name == 'export_pdf' else base.format
        if call.name in ('edit_document', 'edit_presentation'):
            value = text(resolve(call.args['value']), 2000)
            field = call.args['field']
            if call.name == 'edit_presentation':
                slide = call.args['slide'] - 1
                if not 0 <= slide < len(content['pages']):
                    raise ValueError('report_slide_invalid')
                if field == 'title':
                    content['pages'][slide]['heading'] = text(value, 110)
                elif field == 'body':
                    if len(content['pages'][slide]['blocks']) != 1:
                        raise ValueError('report_use_block_field')
                    content['pages'][slide]['blocks'][0]['text'] = value
                else:
                    raise ValueError('report_field_invalid')
            elif field == 'title':
                content['title'] = text(value, 160)
            elif field == 'section':
                content['pages'][0]['blocks'][0]['text'] = value
            else:
                match = re.fullmatch(r'page:(\d+):block:(\d+)', field)
                if not match:
                    raise ValueError('report_field_invalid')
                try:
                    page, block = int(match[1]) - 1, int(match[2]) - 1
                    if min(page, block) < 0:
                        raise IndexError()
                    content['pages'][page]['blocks'][block]['text'] = value
                except IndexError:
                    raise ValueError('report_field_invalid') from None
            # Revalidate edits against the original consulted evidence. The new
            # run may add sources but cannot invent the prior source excerpts.
            sources = sources_for_artifact(base) | consulted_sources(run)
            args = {'title': content['title'], 'language': content['language'], 'content': content['pages']}
            content = normalize(args, {**request, 'allow_mock': base.execution.job.get('request', {}).get('allow_mock', False)}, sources)
    from front.policy import evaluate_policy
    decision = evaluate_policy(run.prompt.user, [('files', json.dumps(content, ensure_ascii=False))], destinations=[])
    if decision.status != 'allow':
        from .service import AgentDenied
        raise AgentDenied('report_content_policy_denied')
    return {'status': 'authorized', 'title': content['title'], 'format': format, 'narrative': content}


def verify_revision(before, after, extension, *, export=False):
    """Independent gate: revisions may replace text, never styles or geometry."""
    if export:
        if before != after:
            raise ValueError('report_export_source_changed')
        return
    def parts(data):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return {item.filename: archive.read(item) for item in archive.infolist()}
    original, revised = parts(before), parts(after)
    if original.keys() != revised.keys():
        raise ValueError('report_revision_parts_changed')
    tag = '{http://schemas.openxmlformats.org/' + (
        'wordprocessingml/2006/main}t' if extension == 'docx' else 'drawingml/2006/main}t')
    for name, raw in original.items():
        if revised[name] == raw:
            continue
        editable = name == 'word/document.xml' if extension == 'docx' else bool(re.fullmatch(r'ppt/slides/slide\d+\.xml', name))
        if not editable:
            raise ValueError('report_revision_nontext_changed')
        trees = [etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
                 for data in (raw, revised[name])]
        for tree in trees:
            for node in tree.iter(tag):
                node.text = ''
        if etree.tostring(trees[0], method='c14n') != etree.tostring(trees[1], method='c14n'):
            raise ValueError('report_revision_style_changed')


def verify_bundle(encoded, expected, *, format, request, base=None, export=False):
    if format not in ('docx', 'pdf', 'pptx') or len(encoded) > MAX_DOCUMENT_BYTES * 2:
        raise ValueError('report_bundle_size')
    data = base64.b64decode(encoded, validate=True)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if (len(items) > 43 or sum(i.file_size for i in items) > MAX_DOCUMENT_BYTES
                or len({i.filename for i in items}) != len(items) or any(i.flag_bits & 1 for i in items)):
            raise ValueError('report_bundle_size')
        parts = {i.filename: archive.read(i) for i in items}
    extension = 'pptx' if format == 'pptx' else 'docx'
    document = parts['final.' + extension]
    if base is not None:
        verify_revision(base, document, extension, export=export)
    if extension == 'docx':
        package(document)
        doc = Document(io.BytesIO(document))
        actual = [p.text for p in doc.paragraphs]
        if actual != paragraphs(expected) or doc.tables or len(doc.sections) != 1 or doc.inline_shapes:
            raise ValueError('report_docx_content_mismatch')
        if doc.core_properties.language != expected['language']:
            raise ValueError('report_docx_language')
        headings = [p for p in doc.paragraphs if p.style.name == 'Heading 1']
        if len(headings) != len(expected['pages']) or any(not p.paragraph_format.page_break_before for p in headings[1:]):
            raise ValueError('report_docx_page_breaks')
    else:
        _verify_pptx(document, expected)
    pdf = parts['final.pdf']
    count = _verify_pdf(pdf, expected, extension)
    required = {'final.' + extension, 'final.pdf', 'render.json'} | {f'page-{n}.png' for n in range(1, count + 1)}
    if set(parts) != required:
        raise ValueError('report_render_pages')
    manifest = json.loads(parts['render.json'])
    if (manifest.get(extension + '_sha256') != sha256(document) or manifest.get('pdf_sha256') != sha256(pdf)
            or manifest.get('pages') != [sha256(parts[f'page-{n}.png']) for n in range(1, count + 1)]):
        raise ValueError('report_render_digest')
    for number in range(1, count + 1):
        with Image.open(io.BytesIO(parts[f'page-{number}.png'])) as image:
            expected_size = (827, 1168, 1171) if extension == 'docx' else (1334, 749, 751)
            if (image.format != 'PNG' or abs(image.width - expected_size[0]) > 1
                    or not expected_size[1] <= image.height <= expected_size[2]):
                raise ValueError('report_render_dimensions')
            image.verify()
    raw = pdf if format == 'pdf' else document
    proof = {'sha256': sha256(raw), 'narrative': True, 'pages': count,
        extension + '_sha256': sha256(document), 'pdf_sha256': sha256(pdf), 'page_sha256': manifest['pages'],
        'automated_passed': True, 'visual_review': 'required', 'source_ids': [s['id'] for s in expected['references']],
        'checks': ['package_structure', 'content_exact', 'pdf_text', 'page_count', 'embedded_fonts', 'rendered_final_bytes']}
    from .request_contract import check_delivery
    check_delivery(request, format, proof)
    proof['request_checked'] = True
    return raw, proof, json.dumps({name: base64.b64encode(value).decode() for name, value in parts.items()})


def _verify_pptx(data, expected):
    from pptx import Presentation
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if (len(items) > 250 or sum(i.file_size for i in items) > MAX_DOCUMENT_BYTES
                or len({i.filename for i in items}) != len(items)):
            raise ValueError('report_pptx_size')
        for item in items:
            if any(key in item.filename.lower() for key in ('vbaproject', 'embeddings/', 'activex', 'media/')):
                raise ValueError('report_pptx_active_content')
            raw = archive.read(item)
            if item.filename.endswith(('.xml', '.rels')):
                if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:
                    raise ValueError('report_pptx_xml')
                tree = etree.fromstring(raw, etree.XMLParser(resolve_entities=False, no_network=True))
                if item.filename.endswith('.rels') and any(n.get('TargetMode', '').lower() == 'external' for n in tree):
                    raise ValueError('report_pptx_external')
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) != len(expected['pages']):
        raise ValueError('report_slide_count')
    for index, (slide, page) in enumerate(zip(presentation.slides, expected['pages'])):
        body = '\n\n'.join(block_text(b, expected['language']) for b in page['blocks'])
        if index == len(expected['pages']) - 1:
            if expected.get('limitations'):
                body += '\n\n' + ('Alcance y límites' if expected['language'] == 'es' else 'Scope and limitations') + '\n' + '\n'.join(expected['limitations'])
            if expected['references']:
                body += '\n\n' + ('Referencias consultadas' if expected['language'] == 'es' else 'Sources consulted') + '\n' + '\n'.join(reference_text(s) for s in expected['references'])
        texts = [s.text.replace('\v', '\n') for s in slide.shapes if s.has_text_frame]
        if texts != ['REGULAY', page['heading'], body.strip(), expected['title'], str(index + 1)]:
            raise ValueError('report_pptx_content')
        if any(s.left < 0 or s.top < 0 or s.left+s.width > presentation.slide_width or
               s.top+s.height > presentation.slide_height for s in slide.shapes):
            raise ValueError('report_pptx_bounds')


def _verify_pdf(data, expected, extension):
    reader = PdfReader(io.BytesIO(data), strict=True)
    if reader.is_encrypted or len(reader.pages) != len(expected['pages']):
        raise ValueError('report_pdf_page_count')
    root = reader.trailer['/Root']
    if root.get('/OpenAction') is not None:
        opening = root['/OpenAction']
        if not isinstance(opening, ArrayObject) or len(opening) != 5 or opening[1] != '/XYZ':
            raise ValueError('report_pdf_action')
    if any(k in root for k in ('/AA', '/AcroForm')) or root.get('/Names') and root['/Names']:
        raise ValueError('report_pdf_active_content')
    normalize_text = lambda value: re.sub(r'\s+', '', value).replace('\u00ad', '')
    all_text = []
    for number, (page, content) in enumerate(zip(reader.pages, expected['pages']), 1):
        if '/AA' in page:
            raise ValueError('report_pdf_annotations')
        # Reference URLs remain text, never executable actions.
        for annotation in page.get('/Annots', []):
            annotation = annotation.get_object()
            if annotation.get('/Subtype') != '/Link' or any(key in annotation for key in ('/A', '/AA', '/JS', '/JavaScript')):
                raise ValueError('report_pdf_annotations')
        value = normalize_text(page.extract_text() or '')
        required = ['REGULAY', content['heading']] + [block_text(b, expected['language']) for b in content['blocks']]
        if extension == 'docx':
            required.append(('Página' if expected['language'] == 'es' else 'Page') + f' {number} / {len(reader.pages)}')
            if number == 1:
                required.append(expected['title'])
        else:
            required.extend([expected['title'], str(number)])
        if number == len(reader.pages):
            if expected.get('limitations'):
                required.extend(['Alcance y límites' if expected['language'] == 'es' else 'Scope and limitations', *expected['limitations']])
            if expected['references']:
                required.extend(['Referencias consultadas' if expected['language'] == 'es' else 'Sources consulted',
                    *[reference_text(source) for source in expected['references']]])
        if any(normalize_text(item) not in value for item in required) or '\ufffd' in value:
            raise ValueError('report_pdf_missing_content')
        # Also reject additions or duplicates: containing all expected text is
        # insufficient to prove that a PDF agrees with its editable original.
        remainder = value
        for item in sorted((normalize_text(item) for item in required), key=len, reverse=True):
            remainder = remainder.replace(item, '', 1)
        if remainder:
            raise ValueError('report_pdf_unexpected_content')
        for font in page['/Resources'].get('/Font', {}).values():
            font = font.get_object()
            if '/DescendantFonts' in font:
                font = font['/DescendantFonts'][0].get_object()
            descriptor = font.get('/FontDescriptor')
            if not descriptor or not any(k in descriptor.get_object() for k in ('/FontFile', '/FontFile2', '/FontFile3')):
                raise ValueError('report_pdf_font_not_embedded')
        all_text.append(value)
    if any(normalize_text(value) not in ''.join(all_text) for value in paragraphs(expected)):
        raise ValueError('report_pdf_docx_incoherent')
    return len(reader.pages)
