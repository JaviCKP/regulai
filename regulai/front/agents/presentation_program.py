"""Bounded OfficeCLI procedure. No network, model or arbitrary CLI arguments."""
import argparse
from decimal import Decimal
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook

INK, TEAL, PAPER, GOLD = '193B45', '326D72', 'F5F7F5', 'E4C38E'
ENV = {**os.environ, 'OFFICECLI_NO_AUTO_RESIDENT': '1', 'DOTNET_CLI_TELEMETRY_OPTOUT': '1'}
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
C = '{http://schemas.openxmlformats.org/drawingml/2006/chart}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def cli(*args):
    result = subprocess.run(['officecli', *map(str, args)], env=ENV, capture_output=True, text=True, timeout=30)
    if result.returncode or 'UNSUPPORTED' in result.stdout + result.stderr:
        raise ValueError('presentation_cli_failed: ' + result.stdout + result.stderr)
    return result.stdout


def props(file, path, **values):
    args = ['set', file, path]
    for key, value in values.items():
        args += ['--prop', f'{key}={value}']
    cli(*args)


def add(file, slide, kind, **values):
    args = ['add', file, f'/slide[{slide}]' if slide else '/', '--type', kind]
    for key, value in values.items():
        args += ['--prop', f'{key}={value}']
    cli(*args)


def shape(file, slide, name, text, x, y, w, h, size=20, color=INK, font='Liberation Sans'):
    add(file, slide, 'shape', name=name, text=text.replace('\\', '\\\\'), x=f'{x}cm', y=f'{y}cm',
        width=f'{w}cm', height=f'{h}cm', size=size, color=color, font=font, fill='none', line='none')


def money(value):
    return format(Decimal(value) / 100, '.2f').replace('.', ',') + '\u00a0€'


def parts(path):
    with zipfile.ZipFile(path) as source:
        return {n: source.read(n) for n in source.namelist()}


def write_parts(path, data):
    temporary = path.with_suffix('.tmp')
    with zipfile.ZipFile(temporary, 'w', zipfile.ZIP_DEFLATED) as target:
        for name, content in data.items():
            target.writestr(name, content)
    temporary.replace(path)


def stable_parts(file, base, allowed):
    before, after = parts(base), parts(file)
    if not set(allowed) <= set(after):
        raise ValueError('presentation_edit_missing_slide')
    write_parts(file, {name: after[name] if name in allowed else value for name, value in before.items()})


def editable_chart(file, rows):
    """Bind OfficeCLI's native chart to a local XLSX, retaining its style/cache.

    Upstream only writes caches. This small OOXML completion enables Edit Data;
    it is not a second presentation engine. The independent verifier checks both.
    """
    package = parts(file)
    charts = [n for n in package if n.endswith('.xml') and '/charts/chart' in n]
    if len(charts) != 1:
        raise ValueError('presentation_chart_count')
    chart_name = charts[0]
    tree = etree.fromstring(package[chart_name])
    labels = tree.find('.//' + C + 'dLbls')
    for index, (_, amount) in enumerate(rows):
        label = etree.Element(C + 'dLbl')
        etree.SubElement(label, C + 'idx', val=str(index))
        rich = etree.SubElement(etree.SubElement(label, C + 'tx'), C + 'rich')
        etree.SubElement(rich, A + 'bodyPr')
        etree.SubElement(rich, A + 'lstStyle')
        paragraph = etree.SubElement(rich, A + 'p')
        run = etree.SubElement(paragraph, A + 'r')
        style = etree.SubElement(run, A + 'rPr', lang='es-ES', sz='1400')
        etree.SubElement(etree.SubElement(style, A + 'solidFill'), A + 'srgbClr', val=INK)
        etree.SubElement(style, A + 'latin', typeface='Liberation Sans')
        etree.SubElement(run, A + 't').text = money(amount)
        etree.SubElement(label, C + 'dLblPos', val='outEnd')
        etree.SubElement(label, C + 'showVal', val='0')
        labels.insert(index, label)
    lang = tree.find(C + 'lang')
    if lang is None:
        lang = etree.Element(C + 'lang')
        tree.insert(1 if tree.find(C + 'date1904') is not None else 0, lang)
    lang.set('val', 'es-ES')
    for node in tree.iter(C + 'f'):
        # The reviewed one-series chart uses this exact workbook shape.
        parent = node.getparent().getparent().tag
        node.text = "'Datos'!$B$1" if parent == C + 'tx' else (
            f"'Datos'!$A$2:$A${len(rows)+1}" if parent == C + 'cat' else f"'Datos'!$B$2:$B${len(rows)+1}")
    external = etree.SubElement(tree, C + 'externalData', {R + 'id': 'regulayData'})
    etree.SubElement(external, C + 'autoUpdate', {'val': '0'})
    package[chart_name] = etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True)
    parent = str(Path(chart_name).parent).replace('\\', '/')
    rels = etree.Element('{http://schemas.openxmlformats.org/package/2006/relationships}Relationships')
    etree.SubElement(rels, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
        Id='regulayData', Type=R[1:-1] + '/package', Target='data.xlsx')
    package[parent + '/_rels/' + Path(chart_name).name + '.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Datos'
    sheet.append(['Concepto', 'Importe (€)'])
    for item, amount in rows:
        sheet.append([item, Decimal(amount) / 100])
        sheet.cell(sheet.max_row, 1).data_type = 's'
        sheet.cell(sheet.max_row, 2).number_format = '[$-es-ES]0.00'
    stream = io.BytesIO()
    book.save(stream)
    package[parent + '/data.xlsx'] = stream.getvalue()
    types = etree.fromstring(package['[Content_Types].xml'])
    if not any(n.get('Extension') == 'xlsx' for n in types):
        etree.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Default',
            Extension='xlsx', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    package['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8')
    slide = etree.fromstring(package['ppt/slides/slide2.xml'])
    for row in slide.iter(A + 'tr'):
        cell = row.findall(A + 'tc')[1]
        for paragraph in cell.iter(A + 'p'):
            style = paragraph.find(A + 'pPr')
            if style is None:
                style = etree.Element(A + 'pPr')
                paragraph.insert(0, style)
            style.set('algn', 'r')
    package['ppt/slides/slide2.xml'] = etree.tostring(slide, xml_declaration=True, encoding='UTF-8', standalone=True)
    write_parts(file, package)


def build(file, rows, fields, template):
    corporate = fields['template'] == 'corporate'
    if corporate:
        shutil.copyfile(template, file)
        for number, value in ((1, fields['title']), (2, 'Detalle de importes'), (3, 'Distribución por concepto')):
            props(file, f'/slide[{number}]/shape[@name=Title]', text=value.replace('\\', '\\\\'))
    else:
        cli('create', file)
        props(file, '/', slideSize='widescreen', **{'theme.font.major.latin': 'Liberation Serif', 'theme.font.minor.latin': 'Liberation Sans'})
        for number, title in enumerate([fields['title'], 'Detalle de importes', 'Distribución por concepto'], 1):
            add(file, 0, 'slide', layout='blank', background=INK if number == 1 else PAPER)
            shape(file, number, 'Brand', 'REGULAY  /  INFORMES', 1.7, 1, 30, .9, 12, GOLD if number == 1 else TEAL)
            shape(file, number, 'Title', title, 1.7, 3.0, 19 if number == 1 else 30, 5 if number == 1 else 2.5,
                  42 if number == 1 else 36, 'FFFFFF' if number == 1 else INK, 'Liberation Serif')
            shape(file, number, 'Footer', f'{number:02d}  /  REGULAY · Informe económico', 1.7, 17, 30, .9, 11, 'B9CDCE' if number == 1 else '60797E')
            add(file, number, 'notes', text='Importes del archivo de origen; cálculo local en céntimos y representación en euros.')
    shape(file, 1, 'Body', 'Datos económicos de España\nDesglose y revisión de importes', 1.7, 10, 18, 4, 22, 'D8E5E5')
    shape(file, 1, 'Total', money(sum(v for _, v in rows)), 22, 6, 10, 3.7, 42, GOLD, 'Liberation Serif')
    shape(file, 1, 'TotalLabel', 'IMPORTE TOTAL', 22, 8.8, 10, 1.1, 14, 'D8E5E5')
    shape(file, 1, 'Count', f'{len(rows)} conceptos', 22, 11, 10, 1.3, 20, 'D8E5E5')
    add(file, 2, 'table', rows=len(rows)+2, cols=2, x='1.7cm', y='6cm', width='30cm', height='8cm', firstRow='true', headerFill=INK, colWidths='22cm,8cm', **{'border.all': 'none', 'border.horizontal': '0.5pt solid D7E2E1'})
    values = [['Concepto', 'Importe (€)']] + [[item, money(cents)] for item, cents in rows] + [['Total', money(sum(v for _, v in rows))]]
    for number, (item, amount) in enumerate(values, 1):
        props(file, f'/slide[2]/table[1]/tr[{number}]', c1=item.replace('\\', '\\\\'), c2=amount)
    props(file, '/slide[2]/table[1]', font='Liberation Sans', size=20, color=INK)
    props(file, '/slide[2]/table[1]/tr[1]', fill=INK, color='FFFFFF', bold='true')
    props(file, f'/slide[2]/table[1]/tr[{len(values)}]', fill='DDEBE9', bold='true')
    shape(file, 2, 'Body', 'Fuente: archivo aportado · Suma calculada en céntimos', 1.7, 15, 30, 1.5, 18, TEAL)
    add(file, 3, 'chart', chartType='bar', categories=','.join(item for item, _ in rows),
        x='1.7cm', y='6cm', width='18cm', height='10cm', legend='none',
        axisfont='12:193B45:Liberation Sans', valaxisvisible='false', gridlines='none', dataLabels='value',
        axismax=str(max(v for _, v in rows)/100*1.3 if max(v for _, v in rows) > 0 else 1),
        **{'series1.name': 'Importe (€)', 'series1.values': ','.join(str(Decimal(v)/100) for _, v in rows), 'series1.color': TEAL})
    shape(file, 3, 'Body', fields['section'], 21, 6, 11, 10, 18)
    cli('close', file)
    if corporate:
        # Only filled slides and newly-added chart relationships may differ.
        before, after = parts(template), parts(file)
        for name in before:
            if not name.startswith('ppt/slides/') and name != '[Content_Types].xml':
                after[name] = before[name]
        write_parts(file, after)
    editable_chart(file, rows)


def edit(file, source, fields):
    shutil.copyfile(source, file)
    number = fields['slide']
    props(file, f'/slide[{number}]/shape[@name={"Title" if fields["field"] == "title" else "Body"}]',
          text=fields['value'].replace('\\', '\\\\'))
    cli('close', file)
    stable_parts(file, source, (f'ppt/slides/slide{number}.xml',))


def render(output):
    file = output / 'final.pptx'
    cli('close', file)
    cli('validate', file)
    issues = cli('view', file, 'issues')
    if 'Found 0 issue(s)' not in issues:
        raise ValueError('presentation_layout_issues: ' + issues)
    before = hashlib.sha256(file.read_bytes()).hexdigest()
    subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(output/'lo'), '--headless',
        '--convert-to', 'pdf:impress_pdf_Export', '--outdir', str(output), str(file)], capture_output=True, check=True, timeout=90)
    subprocess.run(['pdftoppm', '-r', '100', '-png', str(output/'final.pdf'), str(output/'page')], capture_output=True, check=True, timeout=60)
    if hashlib.sha256(file.read_bytes()).hexdigest() != before:
        raise ValueError('presentation_changed_during_render')


def seal(output):
    from pptx import Presentation
    file = output/'final.pptx'
    if len(Presentation(file).slides) != 3:
        raise ValueError('presentation_slide_count')
    files = [file, output/'final.pdf', *(output/f'page-{n}.png' for n in range(1, 4))]
    sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    (output/'render.json').write_text(json.dumps({'pptx_sha256': sha(file), 'pdf_sha256': sha(files[1]), 'pages': [sha(p) for p in files[2:]]}))
    with zipfile.ZipFile(output/'candidate.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in [*files, output/'render.json']:
            archive.write(path, path.name)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('operation', choices=['build', 'edit', 'render', 'seal'])
    for key in ('output', 'source', 'fields', 'template'):
        parser.add_argument('--'+key, type=Path, required=key == 'output')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    if args.operation in ('render', 'seal'):
        return (render if args.operation == 'render' else seal)(args.output)
    fields = json.loads(args.fields.read_text())
    file = args.output/'final.pptx'
    if args.operation == 'edit':
        return edit(file, args.source, fields)
    import csv
    if args.source.suffix == '.csv':
        with args.source.open(encoding='utf-8-sig', newline='') as stream:
            rows = [(r['item'], int(r['amount_cents'])) for r in csv.DictReader(stream)]
    else:
        book = load_workbook(args.source, data_only=True)
        rows = [(r[0], int(r[1])) for r in list(book.active.values)[1:]]
        book.close()
    build(file, rows, fields, args.template)


if __name__ == '__main__':
    main()
