"""Reviewed, offline template procedure; runs only inside the document sandbox."""
import argparse
from copy import deepcopy
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import subprocess
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from openpyxl import load_workbook

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def repack(source, output, document_xml):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            target.writestr(item, document_xml if item.filename == 'word/document.xml' else original.read(item))


def build(template, rows, fields, output):
    doc = Document(template)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace('{{TITLE}}', fields['title']).replace('{{SUMMARY}}', fields['section'])
    table = doc.tables[0]
    seed = deepcopy(table.rows[1]._tr)
    table._tbl.remove(table.rows[1]._tr)
    amounts = rows + [('Total', sum(v for _, v in rows))]
    for item, cents in amounts:
        table._tbl.append(deepcopy(seed))
        row = table.rows[-1]
        for cell, value in zip(row.cells, (item, format(Decimal(cents) / 100, '.2f').replace('.', ',') + ' €')):
            cell.paragraphs[0].runs[0].text = value
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    temporary = io.BytesIO()
    doc.save(temporary)
    with zipfile.ZipFile(temporary) as generated:
        repack(template, output, generated.read('word/document.xml'))


def edit(base, field, value, output):
    with zipfile.ZipFile(base) as source:
        tree = etree.fromstring(source.read('word/document.xml'))
    paragraphs = list(tree.iter(W + 'p'))
    target = None
    for index, paragraph in enumerate(paragraphs):
        style = paragraph.find('./' + W + 'pPr/' + W + 'pStyle')
        if ((field == 'title' and style is not None and style.get(W + 'val') == 'Title') or
                (field == 'section' and index and ''.join(paragraphs[index - 1].itertext()) == 'Resumen ejecutivo')):
            target = paragraph
            break
    if target is None:
        raise ValueError('Editable field missing')
    texts = list(target.iter(W + 't'))
    if len(texts) != 1:
        raise ValueError('Unsupported template field structure')
    texts[0].text = value
    repack(base, output, etree.tostring(tree, xml_declaration=True, encoding='UTF-8', standalone=True))


def render(path, output):
    # Blocking subprocesses exit before the verifier reopens or hashes outputs.
    result = subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(output / 'lo'),
        '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(output), str(path)],
        capture_output=True, timeout=60)
    pdf = output / 'final.pdf'
    if result.returncode or not pdf.is_file() or not pdf.stat().st_size:
        raise ValueError('document_converter_failed')
    subprocess.run(['pdftoppm', '-r', '100', '-png', str(pdf), str(output / 'render')],
                   capture_output=True, check=True, timeout=60)
    pages = sorted(output.glob('render-*.png'))
    if not pages:
        raise ValueError('document_renderer_failed')
    for index, page in enumerate(pages, 1):
        page.rename(output / f'page-{index}.png')


def seal(output):
    # Separate process invocation, after conversion/rasterization has finished.
    docx, pdf = (output / 'final.docx').read_bytes(), (output / 'final.pdf').read_bytes()
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf), strict=True)
    Document(io.BytesIO(docx))
    pages = [output / f'page-{n}.png' for n in range(1, len(reader.pages) + 1)]
    sha = lambda data: hashlib.sha256(data).hexdigest()
    manifest = {'docx_sha256': sha(docx), 'pdf_sha256': sha(pdf),
                'pages': [sha(page.read_bytes()) for page in pages]}
    (output / 'render.json').write_text(json.dumps(manifest))
    with zipfile.ZipFile(output / 'candidate.zip', 'w', zipfile.ZIP_DEFLATED) as archive:
        for path in [output / 'final.docx', output / 'final.pdf', output / 'render.json', *pages]:
            archive.write(path, path.name)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('operation', choices=['build', 'edit', 'render', 'seal'])
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--source', type=Path)
    parser.add_argument('--template', type=Path)
    parser.add_argument('--fields', type=Path)
    args = parser.parse_args()
    output = args.output
    output.mkdir(parents=True, exist_ok=True)
    if args.operation == 'render':
        render(output / 'final.docx', output)
    elif args.operation == 'seal':
        seal(output)
    else:
        fields = json.loads(args.fields.read_text())
        if args.operation == 'edit':
            edit(args.source, fields['field'], fields['value'], output / 'final.docx')
        else:
            import csv
            if args.source.suffix == '.csv':
                with args.source.open(encoding='utf-8-sig', newline='') as stream:
                    rows = [(r['item'], int(r['amount_cents'])) for r in csv.DictReader(stream)]
            else:
                book = load_workbook(args.source, data_only=True)
                sheet = book.active
                if sheet.title == 'Resumen':
                    rows = [(r[0], int(r[1])) for r in list(sheet.values)[3:-1]]
                else:
                    rows = [(r[0], int(r[1])) for r in list(sheet.values)[1:]]
                book.close()
            build(args.template, rows, fields, output / 'final.docx')


if __name__ == '__main__':
    main()
