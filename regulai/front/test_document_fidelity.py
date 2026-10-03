"""Public extraction regressions; fixtures contain synthetic content only."""
import io
import json
import zipfile
import zlib
from datetime import datetime
from openpyxl import Workbook
from openpyxl.workbook.properties import CalcProperties
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from front.documents import extract_document
from front import test_documents
from front import test_pipeline
from front.connectors import ChatResponse
from front.test_chat_flow import ChatFixture
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject
from pypdf.generic import DecodedStreamObject


def inline_image_pdf(image=None, nested=False):
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(test_documents.DocumentTests().pdf(["Readable evidence"]))))
    page = writer.pages[0]
    stream = DecodedStreamObject()
    image = image if image is not None else b"BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\x00\x00 EI"
    stream.set_data(b"q 10 0 0 10 0 0 cm " + image + b" Q\n")
    if nested:
        stream[NameObject("/Subtype")] = NameObject("/Form")
        stream[NameObject("/Resources")] = DictionaryObject()
        page["/Resources"][NameObject("/XObject")] = DictionaryObject({
            NameObject("/InlineForm"): writer._add_object(stream)})
        body = DecodedStreamObject()
        body.set_data(page.get_contents().get_data() + b"\n/InlineForm Do\n")
    else:
        body = DecodedStreamObject()
        body.set_data(page.get_contents().get_data() + b"\n" + stream.get_data())
    page[NameObject("/Contents")] = writer._add_object(body)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def decimal_numbering_docx(restarts, paragraphs):
    word = WordDocument()
    definition = OxmlElement("w:abstractNum")
    definition.set(qn("w:abstractNumId"), "42")
    for index, restart in enumerate(restarts):
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), str(index))
        attributes = [("start", "1"), ("numFmt", "decimal"),
                      ("lvlText", ".".join(f"%{number}" for number in range(1, index + 2)) + ".")]
        if restart is not None:
            attributes.append(("lvlRestart", restart))
        for name, value in attributes:
            child = OxmlElement("w:" + name)
            child.set(qn("w:val"), value)
            level.append(child)
        definition.append(level)
    word.part.numbering_part.element.append(definition)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), "42")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), "42")
    number.append(reference)
    word.part.numbering_part.element.append(number)
    for index, text in paragraphs:
        properties = word.add_paragraph(text)._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_numId().val = 42
        properties.get_or_add_ilvl().val = index
    output = io.BytesIO()
    word.save(output)
    return output.getvalue()


class FidelityTests(SimpleTestCase):
    def extract(self, name, content):
        return extract_document(SimpleUploadedFile(name, content))

    def archive(self, files):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in files.items():
                archive.writestr(name, content)
        return output.getvalue()

    def test_unsupported_formats_use_the_unreadable_coverage_contract(self):
        document = self.extract("legacy.xls", b"synthetic legacy bytes")
        self.assertEqual(document.status, "unreadable")
        self.assertIn("unsupported_format", document.warnings)

    def test_text_encodings_formfeed_and_bm_are_not_binary(self):
        for content in ("BM invoice\fPrice: 17 €".encode("utf-8"),
                        "Price: 17 €".encode("utf-16"), "Price: 17 €".encode("cp1252")):
            with self.subTest(content=content):
                document = self.extract("invoice.txt", content)
                self.assertEqual(document.status, "complete")
                self.assertIn("Price: 17 €", document.text)
        assumed = self.extract("invoice.txt", "Price: 17 €".encode("cp1252"))
        self.assertEqual(assumed.metadata()["disclosures"], [
            {"code": "encoding_assumed", "count": 1, "locators": ["document"]}])
        self.assertEqual(self.extract("invalid.txt", b"x\x81").status, "unreadable")

    def test_csv_dialects_preserve_cells_in_compact_rows(self):
        for separator in (",", ";", "\t", "|"):
            content = f'name{separator}note\nAda{separator}"first\nsecond"\n'.encode()
            with self.subTest(separator=separator):
                document = self.extract("sheet.csv", content)
                self.assertEqual(document.status, "complete")
                self.assertEqual(len(document.fragments), 2)
                self.assertEqual(document.fragments[1]["locator"], "row:2/cells:A2-B2")
                self.assertEqual(document.fragments[1]["text"], "A2: Ada\tB2: first\nsecond")
        document = self.extract("sheet.csv", "nombre;precio\nCafé;17 €\n".encode("cp1252"))
        self.assertEqual(document.status, "complete")
        self.assertIn("encoding_assumed", document.warnings)

    def test_uneven_csv_rows_keep_the_declared_cell_positions(self):
        for separator in (",", ";", "\t", "|"):
            with self.subTest(separator=separator):
                content = f"name{separator}price{separator}note\nAda{separator}17\nBob{separator}18{separator}ok\n"
                document = self.extract("uneven.csv", content.encode())
                self.assertEqual(document.status, "complete")
                self.assertEqual(document.fragments[1]["text"], "A2: Ada\tB2: 17")
                self.assertEqual(document.fragments[2]["locator"], "row:3/cells:A3-C3")
                self.assertEqual(document.fragments[2]["text"], "A3: Bob\tB3: 18\tC3: ok")

    def test_csv_with_ambiguous_ragged_delimiters_declares_unknown_coverage(self):
        document = self.extract("ambiguous.csv", b"name;price;note,tag\nAda;17\nBob;18;ok,tag\n")
        self.assertEqual(document.status, "partial")
        self.assertEqual(document.fragments, [])
        self.assertEqual(document.disclosures, [
            {"code": "csv_dialect_unknown", "count": 1, "locators": ["document"]}])

    def test_csv_without_delimiter_in_sample_cannot_claim_later_cells_are_complete(self):
        document = self.extract("late-delimiter.csv", b"a" * 65536 + b";17\nb;18\n")
        self.assertEqual(document.status, "partial")
        self.assertIn("csv_dialect_unknown", document.warnings)

    def test_text_groups_lines_without_losing_evidence_or_locations(self):
        content = ("ordinary text\n" * 1499 + "late evidence").encode()
        document = self.extract("large.txt", content)
        self.assertEqual(document.status, "complete")
        self.assertEqual(document.units_read, 1500)
        self.assertLess(len(document.fragments), 50)
        self.assertEqual(document.fragments[0]["locator"], "lines:1-40")
        self.assertEqual(document.fragments[-1]["locator"], "lines:1481-1500")
        self.assertTrue(document.text.endswith("late evidence"))

    def word_bytes(self, word):
        output = io.BytesIO()
        word.save(output)
        return output.getvalue()

    def test_word_accepts_changes_and_link_text_without_field_instructions(self):
        word = WordDocument()
        paragraph = word.add_paragraph()
        deleted = OxmlElement("w:del")
        run = OxmlElement("w:r")
        text = OxmlElement("w:delText")
        text.text = "30"
        run.append(text)
        deleted.append(run)
        paragraph._p.append(deleted)
        inserted = OxmlElement("w:ins")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "17 days "
        run.append(text)
        inserted.append(run)
        paragraph._p.append(inserted)
        instruction = OxmlElement("w:instrText")
        instruction.text = "HYPERLINK mailto:private@example.invalid"
        paragraph._p.append(instruction)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), word.part.relate_to("mailto:private@example.invalid",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Contact"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document = self.extract("changes.docx", self.word_bytes(word))
        self.assertEqual(document.status, "complete")
        self.assertEqual(document.text, "17 days Contact")
        self.assertEqual(document.disclosures, [{"code": "tracked_changes_present", "count": 2,
            "locators": ["word/document.xml/paragraph:1"]}])

    def test_word_accepted_view_excludes_deleted_rows_and_declares_row_revisions(self):
        word = WordDocument()
        word.add_paragraph("Current body")
        table = word.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Deleted row: 30 days"
        table.cell(1, 0).text = "Accepted row: 17 days"
        for index, revision in ((0, "del"), (1, "ins")):
            marker = OxmlElement("w:" + revision)
            marker.set(qn("w:id"), str(index))
            table.rows[index]._tr.get_or_add_trPr().append(marker)
        document = self.extract("rows.docx", self.word_bytes(word))
        self.assertEqual(document.text, "Current body\nAccepted row: 17 days")
        self.assertEqual(document.status, "complete")
        self.assertNotIn("Deleted row", str(document.fragments))
        self.assertEqual(document.disclosures, [{"code": "tracked_changes_present", "count": 2,
            "locators": ["word/document.xml/row:1", "word/document.xml/row:2"]}])

    def test_word_discloses_header_images_attachments_equations_and_symbols(self):
        word = WordDocument()
        paragraph = word.add_paragraph("Readable body")
        symbol = OxmlElement("w:sym")
        symbol.set(qn("w:char"), "F0A7")
        paragraph._p.append(symbol)
        paragraph._p.append(OxmlElement("m:oMath"))
        bitmap = io.BytesIO()
        Image.new("RGB", (2, 2), "white").save(bitmap, "PNG")
        bitmap.seek(0)
        word.sections[0].header.paragraphs[0].add_run().add_picture(bitmap)
        with zipfile.ZipFile(io.BytesIO(self.word_bytes(word))) as source:
            files = {name: source.read(name) for name in source.namelist()}
        files["word/embeddings/object1.bin"] = b"synthetic attachment"
        document = self.extract("visual.docx", self.archive(files))
        self.assertEqual(document.status, "complete_text_only")
        disclosures = {item["code"]: item for item in document.disclosures}
        for code in ("images_unread", "attachments_unread", "equations_unread", "symbols_unread"):
            self.assertEqual(disclosures[code]["count"], 1)
            self.assertTrue(disclosures[code]["locators"])
        self.assertIn("Readable body", document.text)
        self.assertNotIn("synthetic attachment", document.text)

    def test_word_renders_decimal_multilevel_numbering_and_declares_other_formats(self):
        word = WordDocument()
        definition = OxmlElement("w:abstractNum")
        definition.set(qn("w:abstractNumId"), "42")
        for level in range(3):
            node = OxmlElement("w:lvl")
            node.set(qn("w:ilvl"), str(level))
            for name, value in (("start", "1"), ("numFmt", "decimal" if level < 2 else "lowerRoman"),
                                ("lvlText", "%1." if level == 0 else "%1.%2.")):
                child = OxmlElement("w:" + name)
                child.set(qn("w:val"), value)
                node.append(child)
            definition.append(node)
        word.part.numbering_part.element.append(definition)
        numbering = OxmlElement("w:num")
        numbering.set(qn("w:numId"), "42")
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), "42")
        numbering.append(abstract)
        word.part.numbering_part.element.append(numbering)
        for level, text in ((0, "First"), (1, "Child"), (1, "Next child"), (0, "Second"), (1, "Reset child"), (2, "Roman")):
            paragraph = word.add_paragraph(text)
            properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
            properties.get_or_add_numId().val = 42
            properties.get_or_add_ilvl().val = level
        document = self.extract("numbered.docx", self.word_bytes(word))
        self.assertEqual(document.status, "complete_text_only")
        self.assertEqual(document.text, "1. First\n1.1. Child\n1.2. Next child\n2. Second\n2.1. Reset child\nRoman")
        self.assertEqual(document.disclosures, [{"code": "numbering_unsupported", "count": 1,
            "locators": ["word/document.xml/paragraph:6"]}])

    def test_decimal_numbering_obeys_never_restart_and_explicit_restart_levels(self):
        for restart, after_parent, expected in (("0", "2.1.3.", "2.2.4. Last grandchild"),
                                               ("1", "2.1.1.", "2.2.2. Last grandchild"),
                                               ("2", "2.1.3.", "2.2.1. Last grandchild"),
                                               (None, "2.1.3.", "2.2.1. Last grandchild")):
            with self.subTest(restart=restart):
                raw = decimal_numbering_docx([None, "0", restart], [
                    (0, "First"), (1, "Child"), (2, "Grandchild"), (2, "Next grandchild"),
                    (0, "Second"), (2, "Grandchild after parent"), (1, "Continuing child"), (2, "Last grandchild")])
                document = self.extract("restart.docx", raw)
                self.assertEqual(document.status, "complete")
                self.assertIn("2.2. Continuing child", document.text)
                self.assertIn(after_parent + " Grandchild after parent", document.text)
                self.assertTrue(document.text.endswith(expected), document.text)

    def test_decimal_numbering_with_invalid_restart_is_an_explicit_omission(self):
        document = self.extract("invalid-restart.docx", decimal_numbering_docx(["-1"], [(0, "Readable text")]))
        self.assertEqual(document.status, "complete_text_only")
        self.assertEqual(document.text, "Readable text")
        self.assertIn("numbering_unsupported", document.warnings)

    def xlsx_bytes(self, workbook, cache=None):
        output = io.BytesIO()
        workbook.save(output)
        if cache is None:
            return output.getvalue()
        with zipfile.ZipFile(io.BytesIO(output.getvalue())) as source:
            files = {name: source.read(name) for name in source.namelist()}
        files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b"<v></v>", f"<v>{cache}</v>".encode())
        return self.archive(files)

    def test_xlsx_cached_formula_sends_value_but_keeps_expression_for_security(self):
        workbook = Workbook()
        workbook.calculation = CalcProperties(fullCalcOnLoad=False)
        sheet = workbook.active
        sheet["A1"], sheet["B1"], sheet["C1"] = 2, 3, "=A1+B1"
        sheet["A2"] = "Linked evidence"
        sheet["A2"].hyperlink = "https://example.invalid/never-called"
        document = self.extract("cached.xlsx", self.xlsx_bytes(workbook, cache=5))
        self.assertEqual(document.status, "complete")
        self.assertIn("=A1+B1", document.text)
        self.assertEqual(document.fragments[0]["locator"], "sheet:Sheet/row:1/cells:A1-C1")
        self.assertEqual(document.fragments[0]["text"], "A1: 2\tB1: 3\tC1: 5")
        self.assertNotIn("=A1+B1", str(document.fragments))
        self.assertIn("formula_expression_security_only", document.warnings)
        self.assertIn("Linked evidence", document.text)
        for full_calc, cache in ((False, None), (True, 5)):
            workbook.calculation.fullCalcOnLoad = full_calc
            partial = self.extract("uncertain.xlsx", self.xlsx_bytes(workbook, cache=cache))
            self.assertEqual(partial.status, "partial")
            self.assertIn("formula_coverage_unknown", partial.warnings)
            self.assertIn("=A1+B1", partial.text)

    def test_xlsx_discloses_visual_content_and_formats_dates_and_percentages(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = 0.17
        sheet["A1"].number_format = "0.0%"
        sheet["B1"] = datetime(2026, 10, 3)
        sheet["B1"].number_format = "yyyy-mm-dd"
        sheet["A1"].comment = Comment("Private comment", "Synthetic")
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=1))
        sheet.add_chart(chart, "D1")
        image = io.BytesIO()
        Image.new("RGB", (2, 2), "white").save(image, "PNG")
        image.seek(0)
        sheet.add_image(SheetImage(image), "D4")
        with zipfile.ZipFile(io.BytesIO(self.xlsx_bytes(workbook))) as source:
            files = {name: source.read(name) for name in source.namelist()}
        files["xl/embeddings/object1.bin"] = b"synthetic attachment"
        document = self.extract("visual.xlsx", self.archive(files))
        self.assertEqual(document.status, "complete_text_only")
        self.assertEqual(document.fragments[0]["text"], "A1: 17.0%\tB1: 2026-10-03")
        for code in ("images_unread", "charts_unread", "comments_unread", "attachments_unread"):
            notice = next(item for item in document.disclosures if item["code"] == code)
            self.assertEqual(notice["count"], 1)
            self.assertTrue(notice["locators"])
        self.assertNotIn("Private comment", document.text)

    def test_pdf_blank_pages_are_not_scans_and_link_annotations_are_text_complete(self):
        fixture = test_documents.DocumentTests()
        blank = self.extract("blank.pdf", fixture.pdf(["Readable", None]))
        self.assertEqual(blank.status, "complete")
        self.assertEqual(blank.units_read, 2)
        self.assertEqual(blank.disclosures, [{"code": "blank_page", "count": 1, "locators": ["page:2"]}])
        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(fixture.pdf(["Linked text"]))))
        writer.pages[0][NameObject("/Annots")] = ArrayObject([writer._add_object(DictionaryObject({
            NameObject("/Subtype"): NameObject("/Link"),
            NameObject("/A"): DictionaryObject({NameObject("/S"): NameObject("/URI"),
                NameObject("/URI"): TextStringObject("https://example.invalid")})}))])
        output = io.BytesIO()
        writer.write(output)
        linked = self.extract("linked.pdf", output.getvalue())
        self.assertEqual(linked.status, "complete")
        self.assertIn("Linked text", linked.text)

    def test_pdf_inline_images_are_declared_with_their_page(self):
        document = self.extract("inline.pdf", inline_image_pdf())
        self.assertEqual(document.status, "complete_text_only")
        self.assertIn("Readable evidence", document.text)
        self.assertEqual(document.disclosures, [
            {"code": "images_unread", "count": 1, "locators": ["page:1"]}])

    def test_pdf_inline_images_in_forms_are_declared_without_expanding_pixels(self):
        image = (b"BI /W 33554433 /H 1 /CS /G /BPC 8 /F /Fl ID " +
                 zlib.compress(b"x" * 33554433) + b" EI")
        for nested in (False, True):
            with self.subTest(nested=nested):
                document = self.extract("compressed-inline.pdf", inline_image_pdf(image=image, nested=nested))
                self.assertEqual(document.status, "complete_text_only")
                self.assertIn("Readable evidence", document.text)
                self.assertEqual(document.disclosures, [
                    {"code": "images_unread", "count": 1, "locators": ["page:1"]}])

    def test_pdf_declares_images_forms_attachments_and_detects_scan_pages(self):
        fixture = test_documents.DocumentTests()
        visible = self.extract("logo.pdf", fixture.pdf(["Readable"], image=True))
        self.assertEqual(visible.status, "complete_text_only")
        self.assertEqual(visible.disclosures, [{"code": "images_unread", "count": 1, "locators": ["page:1"]}])
        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(fixture.pdf(["Readable", None]))))
        writer.pages[1][NameObject("/Resources")] = writer.pages[0][NameObject("/Resources")]
        source = PdfReader(io.BytesIO(fixture.pdf(["Scanned"], image=True)))
        writer.pages[1][NameObject("/Resources")] = source.pages[0]["/Resources"].clone(writer)
        output = io.BytesIO()
        writer.write(output)
        scan = self.extract("scan.pdf", output.getvalue())
        self.assertEqual(scan.status, "partial")
        self.assertIn("page_without_text", scan.warnings)
        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(fixture.pdf(["Readable"]))))
        writer.add_attachment("synthetic.txt", b"private attachment")
        writer._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/Fields"): ArrayObject([
            writer._add_object(DictionaryObject({NameObject("/FT"): NameObject("/Tx"),
                NameObject("/T"): TextStringObject("synthetic"), NameObject("/V"): TextStringObject("private field")}))])})
        output = io.BytesIO()
        writer.write(output)
        attached = self.extract("attached.pdf", output.getvalue())
        self.assertEqual(attached.status, "complete_text_only")
        for code in ("attachments_unread", "form_fields_unread"):
            notice = next(item for item in attached.disclosures if item["code"] == code)
            self.assertEqual(notice["count"], 1)
        self.assertNotIn("private attachment", attached.text)
        self.assertNotIn("private field", attached.text)

    def test_pdf_with_only_blank_pages_has_no_readable_evidence(self):
        document = self.extract("empty.pdf", test_documents.DocumentTests().pdf([None, None]))
        self.assertEqual(document.status, "unreadable")
        self.assertIn("empty_document", document.warnings)
        self.assertEqual(document.disclosures[0]["count"], 2)

    def test_word_numbering_missing_format_is_an_explicit_omission(self):
        word = WordDocument()
        paragraph = word.add_paragraph("Text")
        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_numId().val = 1
        with zipfile.ZipFile(io.BytesIO(self.word_bytes(word))) as source:
            files = {name: source.read(name) for name in source.namelist()}
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        import xml.etree.ElementTree as ET
        root = ET.fromstring(files["word/numbering.xml"])
        for level in root.iter(namespace + "lvl"):
            fmt = level.find(namespace + "numFmt")
            if fmt is not None:
                level.remove(fmt)
        files["word/numbering.xml"] = ET.tostring(root)
        document = self.extract("missing-format.docx", self.archive(files))
        self.assertEqual(document.status, "complete_text_only")
        self.assertIn("numbering_unsupported", document.warnings)
        self.assertNotIn("corrupt_document", document.warnings)

    def test_document_character_and_row_limits_remain_declared_after_grouping(self):
        text = self.extract("blocks.txt", ("x" * 1000 + "\n").encode() * 201)
        self.assertEqual(text.status, "partial")
        self.assertIn("text_limit", text.warnings)
        self.assertLess(text.units_read, text.units_total)
        workbook = Workbook()
        workbook.calculation = CalcProperties(fullCalcOnLoad=False)
        workbook.active["A1"] = "=1+1"
        workbook.active["A10001"] = "late evidence"
        rows = self.extract("rows.xlsx", self.xlsx_bytes(workbook, cache=2))
        self.assertEqual(rows.status, "partial")
        self.assertIn("unit_limit", rows.warnings)
        self.assertEqual(rows.units_read, 10000)
        self.assertIn("=1+1", rows.text)
        self.assertNotIn("late evidence", rows.text)

    def test_word_declares_alternate_chunks_outside_paragraphs(self):
        word = WordDocument()
        word.add_paragraph("Readable text")
        chunk = OxmlElement("w:altChunk")
        chunk.set(qn("r:id"), "synthetic-chunk")
        word._element.body.insert(1, chunk)
        with zipfile.ZipFile(io.BytesIO(self.word_bytes(word))) as source:
            files = {name: source.read(name) for name in source.namelist()}
        files["word/afchunk.html"] = b"<p>private omitted content</p>"
        document = self.extract("chunk.docx", self.archive(files))
        self.assertEqual(document.status, "complete_text_only")
        self.assertEqual(document.disclosures, [{"code": "attachments_unread", "count": 1,
            "locators": ["word/document.xml/element:altChunk"]}])
        self.assertNotIn("private omitted content", document.text)

    def test_textual_bm_upload_passes_shared_file_validation_but_bitmaps_do_not_extract(self):
        from front.file_security import validate_document_upload
        from django.core.exceptions import ValidationError
        try:
            content, metadata = validate_document_upload(SimpleUploadedFile("notes.txt", b"BM customer notes"))
        except ValidationError:
            self.fail("Text beginning with BM must pass document upload validation")
        self.assertEqual(content, b"BM customer notes")
        self.assertFalse(metadata["is_image"])
        bitmap = io.BytesIO()
        Image.new("RGB", (2, 2), "white").save(bitmap, "BMP")
        document = self.extract("renamed.txt", bitmap.getvalue())
        self.assertEqual(document.status, "unreadable")
        self.assertIn("signature_mismatch", document.warnings)


class DocumentHTTPTests(ChatFixture):
    configure = test_pipeline.PipelineTests.configure

    def test_decimal_numbering_never_restart_reaches_the_generator_and_cited_evidence(self):
        self.configure()
        raw = decimal_numbering_docx([None, "0"], [
            (0, "First"), (1, "Child"), (0, "Second"), (1, "Continuing child")])
        self.provider.return_value = ChatResponse(json.dumps({"claims": [{"text": "2.2. Continuing child",
            "citations": [{"id": "S4", "quote": "2.2. Continuing child"}]}]}), 30, "synthetic-v1", "stop", 20, 10)
        result = self.send("/api/chat/send/", "What is the final item?",
            file=SimpleUploadedFile("restart.docx", raw)).json()
        self.assertIn("2.2. Continuing child", str(self.provider.call_args))
        self.assertNotIn("2.1. Continuing child", str(self.provider.call_args))
        self.assertEqual(result["status"], "completed", result)
        self.assertEqual(result["evidence"]["items"][0]["excerpt"], "2.2. Continuing child")

    def test_uneven_csv_sends_correct_cells_and_ambiguous_csv_blocks_without_network(self):
        self.configure()
        self.provider.return_value = ChatResponse(json.dumps({"claims": [{"text": "Ada pays 17",
            "citations": [{"id": "S2", "quote": "A2: Ada\tB2: 17"}]}]}), 30, "synthetic-v1", "stop", 20, 10)
        result = self.send("/api/chat/send/", "What is Ada's price?",
            file=SimpleUploadedFile("uneven.csv", b"name;price;note\nAda;17\nBob;18;ok\n")).json()
        self.assertEqual(result["status"], "completed", result)
        self.assertEqual(result["evidence"]["items"][0]["locator"], "row:2/cells:A2-B2")
        self.assertIn("B2: 17", str(self.provider.call_args))
        self.http.reset_mock()
        self.provider.reset_mock()
        for path in ("/chat/", "/chat/send/", "/api/chat/send/"):
            with self.subTest(path=path):
                result = self.send(path, "Summarize", file=SimpleUploadedFile("ambiguous.csv",
                    b"name;price;note,tag\nAda;17\nBob;18;ok,tag\n")).json()
                self.assertEqual(result["error_type"], "input_unreadable")
                self.assertIn("csv_dialect_unknown", result["receipt"]["coverage"]["warnings"])
                self.http.assert_not_called()
                self.provider.assert_not_called()

    def test_deleted_word_row_is_absent_from_every_external_document_payload(self):
        self.configure()
        word = WordDocument()
        table = word.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "Deleted row: 30 days"
        table.cell(1, 0).text = "Accepted row: 17 days"
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:del"))
        output = io.BytesIO()
        word.save(output)
        self.provider.return_value = ChatResponse(json.dumps({"claims": [{"text": "Accepted row: 17 days",
            "citations": [{"id": "S1", "quote": "Accepted row: 17 days"}]}]}), 30, "synthetic-v1", "stop", 20, 10)
        for path in ("/chat/", "/chat/send/", "/api/chat/send/"):
            with self.subTest(path=path):
                self.http.reset_mock()
                self.provider.reset_mock()
                response = self.send(path, "What is the period?", file=SimpleUploadedFile("rows.docx", output.getvalue()))
                self.assertIn("Accepted row: 17 days", str(self.provider.call_args))
                self.assertNotIn("Deleted row: 30 days", str(self.provider.call_args))
                self.assertNotIn("Deleted row: 30 days", str(self.http.call_args_list))
                self.assertEqual(response.status_code, 302 if path == "/chat/" else 200)
                if path != "/chat/":
                    self.assertEqual(response.json()["status"], "completed", response.json())
                    self.assertIn("tracked_changes_present", response.json()["receipt"]["coverage"]["warnings"])

    def test_inline_pdf_requires_coverage_consent_before_all_three_send_entries(self):
        self.configure()
        self.provider.return_value = ChatResponse(json.dumps({"claims": [{"text": "Readable evidence",
            "citations": [{"id": "S1", "quote": "Readable evidence"}]}]}), 30, "synthetic-v1", "stop", 20, 10)
        for path in ("/chat/", "/chat/send/", "/api/chat/send/"):
            with self.subTest(path=path):
                response = self.send(path, "Summarize", file=SimpleUploadedFile("inline.pdf", inline_image_pdf()))
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["status"], "pending_consent")
                self.assertEqual(response.json()["consent_reason"], "coverage_incomplete")
                self.assertEqual(response.json()["receipt"]["coverage"]["disclosures"], [
                    {"code": "images_unread", "count": 1, "locators": ["page:1"]}])
                self.provider.assert_not_called()
                self.http.assert_not_called()

    def test_partial_document_blocks_all_three_send_entries_without_network(self):
        for path in ("/chat/", "/chat/send/", "/api/chat/send/"):
            with self.subTest(path=path):
                response = self.send(path, "Summarize", file=SimpleUploadedFile("large.txt", b"a" * 200001))
                self.assertEqual(response.json()["error_type"], "input_unreadable")
                self.assertEqual(response.json()["receipt"]["coverage"]["status"], "partial")
                self.provider.assert_not_called()
                self.http.assert_not_called()
