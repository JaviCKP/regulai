"""Bounded textual extraction. Coverage describes admitted units, never visual fidelity."""
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from itertools import zip_longest
import hashlib
import csv
import io
import json
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
import zlib


EXTRACTOR_VERSION = "2"
MAX_TEXT_CHARS = 200_000
MAX_UNITS = 10_000
MAX_PAGES = 100
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000
MAX_COMPRESSION_RATIO = 100
MAX_COLUMNS = 256
MAX_CELLS = 100_000
MAX_SHEETS = 100
MAX_XML_ELEMENTS = 100_000
MAX_XML_DEPTH = 128
INFORMATIONAL = {"encoding_assumed", "blank_page", "tracked_changes_present", "formula_expression_security_only"}
TEXT_ONLY = {"images_unread", "attachments_unread", "charts_unread", "comments_unread", "form_fields_unread",
             "annotations_unread", "equations_unread", "symbols_unread", "numbering_unsupported", "external_content_unread"}
TEXT_FORMATS = {"txt", "md", "json", "py", "js", "csv"}
MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass
class Document:
    text: str = ""
    fragments: list = field(default_factory=list)
    digest: str = ""
    format: str = ""
    status: str = "unreadable"
    units_read: int = 0
    units_total: int | None = None
    warnings: list = field(default_factory=list)
    disclosures: list = field(default_factory=list)
    file_name: str = ""
    file_size: int = 0
    _characters: int = field(default=0, init=False, repr=False)

    def metadata(self):
        return {
            "extractor_version": EXTRACTOR_VERSION,
            "digest": self.digest,
            "format": self.format,
            "status": self.status,
            "units_read": self.units_read,
            "units_total": self.units_total,
            "warnings": list(self.warnings),
            "disclosures": self.disclosures,
            "file_name": self.file_name,
            "file_size": self.file_size,
            "fragment_hashes": [fragment["id"] for fragment in self.fragments],
        }


def _append(document, locator, text):
    if len(document.fragments) >= MAX_UNITS:
        document.warnings.append("fragment_limit")
        return False
    characters = len(text) + bool(document.fragments)
    if document._characters + characters > MAX_TEXT_CHARS:
        document.warnings.append("text_limit")
        return False
    fragment_id = hashlib.sha256(
        f"{EXTRACTOR_VERSION}:{document.digest}:{locator}".encode()
    ).hexdigest()
    document.fragments.append({"id": fragment_id, "locator": locator, "text": text})
    document._characters += characters
    return True


def _disclose(document, code, locator, count=1):
    document.warnings.append(code)
    notice = next((item for item in document.disclosures if item["code"] == code), None)
    if notice is None:
        notice = {"code": code, "count": 0, "locators": []}
        document.disclosures.append(notice)
    notice["count"] += count
    if locator not in notice["locators"]:
        notice["locators"].append(locator)


def _decode_text(document, content):
    try:
        if content.startswith((b"\xff\xfe", b"\xfe\xff")):
            text = content.decode("utf-16")
        else:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")  # strict: undefined bytes remain invalid
                _disclose(document, "encoding_assumed", "document")
        if any(ord(char) < 32 and char not in "\t\n\r\f" for char in text):
            raise ValueError("binary text")
        if content and not text:
            raise ValueError("empty encoded text")
        return text
    except (UnicodeDecodeError, ValueError) as error:
        raise _Rejected("invalid_text") from error


def _parse_text(document, content):
    try:
        text = _decode_text(document, content)
    except _Rejected as error:
        document.warnings.append(str(error))
        return
    lines = text.splitlines()
    document.units_total = len(lines)
    for start in range(0, min(len(lines), MAX_UNITS), 40):
        block = lines[start:min(start + 40, MAX_UNITS)]
        end = start + len(block)
        locator = f"line:{end}" if len(block) == 1 else f"lines:{start + 1}-{end}"
        if not _append(document, locator, "\n".join(block)):
            break
        document.units_read += len(block)
    if len(lines) > MAX_UNITS:
        document.warnings.append("unit_limit")
    document.text = "\n".join(fragment["text"] for fragment in document.fragments)
    document.status = "partial" if document.units_read != document.units_total else "complete"
    if not document.text.strip() and document.status == "complete":
        document.warnings.append("empty_document")
        document.status = "unreadable"


def _parse_csv(document, content):
    from openpyxl.utils import get_column_letter

    text = _decode_text(document, content)
    csv.field_size_limit(1024 * 1024)
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    for row_number, row in enumerate(csv.reader(io.StringIO(text), dialect=dialect, strict=True), 1):
        if row_number > MAX_UNITS:
            document.warnings.append("unit_limit")
            return
        if len(row) > MAX_COLUMNS:
            document.warnings.append("column_limit")
            return
        values = [f"{get_column_letter(column)}{row_number}: {value}" for column, value in enumerate(row, 1)]
        if not _append(document, f"row:{row_number}/cells:A{row_number}-{get_column_letter(max(len(row), 1))}{row_number}", "\t".join(values)):
            return
        document.units_read += 1
    document.units_total = document.units_read


class _Rejected(Exception):
    pass


def _office_archive(document, content):
    archive = zipfile.ZipFile(io.BytesIO(content))
    entries = archive.infolist()
    if len(entries) > MAX_ARCHIVE_ENTRIES or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
        archive.close()
        raise _Rejected("expansion_limit")
    names = set()
    for entry in entries:
        if len(entry.filename) > 256 or entry.filename in names or ".." in entry.filename.replace("\\", "/").split("/"):
            archive.close()
            raise _Rejected("invalid_archive")
        names.add(entry.filename)
        if entry.filename.lower().endswith("vbaproject.bin"):
            archive.close()
            raise _Rejected("unsupported_active_content")
        if entry.flag_bits & 1:
            archive.close()
            raise _Rejected("encrypted_document")
        if entry.file_size > max(entry.compress_size, 1) * MAX_COMPRESSION_RATIO:
            archive.close()
            raise _Rejected("expansion_limit")
        if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
            archive.close()
            raise _Rejected("unsupported_compression")
    required = "word/document.xml" if document.format == "docx" else "xl/workbook.xml"
    if required not in names:
        archive.close()
        raise _Rejected("signature_mismatch")
    for entry in entries:
        if entry.filename.endswith((".xml", ".rels")):
            xml = archive.read(entry)
            markup = xml.replace(b"\x00", b"").upper()
            if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
                archive.close()
                raise _Rejected("unsafe_xml")
            elements = depth = 0
            try:
                for event, element in ET.iterparse(io.BytesIO(xml), events=("start", "end")):
                    if event == "start":
                        elements += 1
                        depth += 1
                        if elements > MAX_XML_ELEMENTS or depth > MAX_XML_DEPTH:
                            raise _Rejected("xml_structure_limit")
                        if (element.tag.endswith("}Relationship") and element.attrib.get("TargetMode") == "External"
                                and not element.attrib.get("Type", "").endswith("/hyperlink")):
                            _disclose(document, "external_content_unread", entry.filename)
                    else:
                        depth -= 1
                        element.clear()
            except Exception:
                archive.close()
                raise
    return archive


def _parse_xlsx(document, content):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

    with _office_archive(document, content) as archive:
        names = archive.namelist()
        for name in names:
            for prefix, code in (("xl/media/", "images_unread"), ("xl/embeddings/", "attachments_unread"),
                                 ("xl/charts/chart", "charts_unread")):
                if name.startswith(prefix) and not name.endswith(".rels"):
                    _disclose(document, code, name)
            if name.startswith("xl/comments") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                for comment in root.iter():
                    if comment.tag.endswith("}comment"):
                        _disclose(document, "comments_unread", f"{name}/cell:{comment.get('ref', '?')}")
        if any(name.startswith("xl/externalLinks/") for name in names):
            document.warnings.append("external_content_unread")
        cells = strings = sheets = 0
        for name in names:
            if not (name.startswith("xl/worksheets/") and name.endswith(".xml")) and name not in {"xl/sharedStrings.xml", "xl/workbook.xml"}:
                continue
            with archive.open(name) as source:
                for event, element in ET.iterparse(source):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "c":
                        cells += 1
                        column, row = coordinate_from_string(element.attrib["r"])
                        if column_index_from_string(column) > MAX_COLUMNS:
                            raise _Rejected("column_limit")
                    elif tag == "si":
                        strings += 1
                    elif tag == "sheet":
                        sheets += 1
                        if len(element.attrib.get("name", "")) > 31:
                            raise _Rejected("sheet_name_limit")
                        if sheets > MAX_SHEETS:
                            raise _Rejected("sheet_limit")
                    if max(cells, strings) > MAX_CELLS:
                        raise _Rejected("cell_limit")
                    element.clear()
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    cached_workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        for sheet, cached_sheet in zip_longest(workbook.worksheets, cached_workbook.worksheets):
            if sheet is None or cached_sheet is None or sheet.title != cached_sheet.title:
                raise _Rejected("worksheet_cache_mismatch")
            # Ignore producer-supplied dimensions: they can hide later cells.
            sheet.reset_dimensions()
            cached_sheet.reset_dimensions()
            for row_number, (row, cached_row) in enumerate(zip_longest(sheet.iter_rows(), cached_sheet.iter_rows()), 1):
                if document.units_read >= MAX_UNITS:
                    document.warnings.append("unit_limit")
                    return
                values, coordinates = [], []
                if row is None or cached_row is None:
                    raise _Rejected("worksheet_cache_mismatch")
                for cell, cached_cell in zip_longest(row, cached_row):
                    if cell is None or cached_cell is None or getattr(cell, "coordinate", None) != getattr(cached_cell, "coordinate", None):
                        raise _Rejected("worksheet_cache_mismatch")
                    if cell.value is None:
                        continue
                    locator = f"sheet:{sheet.title}/cell:{cell.coordinate}"
                    value = cell.value
                    if cell.data_type == "f":
                        formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", str(cell.value))
                        if document._characters + len(formula) + 1 > MAX_TEXT_CHARS:
                            document.warnings.append("text_limit")
                            return
                        document.text += "\n" + formula
                        document._characters += len(formula) + 1
                        _disclose(document, "formula_expression_security_only", locator)
                        value = cached_cell.value
                        if value is None or (workbook.calculation and workbook.calculation.fullCalcOnLoad):
                            _disclose(document, "formula_coverage_unknown", locator)
                        if value is None:
                            continue
                    coordinates.append(cell.coordinate)
                    values.append(f"{cell.coordinate}: {_cell_text(value, cell.number_format)}")
                if values and not _append(document, f"sheet:{sheet.title}/row:{row_number}/cells:{coordinates[0]}-{coordinates[-1]}", "\t".join(values)):
                    return
                document.units_read += 1
        document.units_total = document.units_read
    finally:
        workbook.close()
        cached_workbook.close()


def _cell_text(value, number_format):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    pattern = re.sub(r'"[^\"]*"|\\.', "", number_format.split(";")[0])
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in pattern:
        decimals = re.search(r"\.([0#?]+)", pattern)
        precision = min(len(decimals[1]), 15) if decimals else 0
        return f"{value * 100:.{precision}f}%"
    return str(value)


def _parse_docx(document, content):
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with _office_archive(document, content) as archive:
        names = archive.namelist()
        for name in names:
            if name.startswith("word/embeddings/"):
                _disclose(document, "attachments_unread", name)
        parts = ["word/document.xml"] + sorted(
            name for name in names if name.startswith((
                "word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments",
            )) and name.endswith(".xml")
        )
        definitions, numbers, styles, counters = {}, {}, {}, {}
        if "word/numbering.xml" in names:
            numbering = ET.fromstring(archive.read("word/numbering.xml"))
            for node in numbering.findall(namespace + "abstractNum"):
                definitions[node.get(namespace + "abstractNumId")] = {
                    int(level.get(namespace + "ilvl")): level for level in node.findall(namespace + "lvl")}
            for node in numbering.findall(namespace + "num"):
                reference = node.find(namespace + "abstractNumId")
                levels = dict(definitions.get(reference.get(namespace + "val"), {})) if reference is not None else {}
                starts = {}
                for override in node.findall(namespace + "lvlOverride"):
                    index = int(override.get(namespace + "ilvl"))
                    level, start = override.find(namespace + "lvl"), override.find(namespace + "startOverride")
                    if level is not None:
                        levels[index] = level
                    if start is not None:
                        starts[index] = int(start.get(namespace + "val"))
                numbers[node.get(namespace + "numId")] = (levels, starts)
        if "word/styles.xml" in names:
            styles = {node.get(namespace + "styleId"): node
                      for node in ET.fromstring(archive.read("word/styles.xml")).findall(namespace + "style")}
        count = 0
        for part in parts:
            paragraph = 0
            root_seen = False
            with archive.open(part) as source:
                for event, element in ET.iterparse(source, events=("start", "end")):
                    if not root_seen:
                        root_seen = True
                        if part == "word/document.xml" and element.tag != namespace + "document":
                            raise _Rejected("signature_mismatch")
                    if event == "start" and element.tag == namespace + "altChunk":
                        _disclose(document, "attachments_unread", f"{part}/element:altChunk")
                    if event != "end" or element.tag != namespace + "p":
                        continue
                    paragraph += 1
                    count += 1
                    if count > MAX_UNITS:
                        document.warnings.append("unit_limit")
                        return
                    chunks = []
                    locator = f"{part}/paragraph:{paragraph}"
                    for node in element.iter():
                        omission = {namespace + "drawing": "images_unread", namespace + "pict": "images_unread",
                                    namespace + "sym": "symbols_unread",
                                    "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath": "equations_unread"}.get(node.tag)
                        if omission:
                            _disclose(document, omission, locator)
                        if node.tag in {namespace + tag for tag in ("del", "ins", "moveFrom", "moveTo")}:
                            _disclose(document, "tracked_changes_present", locator)

                    def accepted_text(node):
                        if node.tag in {namespace + "del", namespace + "moveFrom", namespace + "instrText", namespace + "delText"}:
                            return
                        if node.tag == namespace + "t":
                            chunks.append(node.text or "")
                        elif node.tag == namespace + "tab":
                            chunks.append("\t")
                        elif node.tag in {namespace + "br", namespace + "cr"}:
                            chunks.append("\n")
                        for child in node:
                            accepted_text(child)

                    accepted_text(element)
                    text = "".join(chunks)
                    properties = element.find(namespace + "pPr")
                    num_properties = properties.find(namespace + "numPr") if properties is not None else None
                    style = properties.find(namespace + "pStyle") if properties is not None else None
                    seen = set()
                    while num_properties is None and style is not None and style.get(namespace + "val") not in seen:
                        style_id = style.get(namespace + "val")
                        seen.add(style_id)
                        definition = styles.get(style_id)
                        if definition is None:
                            break
                        num_properties = definition.find(f"{namespace}pPr/{namespace}numPr")
                        style = definition.find(namespace + "basedOn")
                    if num_properties is not None:
                        number_id = num_properties.find(namespace + "numId")
                        level_id = num_properties.find(namespace + "ilvl")
                        number_id = number_id.get(namespace + "val") if number_id is not None else None
                        level_id = int(level_id.get(namespace + "val")) if level_id is not None else 0
                        if number_id and number_id != "0":
                            levels, starts = numbers.get(number_id, ({}, {}))
                            level = levels.get(level_id)
                            fmt = level.find(namespace + "numFmt") if level is not None else None
                            if fmt is None or fmt.get(namespace + "val") != "decimal" or not 0 <= level_id <= 8:
                                _disclose(document, "numbering_unsupported", locator)
                            else:
                                values = counters.setdefault((part, number_id), {})
                                start = level.find(namespace + "start")
                                initial = starts.get(level_id, int(start.get(namespace + "val")) if start is not None else 1)
                                values[level_id] = values.get(level_id, initial - 1) + 1
                                for deeper in list(values):
                                    if deeper > level_id:
                                        del values[deeper]
                                pattern = level.find(namespace + "lvlText")
                                pattern = pattern.get(namespace + "val") if pattern is not None else f"%{level_id + 1}."
                                indices = [int(index) - 1 for index in re.findall(r"%([1-9])", pattern)]
                                if any(index not in levels or levels[index].find(namespace + "numFmt") is None
                                       or levels[index].find(namespace + "numFmt").get(namespace + "val") != "decimal" for index in indices):
                                    _disclose(document, "numbering_unsupported", locator)
                                else:
                                    prefix = re.sub(r"%([1-9])", lambda match: str(values.get(int(match[1]) - 1,
                                        starts.get(int(match[1]) - 1, 1))), pattern)
                                    text = prefix + " " + text
                    if not _append(document, f"{part}/paragraph:{paragraph}", text):
                        return
                    document.units_read += 1
                    element.clear()
        document.units_total = count


def _parse_pdf(document, content):
    from pypdf import PdfReader, filters

    remaining = MAX_EXPANDED_BYTES

    def bounded_decompress(data):
        nonlocal remaining
        decoder = zlib.decompressobj()
        decoded = decoder.decompress(data, remaining + 1)
        if len(decoded) > remaining or decoder.unconsumed_tail:
            raise _Rejected("expansion_limit")
        if not decoder.eof:
            raise _Rejected("corrupt_document")
        remaining -= len(decoded)
        return decoded

    def unsupported_filter(*args, **kwargs):
        raise _Rejected("unsupported_pdf_filter")

    # This process owns the parser: bounded Flate streams, no unbounded LZW decoder.
    filters.decompress = bounded_decompress
    filters.LZWDecode.decode = unsupported_filter
    reader = PdfReader(io.BytesIO(content), strict=True)
    if reader.is_encrypted:
        raise _Rejected("encrypted_document")
    document.units_total = len(reader.pages)
    for number, attachment in enumerate(reader.attachment_list, 1):
        if number > MAX_ARCHIVE_ENTRIES:
            raise _Rejected("attachment_limit")
        _disclose(document, "attachments_unread", f"document/attachment:{number}")
    form = reader.trailer["/Root"].get("/AcroForm")
    if form and (form.get_object().get("/Fields") or form.get_object().get("/XFA")):
        fields = form.get_object().get("/Fields", [])
        _disclose(document, "form_fields_unread", "document/form", max(len(fields), 1))
    for number in range(min(document.units_total, MAX_PAGES)):
        page = reader.pages[number]
        resources = page.get("/Resources", {})
        if hasattr(resources, "get_object"):
            resources = resources.get_object()
        locator = f"page:{number + 1}"
        image_count = 0
        pending = [(resources, 0)]
        seen = set()
        while pending:
            current, depth = pending.pop()
            if depth > MAX_XML_DEPTH or len(seen) > MAX_XML_ELEMENTS:
                raise _Rejected("pdf_structure_limit")
            xobjects = current.get("/XObject", {})
            xobjects = xobjects.get_object() if hasattr(xobjects, "get_object") else xobjects
            for reference in xobjects.values():
                item = reference.get_object()
                if id(item) in seen:
                    continue
                seen.add(id(item))
                if item.get("/Subtype") == "/Image":
                    image_count += 1
                elif item.get("/Subtype") == "/Form":
                    nested = item.get("/Resources", {})
                    pending.append((nested.get_object() if hasattr(nested, "get_object") else nested, depth + 1))
        if image_count:
            _disclose(document, "images_unread", locator, image_count)
        for reference in page.get("/Annots", []):
            annotation = reference.get_object()
            subtype = annotation.get("/Subtype")
            if subtype == "/FileAttachment":
                _disclose(document, "attachments_unread", locator)
            elif subtype not in ("/Link", "/Widget"):
                _disclose(document, "annotations_unread", locator)
        text = page.extract_text() if page.get("/Contents") else ""
        if not text or not text.strip():
            # Resources containing images are conservatively a scan, even if a
            # producer omits the drawing operator. Truly empty pages are readable.
            if image_count or (page.get_contents() and page.get_contents().get_data().strip()):
                _disclose(document, "page_without_text", locator)
            else:
                _disclose(document, "blank_page", locator)
                document.units_read += 1
            continue
        if not _append(document, locator, text):
            break
        document.units_read += 1
    if document.units_total > MAX_PAGES:
        document.warnings.append("page_limit")
    if not document.units_total:
        document.warnings.append("empty_document")


def _worker(format, content):
    document = Document(format=format, digest=hashlib.sha256(content).hexdigest())
    try:
        if format == "pdf":
            _parse_pdf(document, content)
        elif format == "csv":
            _parse_csv(document, content)
        elif format == "xlsx":
            _parse_xlsx(document, content)
        else:
            _parse_docx(document, content)
        document.status = ("partial" if set(document.warnings) - INFORMATIONAL - TEXT_ONLY else
                           "complete_text_only" if set(document.warnings) & TEXT_ONLY else "complete")
    except _Rejected as error:
        document.warnings.append(str(error))
        document.status = "partial" if document.units_read else "unreadable"
    except UnicodeDecodeError:
        document.warnings.append("invalid_text")
    except Exception:
        document.warnings.append("corrupt_document")
        document.status = "partial" if document.units_read else "unreadable"
    document.text = "\n".join(fragment["text"] for fragment in document.fragments) + document.text
    if document.status in {"complete", "complete_text_only"} and not document.text.strip():
        document.warnings.append("empty_document")
        document.status = "unreadable"
    document.warnings = list(dict.fromkeys(document.warnings))
    return document


def extract_document(uploaded_file):
    from django.conf import settings
    from django.core.exceptions import ValidationError
    from .file_security import (
        FILE_SIGNATURES, MAX_FILE_SIZES, get_file_extension, has_file_signature, sanitize_filename,
        validate_file_signature, validate_file_size,
    )

    document = Document(
        format=get_file_extension(uploaded_file.name),
        file_name=sanitize_filename(uploaded_file.name),
        file_size=uploaded_file.size,
    )
    if document.format not in TEXT_FORMATS | MIME_TYPES.keys():
        document.warnings.append("unsupported_format")
        return document
    mime = MIME_TYPES.get(document.format, "text/plain")
    try:
        validate_file_size(document.file_size, mime)
        uploaded_file.seek(0)
        content = uploaded_file.read(MAX_FILE_SIZES[mime] + 1)
        document.file_size = len(content)
        validate_file_size(document.file_size, mime)
    except ValidationError:
        document.warnings.append("file_size_limit")
        return document
    except (OSError, ValueError):
        document.warnings.append("file_read_error")
        return document
    document.digest = hashlib.sha256(content).hexdigest()
    try:
        validate_file_signature(content, mime)
        if document.format in MIME_TYPES and not any(
            content.startswith(signature) for signature in FILE_SIGNATURES[mime]
        ):
            raise ValidationError("missing signature")
        if document.format in TEXT_FORMATS and any(
            has_file_signature(content, known_mime)
            for known_mime in FILE_SIGNATURES
        ):
            raise ValidationError("binary signature")
    except ValidationError:
        document.warnings.append("signature_mismatch")
        return document
    if document.format in TEXT_FORMATS - {"csv"}:
        _parse_text(document, content)
    else:
        try:
            result = subprocess.run(
                [sys.executable, __file__, "--parse", document.format],
                input=content, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                timeout=getattr(settings, "DOCUMENT_PARSER_TIMEOUT", 5), check=True,
            )
            parsed = json.loads(result.stdout)
            for attribute in ("text", "fragments", "status", "units_read", "units_total", "warnings", "disclosures"):
                setattr(document, attribute, parsed[attribute])
        except subprocess.TimeoutExpired:
            document.warnings.append("parser_timeout")
        except (subprocess.SubprocessError, OSError, ValueError, KeyError):
            document.warnings.append("parser_error")
    return document


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--parse":
        data = sys.stdin.buffer.read(10 * 1024 * 1024 + 1)
        print(json.dumps(asdict(_worker(sys.argv[2], data))))
