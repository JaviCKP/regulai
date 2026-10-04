"""HTTP privacy and monetary boundaries, with synthetic transports only."""
import json
import uuid
from datetime import timedelta
from unittest.mock import patch
from types import SimpleNamespace
import re
from decimal import Decimal
import io

from django.utils import timezone
from django.urls import reverse

from front.models import AuditLog, Conversation, Message, ObfuscationConfig, Prompt
from front.test_chat_flow import ChatFixture
from front import test_pipeline
from front.connectors import ChatResponse, DOCUMENT_RESPONSE_SCHEMA
from front.models import CompanyEngine, Engine, EngineModel, Role


class PipelinePrivacyTests(ChatFixture):
    configure = test_pipeline.PipelineTests.configure

    def test_reclassification_cannot_generate_without_the_mandatory_judge(self):
        self.configure()
        self.company.ai_config['destinations']['confidential'] = [f'engine:{self.credential.pk}']
        self.company.save()
        self.http.side_effect = test_pipeline.jev_http(choices={'category': 'confidential'})
        for path in ('/api/chat/send/', '/chat/send/', '/chat/'):
            with self.subTest(path=path):
                self.provider.reset_mock()
                self.http.reset_mock()
                response = self.send(path, 'Synthetic report')
                self.assertEqual(response.json()['error_type'], 'security_check_failed')
                self.provider.assert_not_called()
                self.assertEqual(self.http.call_count, 2)
                receipt = response.json()['receipt']
                self.assertEqual([call['purpose'] for call in receipt['calls']], ['obfuscation', 'classification'])
                self.assertEqual(receipt['calls'][0]['status'], 'completed')
                self.assertEqual(Decimal(receipt['total']['cost']), Decimal('0.00014'))
                self.assertEqual(receipt['total']['attempts'], 0)
                self.assertTrue(receipt['total']['complete'])

    def test_reclassification_cannot_generate_without_the_response_mask(self):
        self.configure(verification={'mode': 'local'})
        self.mask()
        self.company.ai_config['destinations']['confidential'] = [f'engine:{self.credential.pk}', 'jev']
        self.company.save()
        evaluator = test_pipeline.jev_http(choices={'category': 'confidential'})
        def post(url, **kwargs):
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': kwargs['json']['text']})
            return evaluator(url, **kwargs)
        self.http.side_effect = post
        for path in ('/api/chat/send/', '/chat/send/', '/chat/'):
            with self.subTest(path=path):
                self.provider.reset_mock()
                self.http.reset_mock()
                response = self.send(path, 'Synthetic report')
                self.assertEqual(response.json()['error_type'], 'security_check_failed')
                self.provider.assert_not_called()
                self.assertEqual(self.http.call_count, 2)
                receipt = response.json()['receipt']
                self.assertEqual([call['purpose'] for call in receipt['calls']], ['obfuscation', 'classification'])
                self.assertTrue(all(call['status'] == 'completed' for call in receipt['calls']))
                self.assertEqual(Decimal(receipt['total']['cost']), Decimal('0.00014'))
                self.assertEqual(receipt['total']['attempts'], 0)
                self.assertTrue(receipt['total']['complete'])

    def test_unknown_floor_matches_chat_and_semantic_simulation(self):
        self.profile.role = Role.objects.create(name='Admin')
        self.profile.save()
        self.configure(local_sensitivity='unknown', destinations={
            'confidential': [f'engine:{self.credential.pk}', 'jev']},
            routing={'mode': 'auto', 'candidates': [{'company_engine_id': self.credential.pk,
                'model_id': self.model.pk, 'level': 'economy', 'quality': {'general': 1.0}}]})
        self.http.side_effect = test_pipeline.jev_http(choices={'category': 'public'})
        for path in ('/api/chat/send/', '/chat/send/', '/chat/'):
            with self.subTest(path=path):
                response = self.send(path, 'Synthetic report')
                self.assertEqual(response.status_code, 302 if path == '/chat/' else 200)
                if path != '/chat/':
                    result = response.json()
                    self.assertEqual(result['status'], 'completed')
                    self.assertEqual(result['receipt']['policy']['sensitivity'], 'confidential')
                    self.assertEqual(Decimal(result['receipt']['total']['cost']), Decimal('0.00032'))
        self.http.reset_mock()
        self.provider.reset_mock()
        response = self.client.post(reverse('front:policy_simulator'), {'mode': 'semantic',
            'sample': 'Synthetic report', 'proposed_config': json.dumps(self.company.ai_config)})
        comparison = response.context['comparison']
        self.assertEqual(comparison['semantic']['status'], 'completed')
        for name in ('current', 'proposed'):
            self.assertEqual(comparison[name]['policy']['sensitivity'], 'confidential')
            self.assertEqual(comparison[name]['policy']['status'], 'allow')
            self.assertEqual(comparison[name]['route']['status'], 'selected')
        self.assertEqual(self.http.call_count, 2)
        self.provider.assert_not_called()
        self.assertEqual(comparison['calls'][-1]['purpose'], 'simulation_classification')
        self.assertEqual(Decimal(comparison['calls'][-1]['cost']), Decimal('0.00014'))

    def test_mask_deactivated_during_classification_stops_before_generation(self):
        self.configure(verification={'mode': 'local'})
        self.mask()
        mask = ObfuscationConfig.objects.get(company=self.company)
        evaluator = test_pipeline.jev_http()
        def post(url, **kwargs):
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': kwargs['json']['text']})
            ObfuscationConfig.objects.filter(pk=mask.pk).update(active=False)
            return evaluator(url, **kwargs)
        self.http.side_effect = post
        for path in ('/api/chat/send/', '/chat/send/', '/chat/'):
            with self.subTest(path=path):
                ObfuscationConfig.objects.filter(pk=mask.pk).update(active=True)
                self.provider.reset_mock()
                self.http.reset_mock()
                response = self.send(path, 'Synthetic report')
                self.assertEqual(response.json()['error_type'], 'authorization_changed')
                self.provider.assert_not_called()
                self.assertEqual(self.http.call_count, 2)
                receipt = response.json()['receipt']
                self.assertEqual([call['purpose'] for call in receipt['calls']], ['obfuscation', 'classification'])
                self.assertEqual(Decimal(receipt['total']['cost']), Decimal('0.00014'))
                self.assertEqual(receipt['total']['attempts'], 0)
                self.assertTrue(receipt['total']['complete'])

    def mask(self, *, bound='0', malformed=False):
        config = self.required_mask
        if f'obfuscation:{config.pk}' not in self.company.ai_config['destinations']['internal']:
            self.company.ai_config['destinations']['internal'].append(f'obfuscation:{config.pk}')
        self.company.ai_config['masking'] = {'max_call_cost': bound}
        self.company.save()
        evaluator = test_pipeline.jev_http()
        def post(url, **kwargs):
            if url.endswith('/api/mask'):
                pending = Prompt.objects.get(status='processing')
                self.assertEqual(pending.decision['calls'][-1]['status'], 'started')
                text = kwargs['json']['text']
                self.assertIsInstance(json.loads(text), list)
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': '[]' if malformed else text.replace('12345678Z', '<DNI_1>')})
            self.assertNotIn('12345678Z', kwargs['data'].decode())
            return evaluator(url, **kwargs)
        self.http.side_effect = post

    def test_structured_masking_preserves_escapes_and_all_segments_in_every_http_entry(self):
        self.configure()
        self.company.system_prompt = 'Policy "12345678Z" \\ [END]'
        self.company.save()
        self.mask()
        for path in ('/chat/send/', '/api/chat/send/', '/chat/'):
            response = self.send(path, '12345678Z\n[END] " ], [ " 12345678Z')
            self.assertEqual(response.status_code, 302 if path == '/chat/' else 200)
            messages = self.provider.call_args.args[0]
            self.assertEqual(messages[0].content, 'Policy "<DNI_1>" \\ [END]')
            self.assertEqual(messages[-1].content, '<DNI_1>\n[END] " ], [ " <DNI_1>')
            completed = Prompt.objects.latest('pk')
            self.assertEqual(completed.decision['calls'][0]['cost'], '0')
            self.assertTrue(completed.decision['total']['complete'])
            self.assertNotIn('12345678Z', completed.user_message.content)

    def test_invalid_mask_segments_stop_before_classifier_or_generator(self):
        self.configure()
        self.mask(malformed=True)
        result = self.send('/api/chat/send/', '12345678Z').json()
        self.assertEqual(result['status'], 'failed')
        self.assertEqual(self.http.call_count, 1)
        self.provider.assert_not_called()

    def test_mask_cost_reserve_prevents_any_call_when_budget_is_insufficient(self):
        self.configure(limits={'budget': '0.01'})
        self.mask(bound='0.02')
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['error_type'], 'budget_exhausted')
        self.http.assert_not_called()
        self.provider.assert_not_called()

    def test_history_window_omits_whole_old_pairs_from_every_destination(self):
        self.configure(history={'max_turns': 1, 'max_bytes': 100})
        first = self.send('/api/chat/send/', 'OLD_SECRETLESS_TURN').json()
        second = self.send('/api/chat/send/', 'RECENT_TURN', conversation_id=first['conversation_id']).json()
        self.mask()
        result = self.send('/api/chat/send/', 'NEXT_TURN', conversation_id=first['conversation_id']).json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(result['receipt']['history'], {'sent_turns': 1, 'omitted_turns': 1,
            'sent_bytes': 25, 'omitted_bytes': 33})
        self.assertNotIn('OLD_SECRETLESS_TURN', str(self.provider.call_args))
        self.assertNotIn('OLD_SECRETLESS_TURN', str(self.http.call_args_list[-3:]))
        self.assertIn('RECENT_TURN', str(self.provider.call_args))

    def test_new_cancel_is_idempotent_without_conversation_or_message(self):
        request_id = str(uuid.uuid4())
        for path in ('/api/chat/send/', '/chat/send/'):
            result = self.send(path, '', action='cancel', request_id=request_id).json()
            self.assertEqual(result['status'], 'cancelled', result)
        self.assertEqual(Conversation.objects.count(), 0)
        self.assertEqual(Message.objects.count(), 0)
        self.provider.assert_not_called()

    def test_mask_expansion_cannot_exceed_effective_history_byte_limit(self):
        self.configure(history={'max_turns': 1, 'max_bytes': 25})
        first = self.send('/api/chat/send/', 'RECENT_TURN').json()
        self.mask()
        original_transport = self.http.side_effect
        def expanded(url, **kwargs):
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': kwargs['json']['text'].replace('Correct answer', 'X' * 50)})
            return original_transport(url, **kwargs)
        self.http.side_effect = expanded
        self.http.reset_mock()
        self.provider.reset_mock()
        result = self.send('/api/chat/send/', 'NEXT_TURN', conversation_id=first['conversation_id']).json()
        self.assertEqual(result.get('error_type'), 'context_overflow', result)
        self.assertEqual(self.http.call_count, 1)
        self.provider.assert_not_called()

    def test_processing_is_not_marked_uncertain_while_max_deadline_can_still_run(self):
        result = self.send('/api/chat/send/').json()
        pending = Prompt.objects.get(request_id=result['request_id'])
        Prompt.objects.filter(pk=pending.pk).update(status='processing', updated_at=timezone.now()-timedelta(seconds=181))
        repeated = self.send('/api/chat/send/', request_id=result['request_id']).json()
        self.assertEqual(repeated['status'], 'processing')
        Prompt.objects.filter(pk=pending.pk).update(updated_at=timezone.now()-timedelta(seconds=241))
        repeated = self.send('/api/chat/send/', request_id=result['request_id']).json()
        self.assertEqual(repeated['status'], 'outcome_unknown')
        self.assertEqual(self.provider.call_count, 1)

    def fallback(self, *, cross_engine=False, allow=False):
        engine = Engine.objects.create(name='Second engine', connector_type='openai') if cross_engine else self.engine
        credential = CompanyEngine.objects.create(company=self.company, engine=engine, api_key='synthetic-other') if cross_engine else self.credential
        stronger = EngineModel.objects.create(engine=engine, name='Strong', model_id='synthetic-strong',
            context_window=32768, input_price=2, output_price=3)
        self.company.ai_config['destinations']['internal'] += [f'engine:{credential.pk}']
        self.company.ai_config['routing'] = {'mode': 'manual', 'allow_cross_engine_fallback': allow, 'candidates': [
            {'company_engine_id': self.credential.pk, 'model_id': self.model.pk, 'level': 'economy', 'quality': {'general': .5}},
            {'company_engine_id': credential.pk, 'model_id': stronger.pk, 'level': 'strong', 'quality': {'general': .9}}]}
        self.company.save()
        return credential

    def test_retry_sends_previous_answer_and_concrete_failed_checks_with_valid_roles(self):
        self.configure()
        self.fallback()
        self.http.side_effect = test_pipeline.jev_http(failures={'failure_contradiction': 1})
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'failed', result)
        self.assertEqual(self.provider.call_count, 2)
        messages = self.provider.call_args.args[0]
        self.assertEqual([m.role for m in messages], ['user', 'assistant', 'user'])
        self.assertEqual(messages[-2].content, 'Correct answer')
        self.assertIn('failure_contradiction', messages[-1].content)
        self.assertNotIn('failure_omission', messages[-1].content)

    def test_manual_fallback_stays_with_selected_engine_unless_enabled(self):
        self.configure()
        other = self.fallback(cross_engine=True)
        self.http.side_effect = test_pipeline.jev_http(failures={'failure_contradiction': 1})
        result = self.send('/api/chat/send/').json()
        self.assertEqual(self.provider.call_count, 1, result)
        self.company.ai_config['routing']['allow_cross_engine_fallback'] = True
        self.company.save()
        self.provider.reset_mock()
        result = self.send('/api/chat/send/').json()
        self.assertEqual(self.provider.call_count, 2, result)
        self.assertEqual(result['receipt']['route']['company_engine_id'], other.pk)

    def test_busy_provider_retry_does_not_invent_failed_checks_or_an_answer(self):
        self.configure()
        self.fallback()
        error = RuntimeError('sensitive provider message')
        error.status_code = 503
        self.provider.side_effect = [error, ChatResponse('Recovered', 30, 'synthetic-v1', 'stop', 20, 10)]
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual([m.role for m in self.provider.call_args.args[0]], ['user'])
        self.assertEqual(self.provider.call_args.args[0][0].content, 'Hola')

    def test_semantic_simulator_masks_sample_and_audits_each_intention_before_network(self):
        self.user.is_superuser = True
        self.user.save()
        self.configure()
        self.mask()
        evaluator = test_pipeline.jev_http()
        def post(url, **kwargs):
            audit = json.loads(AuditLog.objects.latest('pk').details)
            self.assertEqual(audit['call']['status'], 'started')
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': kwargs['json']['text'].replace('12345678Z', '<DNI_1>')})
            self.assertNotIn('12345678Z', kwargs['data'].decode())
            return evaluator(url, **kwargs)
        self.http.side_effect = post
        result = self.client.post(reverse('front:policy_simulator'), {'mode': 'semantic',
            'sample': 'My id is 12345678Z', 'proposed_config': json.dumps(self.company.ai_config)})
        comparison = result.context['comparison']
        self.assertEqual(comparison['semantic']['status'], 'completed', comparison)
        self.assertEqual([item['purpose'] for item in comparison['calls']], ['simulation_obfuscation', 'simulation_classification'])
        self.assertEqual(comparison['calls'][0]['cost'], '0')
        self.assertEqual(AuditLog.objects.filter(action='policy_simulate').count(), 2)
        self.assertNotIn('12345678Z', str(list(AuditLog.objects.values_list('details', flat=True))))
        self.provider.assert_not_called()

    def test_masking_cannot_launder_an_invented_document_quote(self):
        self.configure(limits={'budget': None})
        self.mask()
        evaluator = test_pipeline.jev_http(choices={'task': 'document_question'})
        def post(url, **kwargs):
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': re.sub(r'Code: \d+', 'Code: <NUM_1>', kwargs['json']['text'])})
            return evaluator(url, **kwargs)
        self.http.side_effect = post
        self.provider.return_value = ChatResponse(json.dumps({'claims': [{'text': 'The code is valid.',
            'citations': [{'id': 'S1', 'quote': 'Code: 99'}]}]}), 30, 'synthetic-v1', 'stop', 20, 10)
        from django.core.files.uploadedfile import SimpleUploadedFile
        result = self.send('/api/chat/send/', 'What is the code?', file=SimpleUploadedFile('code.txt', b'Code: 42')).json()
        self.assertEqual(result['status'], 'failed', result)
        self.assertEqual(result['receipt']['verification']['local_status'], 'failed')
        self.assertEqual([c['purpose'] for c in result['receipt']['calls']], ['obfuscation', 'classification', 'generation', 'obfuscation'])
        self.assertNotIn('assistant_message', result)

    def test_judge_capacity_accepts_exact_limit_and_rejects_one_more_byte_before_network(self):
        self.configure()
        self.model.context_window = 1_000_000
        self.model.save()
        base = self.send('/api/chat/send/', 'x').json()['receipt']['limits']['judge_input_bytes_bound']
        amount = 1 + 98304 - base
        self.http.reset_mock()
        self.provider.reset_mock()
        result = self.send('/api/chat/send/', 'x' * amount).json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(result['receipt']['limits']['judge_input_bytes_bound'], 98304)
        self.http.reset_mock()
        self.provider.reset_mock()
        result = self.send('/api/chat/send/', 'x' * (amount + 1)).json()
        self.assertEqual(result['error_type'], 'evaluation_input_too_large', result)
        self.http.assert_not_called()
        self.provider.assert_not_called()

    def test_generated_output_limit_is_explicit_and_accepts_boundary_with_utf8(self):
        self.configure(classification='local', verification={'mode': 'local'}, routing={'max_output_tokens': 2})
        self.provider.return_value = ChatResponse('é' * 8, 30, 'synthetic-v1', 'stop', 20, 10)
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(result['receipt']['limits']['response_bytes'], 16)
        self.provider.return_value = ChatResponse('é' * 8 + 'x', 30, 'synthetic-v1', 'stop', 20, 10)
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['error_type'], 'evaluation_input_too_large', result)
        self.assertNotIn('assistant_message', result)

    def test_internal_failure_after_completed_generation_keeps_ledger_and_safe_logs(self):
        self.provider.return_value = ChatResponse(None, 10, 'synthetic-v1', 'stop')
        with self.assertLogs('front.chat_service', level='ERROR') as captured:
            result = self.send('/api/chat/send/', 'PRIVATE_PROMPT_FOR_LOG_TEST').json()
        self.assertEqual(result['error_type'], 'internal_error', result)
        self.assertEqual(result['receipt']['calls'][0]['status'], 'completed')
        self.assertNotIn('PRIVATE_PROMPT_FOR_LOG_TEST', str(captured.output))
        self.assertEqual(captured.records[0].code, 'internal_error')
        self.assertEqual(captured.records[0].request_id, result['request_id'])

    def test_anthropic_document_schema_counts_before_paid_boundaries(self):
        self.configure(classification='local')
        self.engine.connector_type = 'anthropic'
        self.engine.save()
        answer = ChatResponse(json.dumps({'claims': [{'text': 'The approved period is seventeen days.',
            'citations': [{'id': 'S1', 'quote': 'The approved period is seventeen days.'}]}]}),
            30, 'synthetic-v1', 'stop', 20, 10)
        from django.core.files.uploadedfile import SimpleUploadedFile
        with patch('front.connectors.AnthropicConnector.chat', return_value=answer) as provider:
            result = self.send('/api/chat/send/', 'What is the period?',
                file=SimpleUploadedFile('period.txt', b'The approved period is seventeen days.')).json()
            self.assertEqual(result['status'], 'completed', result)
            messages = provider.call_args.args[0]
            input_bound = result['receipt']['route']['input_tokens_bound']
            self.assertEqual(input_bound, sum(len(m.content.encode()) + len(m.role.encode()) + 256 for m in messages)
                + len(json.dumps({'output_config': {'format': {'type': 'json_schema', 'schema': DOCUMENT_RESPONSE_SCHEMA}}}).encode()))
            self.assertTrue(provider.call_args.kwargs['supports_json'])
            self.assertFalse(provider.call_args.kwargs['reasoning'])
            self.assertEqual(provider.call_args.kwargs['temperature'], .1)
            # Input includes a known 496-byte envelope. A budget one micro-dollar
            # below that complete generation/checking bound must make zero calls.
            self.company.ai_config['limits']['budget'] = str(Decimal(input_bound + 2000) / Decimal(1_000_000)
                + Decimal('.01') - Decimal('.000001'))
            self.company.save()
            self.http.reset_mock()
            provider.reset_mock()
            result = self.send('/api/chat/send/', 'What is the period?',
                file=SimpleUploadedFile('period.txt', b'The approved period is seventeen days.')).json()
            self.assertEqual(result['error_type'], 'budget_exhausted', result)
            self.http.assert_not_called()
            provider.assert_not_called()

    def test_zero_mask_cost_is_known_for_every_decimal_representation(self):
        self.configure()
        self.mask(bound='0.000')
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertTrue(result['receipt']['total']['complete'], result)
        self.assertEqual([call['cost'] for call in result['receipt']['calls'] if call['purpose']=='obfuscation'], ['0', '0'])

    def test_text_only_document_coverage_requires_confirmation_without_sending_omitted_image(self):
        self.configure()
        from docx import Document
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        doc = Document()
        doc.add_paragraph('The approved period is seventeen days.')
        image = io.BytesIO()
        Image.new('RGB', (1, 1), 'white').save(image, format='PNG')
        image.seek(0)
        doc.add_picture(image)
        raw = io.BytesIO()
        doc.save(raw)
        request_id = str(uuid.uuid4())
        first = self.send('/api/chat/send/', 'What is the period?', request_id=request_id,
            file=SimpleUploadedFile('period.docx', raw.getvalue())).json()
        self.assertEqual(first['status'], 'pending_consent', first)
        self.assertEqual(first['receipt']['coverage']['status'], 'complete_text_only')
        self.assertEqual(first['consent_reason'], 'coverage_incomplete')
        self.http.assert_not_called()
        self.provider.assert_not_called()
        self.provider.return_value = ChatResponse(json.dumps({'claims': [{'text': 'The approved period is seventeen days.',
            'citations': [{'id': 'S1', 'quote': 'The approved period is seventeen days.'}]}]}),
            30, 'synthetic-v1', 'stop', 20, 10)
        result = self.send('/api/chat/send/', 'What is the period?', request_id=request_id,
            consent_token=first['consent_token'], file=SimpleUploadedFile('period.docx', raw.getvalue())).json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(result['user_message']['id'], first['user_message']['id'])
        self.assertNotIn('image/png', str(self.provider.call_args))

    def test_timeout_configuration_reaches_all_network_boundaries(self):
        self.configure(limits={'budget': '1', 'deadline_seconds': 120, 'call_timeout_seconds': 70})
        self.mask()
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(self.provider.call_args.kwargs['timeout'], 70)
        for call in self.http.call_args_list:
            timeout = call.kwargs['timeout']
            self.assertEqual(sum(timeout) if isinstance(timeout, tuple) else timeout, 70)

    def test_uncertain_503_is_not_automatically_retried(self):
        self.configure()
        self.fallback()
        error = RuntimeError('sensitive error')
        error.status_code, error.outcome_unknown = 503, True
        self.provider.side_effect = error
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'outcome_unknown', result)
        self.assertEqual(self.provider.call_count, 1)

    def test_receipts_identify_actual_classification_rubrics_despite_legacy_config_label(self):
        self.configure()
        self.company.ai_config['jev']['rubric_version'] = 'sensitivity-v1'
        self.company.save()
        result = self.send('/api/chat/send/').json()
        self.assertEqual(result['status'], 'completed', result)
        self.assertEqual(result['receipt']['calls'][1]['rubric_version'], 'sensitivity-v2')
        self.assertEqual(result['receipt']['calls'][1]['task_rubric_version'], 'task-v2')
        self.assertEqual(result['decision']['rubric_version'], 'sensitivity-v2')

    def test_company_admin_can_run_masked_semantics_without_global_flags(self):
        self.configure()
        self.mask()
        self.profile.role = Role.objects.create(name='Admin')
        self.profile.save()
        evaluator = test_pipeline.jev_http()
        def post(url, **kwargs):
            if url.endswith('/api/mask'):
                return SimpleNamespace(raise_for_status=lambda: None,
                    json=lambda: {'masked': kwargs['json']['text']})
            return evaluator(url, **kwargs)
        self.http.side_effect = post
        response = self.client.post(reverse('front:policy_simulator'), {'mode': 'semantic',
            'sample': 'Synthetic text', 'proposed_config': json.dumps(self.company.ai_config)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['comparison']['semantic']['status'], 'completed')
        self.assertEqual(response.context['comparison']['calls'][-1]['rubric_version'], 'sensitivity-v2')
        self.assertFalse(self.user.is_superuser)
        self.assertFalse(self.user.is_staff)
