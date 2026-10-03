"""Create small synthetic extraction fixtures; never label these as Office exports.

Use the Codex bundled Python runtime (reportlab, python-docx, openpyxl, Pillow,
pypdf). Output must be a NEW directory. Expected results live separately and
must be reviewed before extracting. OOXML cache/revision injection is deliberate
test construction, not a claim about a native application's behavior.
"""
import argparse
from datetime import datetime
import hashlib
from importlib.metadata import version
import io
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas


STAMP = datetime(2026, 10, 3, 0, 0, 0)


def office_bytes(content, replacements=None):
    """Stable ZIP metadata and optional explicit XML edits for synthetic fixtures."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == 'docProps/core.xml':
                # openpyxl rewrites modified on save; freeze it explicitly.
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]+', rb'\g<1>2026-10-03T00:00:00Z', data)
            if replacements and name in replacements:
                data = replacements[name](data)
            info = zipfile.ZipInfo(name, (2026, 10, 3, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    return output.getvalue()


def pdf(pages):
    target = io.BytesIO()
    canvas = Canvas(target, pagesize=(595, 842), invariant=1, pageCompression=1)
    canvas.setTitle('Synthetic document pilot')
    canvas.setAuthor('regulAI synthetic fixture generator')
    for lines in pages:
        canvas.setFont('Helvetica', 12)
        for line, y in zip(lines, range(790, 70, -25)):
            canvas.drawString(50, y, line)
        canvas.showPage()
    canvas.save()
    return target.getvalue()


def synthetic_image():
    image = Image.new('RGB', (640, 180), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((2, 2, 637, 177), outline='#193452', width=3)
    draw.text((24, 45), 'VISUAL ONLY 73', fill='black', font_size=42)
    draw.text((24, 115), 'Synthetic raster fixture - no OCR', fill='black', font_size=23)
    target = io.BytesIO()
    image.save(target, 'PNG')
    return target.getvalue()


def image_pdf(bitmap, *, text):
    target = io.BytesIO()
    canvas = Canvas(target, pagesize=(595, 842), invariant=1, pageCompression=1)
    canvas.setTitle('Synthetic raster coverage fixture')
    if text:
        canvas.setFont('Helvetica', 12)
        canvas.drawString(50, 790, 'El contrato cuesta 10 euros.')
    canvas.drawImage(ImageReader(io.BytesIO(bitmap)), 50, 570, width=480, height=135)
    canvas.showPage()
    canvas.save()
    return target.getvalue()


def word(title):
    document = Document()
    document.core_properties.title = title
    document.core_properties.author = 'regulAI synthetic fixture generator'
    document.core_properties.created = document.core_properties.modified = STAMP
    document.add_paragraph(title, style='Title')
    return document


def word_bytes(document):
    target = io.BytesIO()
    document.save(target)
    return office_bytes(target.getvalue())


def revision(kind, value, revision_id):
    node = OxmlElement('w:' + kind)
    node.set(qn('w:id'), str(revision_id))
    node.set(qn('w:author'), 'Synthetic fixture')
    node.set(qn('w:date'), '2026-10-03T00:00:00Z')
    run = OxmlElement('w:r')
    text = OxmlElement('w:delText' if kind == 'del' else 'w:t')
    text.text = value
    run.append(text)
    node.append(run)
    return node


def numbering_word():
    document = word('Contrato con numeracion decimal')
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), '42')
    for index in (0, 1):
        level = OxmlElement('w:lvl')
        level.set(qn('w:ilvl'), str(index))
        for name, value in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.' if index == 0 else '%1.%2.')]:
            child = OxmlElement('w:' + name)
            child.set(qn('w:val'), value)
            level.append(child)
        abstract.append(level)
    document.part.numbering_part.element.append(abstract)
    number = OxmlElement('w:num')
    number.set(qn('w:numId'), '42')
    reference = OxmlElement('w:abstractNumId')
    reference.set(qn('w:val'), '42')
    number.append(reference)
    document.part.numbering_part.element.append(number)
    for level, text in [(0, 'Servicio'), (1, 'Precio de 10 euros'), (1, 'Revision anual'), (0, 'Terminacion'), (1, 'Preaviso de 17 dias')]:
        properties = document.add_paragraph(text)._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_numId().val = 42
        properties.get_or_add_ilvl().val = level
    return word_bytes(document)


def formula_workbook(*, cache):
    workbook = Workbook()
    workbook.properties.creator = 'regulAI synthetic fixture generator'
    workbook.properties.created = workbook.properties.modified = STAMP
    workbook.calculation = CalcProperties(fullCalcOnLoad=False, forceFullCalc=False)
    sheet = workbook.active
    sheet.title = 'Costes'
    sheet.append(['Unidades', 'Precio EUR', 'Total EUR'])
    sheet.append([2, 5, '=A2*B2'])
    for column in 'ABC':
        sheet.column_dimensions[column].width = 20
        sheet[column + '1'].font = Font(bold=True)
    sheet['B2'].number_format = sheet['C2'].number_format = '0.00'
    target = io.BytesIO()
    workbook.save(target)
    edits = None
    if cache:
        edits = {'xl/worksheets/sheet1.xml': lambda data: data.replace(b'<f>A2*B2</f><v></v>', b'<f>A2*B2</f><v>10</v>')}
    return office_bytes(target.getvalue(), edits)


def build(output):
    output.mkdir(parents=True, exist_ok=False)
    bitmap = synthetic_image()
    assets = {
        'textual.pdf': pdf([['Contrato sintetico', 'Contenido de prueba sin datos personales.'], ['El preaviso final es de 17 dias.']]),
        'table.pdf': pdf([['Tabla sintetica de precios', 'Plan             Precio EUR', 'Basico           10', 'Avanzado         25']]),
        'text-and-image.pdf': image_pdf(bitmap, text=True),
        'simulated-scan.pdf': image_pdf(bitmap, text=False),
        'multilevel-numbering.docx': numbering_word(),
        'quoted.csv': 'Plan;Precio EUR;Nota\nBasico;10;"Primera linea\nSegunda linea"\nAvanzado;25;Fin\n'.encode('utf-8'),
        'formula-cached.xlsx': formula_workbook(cache=True),
        'formula-uncached.xlsx': formula_workbook(cache=False),
    }
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    target = io.BytesIO()
    writer.write(target)
    assets['blank.pdf'] = target.getvalue()
    reader = PdfReader(io.BytesIO(pdf([['Inicio del contrato'], ['Clausula final: el preaviso es de 17 dias.']])))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    writer.add_blank_page(width=595, height=842)
    writer.add_page(reader.pages[1])
    target = io.BytesIO()
    writer.write(target)
    assets['blank-and-final-clause.pdf'] = target.getvalue()
    document = word('Contrato con cambios pendientes')
    paragraph = document.add_paragraph('El preaviso es de ')
    paragraph._p.append(revision('del', '30 dias', 1))
    paragraph._p.append(revision('ins', '17 dias', 2))
    paragraph.add_run('.')
    assets['tracked-changes.docx'] = word_bytes(document)
    document = word('Contrato con imagen sintetica')
    document.add_paragraph('El contrato cuesta 10 euros.')
    document.add_picture(io.BytesIO(bitmap), width=Inches(5))
    assets['text-and-image.docx'] = word_bytes(document)
    manifest = {
        'schema_version': 1,
        'corpus_id': 'pilot-documents-20261003',
        'content_origin': 'entirely synthetic agent-authored fixtures',
        'producer': 'Python libraries and explicit OOXML edits; never opened or saved by Word, Excel or LibreOffice',
        'human_review': 'not performed',
        'libraries': {name: version(name) for name in ['reportlab', 'python-docx', 'openpyxl', 'Pillow', 'pypdf']},
        'generator_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        'expectations_sha256': hashlib.sha256(Path(__file__).with_name('expectations.json').read_bytes()).hexdigest(),
        'files': [],
    }
    for name, content in sorted(assets.items()):
        (output / name).write_bytes(content)
        manifest['files'].append({'file': name, 'bytes': len(content), 'sha256': hashlib.sha256(content).hexdigest(), 'native_application_saved': False})
    (output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'files': len(assets), 'output': str(output), 'expectations_sha256': manifest['expectations_sha256']}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    build(parser.parse_args().output.resolve())
