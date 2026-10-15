"""Deterministic acceptance fixture, NOT evidence of live-model autonomy.

This newly written program is transferred as a tool argument to the actual
isolated runner. No runtime tool invokes it by name or installs it in an image.
The text exercises pagination/editing; factual conclusions require live review.
"""
from pathlib import Path
import csv
import hashlib
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from PIL import Image, ImageDraw, ImageFont

OUT = Path('/workspace/output')
DATA = Path('/workspace/work/dataset.csv')
rows = list(csv.DictReader(DATA.open()))
values = [int(row['load_mwh']) for row in rows]
total, average = sum(values), sum(values) / len(values)
chart = Image.new('RGB', (980, 500), 'white')
draw = ImageDraw.Draw(chart)
font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 22)
draw.text((45, 22), 'Dataset controlado de aceptación · demanda (MWh)', fill='#182c4e', font=font)
draw.line((70, 400, 920, 400), fill='#52617a', width=2)
for index, row in enumerate(rows):
    left, top = 140 + index * 260, 400 - values[index] * 10
    draw.rectangle((left, top, left + 140, 400), fill='#2554a4')
    draw.text((left, top - 32), str(values[index]), fill='#182c4e', font=font)
    draw.text((left, 420), row['region'], fill='#182c4e', font=font)
chart.save(OUT / 'chart.png')

pages = [
('Computación cuántica y ciberseguridad', 'Resumen para la dirección', [
    'Este documento es un fixture de aceptación del agente profesional. Verifica la creación de un Word editable, su PDF de cinco páginas y la conservación de revisiones. No acredita investigación autónoma ni una evaluación de seguridad de una organización real.',
    'La migración a criptografía poscuántica necesita un inventario de dependencias, requisitos de interoperabilidad y pruebas de implementación. El plan debe distinguir los mecanismos de establecimiento de claves de los de firma digital. La elección de algoritmos no sustituye el análisis de protocolos, certificados, servicios y dispositivos.',
    'La dirección necesita responsables claros, criterios de priorización y evidencia reproducible. La primera decisión consiste en identificar dónde se utiliza criptografía y qué información requiere protección prolongada. Las fechas de sustitución deben venir de la política y de los proveedores concretos, sin inventar una fecha de llegada de un ordenador cuántico capaz de romper sistemas desplegados.',
    'Alcance de la prueba: contenido en español, jerarquía editorial, saltos de página, referencias visibles y exportación desde un único archivo fuente. Las conclusiones temáticas se incluyen como texto controlado para comprobar el recorrido documental.']),
('01 · Dependencias y exposición', 'Inventario antes de priorizar', [
    'El inventario debe cubrir aplicaciones, bibliotecas, certificados, equipos de red y conexiones entre servicios. Para cada elemento conviene registrar el propietario, el protocolo, el algoritmo observado, la posibilidad de actualización y la evidencia que respalda el registro.',
    'También hay que distinguir la vida útil del sistema de la vida útil de los datos. Un servicio que se retirará pronto puede seguir almacenando información de larga duración. La prioridad depende de esa exposición y de las restricciones de migración, no de una puntuación inventada.',
    'Los proveedores deben aportar información verificable sobre su hoja de ruta. Una declaración comercial de compatibilidad debe contrastarse con versiones concretas, parámetros y pruebas de interoperabilidad. Si falta documentación, el estado correcto es pendiente de confirmación.',
    'Resultado esperado: un inventario con responsables, lagunas identificadas y fuentes enlazadas. Esta sección se conserva sin cambios cuando la prueba modifica posteriormente la recomendación operativa.']),
('02 · Arquitectura de transición', 'Cambios que se pueden comprobar', [
    'La agilidad criptográfica permite sustituir componentes sin rehacer toda la aplicación. Para evaluar una propuesta hay que revisar dónde se seleccionan algoritmos, cómo se negocian parámetros y cómo se gestionan las versiones incompatibles.',
    'Las pruebas deben medir tamaños de mensajes, compatibilidad, latencia y uso de memoria con los sistemas representativos. No se presentan aquí porcentajes ni mejoras de rendimiento: esos resultados necesitan mediciones realizadas sobre una configuración documentada.',
    'Una transición necesita un mecanismo de reversión y reglas para evitar degradaciones silenciosas. El equipo debe observar qué algoritmo se utilizó realmente, qué fallo ocurrió y qué política autorizó cualquier alternativa.',
    'Las decisiones se documentan con fecha, versión y alcance. Las excepciones temporales deben tener propietario y condición de cierre, de modo que la revisión posterior pueda distinguir una dependencia resuelta de una deuda pendiente.']),
('03 · Verificación y evidencias', 'Del documento a la prueba', [
    'Las verificaciones de integridad, instrucciones, maquetación, fuentes y seguridad responden a preguntas distintas. Abrir correctamente un DOCX demuestra una propiedad del archivo; no demuestra que sus afirmaciones sean correctas ni que todos los párrafos resulten legibles.',
    'La revisión visual requiere observar las páginas renderizadas. La extracción de texto sirve para comprobar coherencia aproximada entre Word y PDF, pero no detecta por sí sola una tabla recortada o un contraste insuficiente.',
    'Las fuentes deben conservar su URL solicitada y final, el instante de recuperación y el fragmento consultado. Una URL descubierta por el buscador no se presenta como una página leída. Las discrepancias y los accesos fallidos se conservan en el resultado.',
    'Referencia temática para la prueba: https://www.nist.gov/pqc . La simulación del proveedor utiliza evidencia controlada y se etiqueta por separado. Esta referencia no implica que un modelo vivo haya contrastado cada afirmación.']),
('04 · Recomendación operativa', 'Siguiente paso verificable', [
    'Recomendación inicial: preparar el inventario criptográfico y asignar responsables de revisión.',
    'Una vez conocido el inventario, el equipo puede seleccionar un piloto de alcance reducido con criterios observables: compatibilidad funcional, errores, rendimiento medido y capacidad de reversión. El resultado debe conservar tanto los éxitos como los fallos.',
    'El cierre del piloto exige una revisión de los cambios y de las dependencias que permanecen abiertas. Las aprobaciones técnicas y de negocio se registran por separado. La ausencia de un dato no se transforma en una cifra estimada sin método y justificación.',
    'Este fixture termina con un plan de trabajo, no con una promesa de seguridad universal. La validación con modelo y buscador reales queda fuera de esta ejecución sin gasto. El informe, sus revisiones y los hashes permiten comprobar el comportamiento local del producto.'])]
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.85)
normal = doc.styles['Normal']
normal.font.name = 'Liberation Sans'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(11)
normal.paragraph_format.line_spacing = 1.15
for index, (title, subtitle, paragraphs) in enumerate(pages):
    if index:
        doc.add_page_break()
    doc.add_paragraph('REGULAY · PRUEBA DOCUMENTAL CONTROLADA', 'Subtitle')
    doc.add_heading(title, 0 if not index else 1)
    doc.add_heading(subtitle, 2)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    note = doc.add_paragraph(f'Página temática {index + 1} de 5 · Evidencia de formato y edición')
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor.from_string('52617A')
doc.save(OUT / 'quantum-v1.docx')
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Recomendación inicial:'):
        paragraph.text = 'Recomendación revisada: validar el inventario con los responsables de cada servicio antes de seleccionar el piloto.'
doc.save(OUT / 'quantum-v2.docx')

book = Workbook()
book.active.title = 'Dataset'
book.active.append(['region', 'load_mwh'])
for row in rows:
    book.active.append([row['region'], int(row['load_mwh'])])
summary = book.create_sheet('Cálculos')
summary.append(['Medida', 'Valor', 'Unidad'])
summary.append(['Total', '=SUM(Dataset!B2:B4)', 'MWh'])
summary.append(['Media', '=AVERAGE(Dataset!B2:B4)', 'MWh'])
metadata = book.create_sheet('Procedencia')
metadata.append(['URL', 'https://dataset.example/energy.csv'])
metadata.append(['SHA256', hashlib.sha256(DATA.read_bytes()).hexdigest()])
metadata.append(['Estado', 'Dataset controlado de aceptación'])
book.save(OUT / 'analysis-v1.xlsx')
summary['B3'] = '=ROUND(AVERAGE(Dataset!B2:B4),2)'
book.save(OUT / 'analysis-v2.xlsx')

deck = Presentation()
for title, body in [('Plan de transición', 'Inventario, responsables y pruebas verificables.'),
        ('Próximo paso', 'Revisar el inventario con los equipos de cada servicio.')]:
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    title_shape, body_shape = slide.shapes.title, slide.placeholders[1]
    for shape, top, height, size in [(title_shape, .5, .9, 34), (body_shape, 2, 3.5, 24)]:
        shape.left, shape.top = SlideInches(.6), SlideInches(top)
        shape.width, shape.height = SlideInches(8.8), SlideInches(height)
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = SlidePt(size)
deck.save(OUT / 'deck-v1.pptx')
deck.slides[1].placeholders[1].text = 'Elegir el piloto después de validar dependencias y capacidad de reversión.'
deck.slides[1].placeholders[1].left = SlideInches(.8)
deck.slides[1].placeholders[1].top = SlideInches(2.2)
deck.slides[1].placeholders[1].width = SlideInches(8.4)
deck.slides[1].placeholders[1].height = SlideInches(3.3)
for paragraph in deck.slides[1].placeholders[1].text_frame.paragraphs:
    paragraph.font.size = SlidePt(24)
deck.save(OUT / 'deck-v2.pptx')
(OUT / 'analysis.json').write_text(json.dumps({'total_mwh': total, 'average_mwh': average,
    'source': 'https://dataset.example/energy.csv', 'input_sha256': hashlib.sha256(DATA.read_bytes()).hexdigest()}))
print(json.dumps({'total_mwh': total, 'average_mwh': average, 'documents': 6}))
