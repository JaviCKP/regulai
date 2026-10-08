"""Only the failing OfficeCLI case: native chart with editable embedded data."""
import io
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from openpyxl import load_workbook
from pptx import Presentation

root = Path('/evidence/chart-comparison')
root.mkdir(exist_ok=False)
engine = '/app/.task51-branddocs-source/scripts/brandkit/cli.py'
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = 'Distribución por concepto'
slide.placeholders[1].text = 'Importes en euros'
deck.save(root/'template.pptx')
content = {'blocks': [{'type': 'heading', 'level': 1, 'runs': [{'t': 'Distribución por concepto'}]},
    {'type': 'chart', 'chart_type': 'barh', 'title': 'Importe (€)',
     'categories': ['Área Norte', 'Piñón español', 'Revisión técnica'],
     'series': [{'name': 'Importe (€)', 'values': [1.25, 2.5, 3.75]}]}]}
(root/'input.json').write_text(json.dumps(content))
log = []
for args in [['doctor'], ['extract', '--name', 'chart', '--template', str(root/'template.pptx'), '--scope', 'project'],
    ['generate', '--name', 'chart', '--input', str(root/'input.json'), '--output', str(root/'final.pptx'), '--scope', 'project', '--qa', 'strict']]:
    result = subprocess.run([sys.executable, engine, *args], cwd=root, capture_output=True, text=True, timeout=180)
    log.append({'args': args, 'code': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr})
(root/'commands.json').write_text(json.dumps(log, indent=2))
result = {'branddocs_commit': '97a6e384cb0664cd44ba8113c0b80c9761700670'}
if (root/'final.pptx').exists():
    deck = Presentation(root/'final.pptx')
    charts = [s.chart for slide in deck.slides for s in slide.shapes if s.has_chart]
    result['charts'] = len(charts)
    if charts:
        data = charts[0].part.chart_workbook.xlsx_part.blob
        result['embedded_workbook_values'] = list(load_workbook(io.BytesIO(data), data_only=True).active.values)
(root/'result.json').write_text(json.dumps(result, indent=2, ensure_ascii=False))
print(json.dumps(result))
