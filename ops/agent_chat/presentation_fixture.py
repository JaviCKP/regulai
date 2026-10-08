"""Author the synthetic corporate shell once; runtime edits it with OfficeCLI."""
from copy import deepcopy
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt


def create(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(33.867), Cm(19.05)
    for number, title in enumerate(['Informe económico', 'Detalle de importes', 'Distribución por concepto'], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('193B45' if number == 1 else 'F5F7F5')
        for name, value, x, y, w, h, size, font, color in [
            ('Title', title, 1.7, 3, 19 if number == 1 else 30, 5 if number == 1 else 2.5, 42 if number == 1 else 36, 'Liberation Serif', 'FFFFFF' if number == 1 else '193B45'),
            ('Footer', f'{number:02d}  /  REGULAY · Informe económico', 1.7, 17, 30, .9, 11, 'Liberation Sans', 'B9CDCE' if number == 1 else '60797E')]:
            shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
            shape.name = name
            shape.text_frame.word_wrap = True
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.text = value
            paragraph.font.name, paragraph.font.size = font, Pt(size)
            paragraph.font.color.rgb = RGBColor.from_string(color)
        slide.notes_slide.notes_text_frame.text = 'Informe económico. Fuente: archivo aportado.'
    brand = deck.slides[0].shapes.add_textbox(Cm(1.7), Cm(1), Cm(8), Cm(.9))
    brand.name = 'MasterBrand'
    brand.fill.solid()
    brand.fill.fore_color.rgb = RGBColor.from_string('193B45')
    paragraph = brand.text_frame.paragraphs[0]
    paragraph.text = 'REGULAY  /  INFORMES'
    paragraph.font.name, paragraph.font.size = 'Liberation Sans', Pt(12)
    paragraph.font.color.rgb = RGBColor.from_string('E4C38E')
    deck.slide_master.shapes._spTree.append(deepcopy(brand._element))
    brand._element.getparent().remove(brand._element)
    deck.save(path)


if __name__ == '__main__':
    import sys
    create(Path(sys.argv[1]))
