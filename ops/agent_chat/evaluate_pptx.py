"""Three pinned-OfficeCLI acceptance cases, final PPTX rendered by Impress."""
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

ROOT = Path('/evidence/officecli-final')
ROOT.mkdir(parents=True, exist_ok=False)
LOG = []
ENV = {**os.environ, 'OFFICECLI_NO_AUTO_RESIDENT': '1', 'DOTNET_CLI_TELEMETRY_OPTOUT': '1'}
BLUE = '17365D'
ROWS = [('Área Norte', 125), ('Piñón español', 250), ('Revisión técnica', 375)]


def cli(*args, strict=True):
    result = subprocess.run(['officecli', *map(str, args)], env=ENV, capture_output=True, text=True, timeout=30)
    LOG.append({'args': list(map(str, args)), 'code': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr})
    if strict and result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    return result


def add(file, slide, kind, **props):
    args = ['add', file, f'/slide[{slide}]' if slide else '/', '--type', kind]
    for key, value in props.items():
        args += ['--prop', f'{key}={value}']
    cli(*args)


def shape(file, slide, name, text, x, y, w, h, size=22, color=BLUE):
    add(file, slide, 'shape', name=name, text=text, x=f'{x}cm', y=f'{y}cm', width=f'{w}cm',
        height=f'{h}cm', font='Liberation Sans', size=size, color=color, fill='none', line='none')


def new_deck(file):
    cli('create', file)
    cli('set', file, '/', '--prop', 'slideSize=widescreen')
    for number, title in enumerate(['Informe económico de España', 'Importes por concepto', 'Distribución y revisión'], 1):
        add(file, 0, 'slide', layout='blank', background='FFFFFF')
        shape(file, number, 'Brand', 'REGULAY · ÁREA TÉCNICA', 1.5, 1.2, 29, .9, 12)
        shape(file, number, 'Title', title, 1.5, 2.7, 30, 3, 36)
        shape(file, number, 'Footer', f'{number} · Informe sintético · 2026', 1.5, 16.8, 29, .9, 12)
        add(file, number, 'notes', text='Datos sintéticos para comprobar conservación, edición y render final.')
    shape(file, 1, 'Body', 'Resumen de importes en euros. Año 2026: España, ñ, €, £ y $.', 1.5, 6, 30, 4, 24)
    shape(file, 1, 'Total', '7,50 €', 1.5, 11, 20, 3, 60)
    shape(file, 1, 'TotalLabel', 'Total de los tres conceptos', 1.5, 14.2, 25, 1.3, 20)
    add(file, 2, 'table', rows=5, cols=2, x='1.5cm', y='6cm', width='30cm', height='8.5cm', firstRow='true', headerFill='D6E2EE')
    values = [['Concepto', 'Importe (€)']] + [[item, f'{amount//100},{amount%100:02d}\u00a0€'] for item, amount in ROWS] + [['Total', '7,50\u00a0€']]
    for n, (a, b) in enumerate(values, 1):
        cli('set', file, f'/slide[2]/table[1]/tr[{n}]', '--prop', 'c1=' + a, '--prop', 'c2=' + b)
    cli('set', file, '/slide[2]/table[1]', '--prop', 'font=Liberation Sans', '--prop', 'size=22')
    shape(file, 2, 'Body', 'Los importes proceden de los datos locales autorizados.', 1.5, 15, 30, 1.8, 20)
    add(file, 3, 'chart', chartType='column', categories=','.join(item for item, _ in ROWS), x='1.5cm', y='6cm', width='18cm', height='10cm',
        **{'series1.name': 'Importe (€)', 'series1.values': '1.25,2.5,3.75', 'series1.color': BLUE})
    shape(file, 3, 'Body', 'Los tres conceptos suman 7,50\u00a0€. El detalle editable se conserva en la tabla y en el gráfico.', 21, 6, 11, 10, 20)
    cli('close', file)


def external_template(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(33.867), Cm(19.05)
    for number, title in enumerate(['Plantilla corporativa', 'Resumen técnico', 'Siguiente revisión'], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for name, value, x, y, w, h, size in [('Title', title, 1.5, 2, 30, 3, 36),
                ('Body', 'Contenido sintético conservado: España, ñ y €.', 1.5, 6, 30, 5, 24),
                ('Footer', f'{number} · Informe interno', 1.5, 16.8, 30, .9, 12)]:
            shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
            shape.name = name
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.text = value
            paragraph.font.name, paragraph.font.size = 'Liberation Sans', Pt(size)
            paragraph.font.color.rgb = RGBColor.from_string(BLUE)
    # Actual master decoration, inherited by all three blank-layout slides.
    brand = deck.slides[0].shapes.add_textbox(Cm(1.5), Cm(.9), Cm(30), Cm(.9))
    brand.name = 'MasterBrand'
    brand.text_frame.paragraphs[0].text = 'REGULAY · PLANTILLA CORPORATIVA'
    brand.text_frame.paragraphs[0].font.name = 'Liberation Sans'
    brand.text_frame.paragraphs[0].font.size = Pt(12)
    deck.slide_master.shapes._spTree.append(deepcopy(brand._element))
    brand._element.getparent().remove(brand._element)
    deck.save(path)


def parts(path):
    with zipfile.ZipFile(path) as source:
        return {name: source.read(name) for name in source.namelist()}


def inspect(file, base=None, allowed=()):
    cli('close', file)
    proof = {}
    package = parts(file)
    if base:
        previous = parts(base)
        proof['changed_parts'] = [name for name in package if package.get(name) != previous.get(name)]
        # OfficeCLI adds its Version/LastModified custom properties and XML namespace
        # declarations. Confirm master/layout semantics before restoring all
        # untouched original members, retaining only the requested slide changes.
        def semantic(data):
            def node(x):
                return (x.tag, sorted(x.attrib.items()), x.text or '', [node(c) for c in x])
            return node(etree.fromstring(data))
        for name in previous:
            if name not in (*allowed, '[Content_Types].xml', '_rels/.rels', 'docProps/custom.xml'):
                assert semantic(package[name]) == semantic(previous[name]) if name.endswith(('.xml', '.rels')) else package[name] == previous[name], name
        target = file.with_suffix('.stable.pptx')
        with zipfile.ZipFile(base) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
            for member in source.infolist():
                output.writestr(member, package[member.filename] if member.filename in allowed else source.read(member))
        target.replace(file)
        package = parts(file)
        assert set(package) == set(previous)
        assert not {name for name in package if package[name] != previous[name]} - set(allowed)
        proof['untouched_parts_preserved_byte_for_byte'] = True
    proof['sha256'] = hashlib.sha256(file.read_bytes()).hexdigest()
    deck = Presentation(file)
    assert len(deck.slides) == 3
    for slide in deck.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width and shape.top + shape.height <= deck.slide_height
    proof['masters'], proof['slides'] = len(deck.slide_masters), len(deck.slides)
    proof['tables'] = sum(s.has_table for slide in deck.slides for s in slide.shapes)
    proof['charts'] = sum(s.has_chart for slide in deck.slides for s in slide.shapes)
    validation = cli('validate', file, strict=False)
    issues = cli('view', file, 'issues', strict=False)
    proof['validate_code'], proof['issues_code'] = validation.returncode, issues.returncode
    proof['issues_text'] = issues.stdout
    assert validation.returncode == 0 and issues.returncode == 0 and 'Found 0 issue(s)' in issues.stdout
    output = file.parent
    subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(output/'lo'), '--headless',
        '--convert-to', 'pdf:impress_pdf_Export', '--outdir', str(output), str(file)], check=True, capture_output=True, timeout=60)
    pdf = file.with_suffix('.pdf')
    subprocess.run(['pdftoppm', '-r', '100', '-png', str(pdf), str(output/'slide')], check=True, timeout=60)
    assert len(list(output.glob('slide-*.png'))) == 3
    assert hashlib.sha256(file.read_bytes()).hexdigest() == proof['sha256']
    proof['pdf_sha256'] = hashlib.sha256(pdf.read_bytes()).hexdigest()
    return proof


report = {'version': cli('--version').stdout.strip(), 'cases': {}}
try:
    for name in ('new', 'template', 'one-slide'):
        (ROOT/name).mkdir()
    new = ROOT/'new/final.pptx'
    new_deck(new)
    report['cases']['new'] = inspect(new)
    template = ROOT/'template/base.pptx'
    external_template(template)
    modified = ROOT/'template/final.pptx'
    shutil.copyfile(template, modified)
    cli('set', modified, '/slide[1]/shape[@name=Title]', '--prop', 'text=Informe corporativo de España')
    cli('set', modified, '/slide[2]/shape[@name=Body]', '--prop', 'text=Resumen técnico revisado, año 2026. Conserva cabeceras, diseños y numeración.')
    report['cases']['template'] = inspect(modified, template, ('ppt/slides/slide1.xml', 'ppt/slides/slide2.xml'))
    single = ROOT/'one-slide/final.pptx'
    shutil.copyfile(new, single)
    cli('set', single, '/slide[3]/shape[@name=Body]', '--prop', 'text=Revisión localizada: análisis de importes de España. Tildes, ñ y símbolos monetarios €, £ y $.')
    report['cases']['one-slide'] = inspect(single, new, ('ppt/slides/slide3.xml',))
    report['structural_passed'] = True
finally:
    (ROOT/'commands.json').write_text(json.dumps(LOG, indent=2, ensure_ascii=False))
    (ROOT/'result.json').write_text(json.dumps(report, indent=2, ensure_ascii=False))
print(json.dumps(report))
