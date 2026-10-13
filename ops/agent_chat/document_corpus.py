"""Sixteen local tasks, final-byte verification in another process, no inference."""
import base64
import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject

sys.path.insert(0, '/source/regulai')
from front.agents import excel_contract as excel, document_contract as docx, presentation_contract as pptx

TITLE = 'Informe económico de España'
SUMMARY = 'Análisis sintético: año, acción, piñón y 1.250,00 €. Revisar los importes antes de utilizar el informe.'
CASES = [('docx-template','docx','Informe con plantilla'),('docx-pages','docx','Tabla que cruza páginas'),
    ('docx-title','docx','Cambiar título conservando estilos'),('docx-comments','docx','Editar comentarios/control de cambios'),
    ('xlsx-dates','xlsx','CSV con identificadores y fechas'),('xlsx-cross-sheet','xlsx','Fórmulas entre hojas y nulos'),
    ('xlsx-formula-fill','xlsx','Relleno que cruza celdas de fórmula'),('xlsx-chart-print','xlsx','Gráfico y vista de impresión'),
    ('pptx-masters','pptx','Plantilla con maestros'),('pptx-new','pptx','Deck nuevo'),
    ('pptx-local-edit','pptx','Cambio localizado'),('pptx-long','pptx','Gráfico/tabla editables y texto largo'),
    ('pdf-spanish','pdf','Informe español con fuentes'),('pdf-form','pdf','Formulario'),
    ('pdf-ocr','pdf','Extracción de escaneo con OCR'),('pdf-redact','pdf','Eliminación verificable de información')]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify(case):
    spec = json.loads((case/'spec.json').read_text())
    source = (case/'source.csv').read_bytes() if (case/'source.csv').exists() else (case/'source.xlsx').read_bytes()
    base = Path(spec['base']).read_bytes() if spec.get('base') else None
    if spec['format'] in ('docx','pdf'):
        _, proof, _ = docx.verify_bundle(base64.b64encode((case/'candidate.zip').read_bytes()).decode(), source,
            'csv', spec['expected'], format=spec['format'], base=base, changed=spec.get('changed'))
    elif spec['format'] == 'pptx':
        _, proof, _ = pptx.verify_bundle(base64.b64encode((case/'candidate.zip').read_bytes()).decode(), source,
            'csv', spec['expected'], base=base, changed=tuple(spec['changed']) if spec.get('changed') else None,
            template=spec['template'])
    else:
        proof, _ = excel.verify((case/'final.xlsx').read_bytes(), source, 'xlsx', title=TITLE)
        from openpyxl import load_workbook
        book = load_workbook(case/'final.xlsx')
        if not book.active._charts:
            raise ValueError('corpus_requested_chart_not_preserved')
    (case/'verification.json').write_text(json.dumps(proof, indent=2))


def run(output):
    output.mkdir(parents=True, exist_ok=False)
    records = []
    def execute(case, command):
        started = time.monotonic()
        result = subprocess.run(command, capture_output=True, timeout=180)
        with (case/'commands.log').open('ab') as log:
            log.write(result.stdout+result.stderr)
        record['commands'].append({'argv':command,'exit_code':result.returncode,'seconds':round(time.monotonic()-started,3)})
        if result.returncode:
            record['failure_detail'] = result.stderr.decode(errors='replace').splitlines()[-1] if result.stderr else 'Command failed'
            raise ValueError('local_command_failed')
    def csv(case, count=3, long=False):
        path = case/'source.csv'
        path.write_text('item,amount_cents\n'+''.join(f'Concepto español {n:02d}, {n*125}\n'.replace(', ', ',') for n in range(1,count+1)),encoding='utf-8')
        if not long:
            path.write_text('item,amount_cents\nAño,125\nAcción,250\nPiñón,375\n',encoding='utf-8')
        return path
    def reject(case, path, kind):
        try:
            excel.input_rows(path.read_bytes(),kind)
        except ValueError as error:
            record.update(status='not_supported',reason=str(error),gate='production input_rows')
            return
        raise ValueError('unsupported_fixture_unexpectedly_accepted')
    for key, format, task in CASES:
        case = output/key; case.mkdir()
        record = {'id':key,'format':format,'task':task,'status':'failed','commands':[],
            'repairs':0,'model_calls':0,'cost_usd':'0','visual_review':'not_applicable'}
        started = time.monotonic()
        try:
            fields = {'title':TITLE,'section':SUMMARY}
            spec = {'format':format,'expected':fields.copy()}
            if key in ('docx-template','docx-pages','docx-title','pdf-spanish'):
                source = csv(case,65,key == 'docx-pages')
                program = '/skills/document_program.py'
                operation = 'build'
                if key == 'docx-title':
                    source = output/'docx-template/final.docx'
                    fields = {'field':'title','value':'Informe económico revisado'}
                    spec.update(base=str(source),changed='title',expected={'title':fields['value'],'section':SUMMARY})
                operation = 'edit' if key == 'docx-title' else 'build'
                (case/'fields.json').write_text(json.dumps(fields,ensure_ascii=False))
                execute(case,[sys.executable,'-I',program,operation,'--output',str(case),'--source',str(source),
                    '--fields',str(case/'fields.json'),'--template','/skills/corporate.docx'])
                for phase in ('render','seal'): execute(case,[sys.executable,'-I',program,phase,'--output',str(case)])
            elif key.startswith('pptx-') and key != 'pptx-long':
                source = csv(case)
                fields['template'] = 'corporate' if key == 'pptx-masters' else 'new'
                spec.update(template=fields['template'],expected=pptx.expected_fields(TITLE,SUMMARY))
                operation = 'build'
                if key == 'pptx-local-edit':
                    source = output/'pptx-new/final.pptx'
                    fields = {'slide':2,'field':'title','value':'Detalle revisado de importes'}
                    spec.update(base=str(source),changed=[2,'title'])
                    spec['expected'][1]['title'] = fields['value']
                    operation = 'edit'
                (case/'fields.json').write_text(json.dumps(fields,ensure_ascii=False))
                program='/skills/presentation_program.py'
                execute(case,[sys.executable,'-I',program,operation,'--output',str(case),'--source',str(source),
                    '--fields',str(case/'fields.json'),'--template','/skills/corporate.pptx'])
                for phase in ('render','seal'): execute(case,[sys.executable,'-I',program,phase,'--output',str(case)])
            elif key == 'xlsx-chart-print':
                book = Workbook(); sheet = book.active
                for row in [('item','amount_cents'),('Año',125),('Acción',250),('Piñón',375)]: sheet.append(row)
                chart=BarChart(); chart.add_data(Reference(sheet,min_col=2,min_row=1,max_row=4),titles_from_data=True)
                sheet.add_chart(chart,'D2'); sheet.print_area='A1:L20'; book.save(case/'source.xlsx')
                (case/'fields.json').write_text(json.dumps({'title':TITLE}))
                execute(case,[sys.executable,'-I','/skills/spreadsheet_program.py','build',str(case/'source.xlsx'),
                    '--output',str(case/'candidate.xlsx'),'--fields',str(case/'fields.json')])
                # Converter cannot overwrite input; use a separate output directory.
                recalc=case/'recalc'; recalc.mkdir()
                execute(case,['libreoffice','-env:UserInstallation=file://'+str(case/'lo2'),'--headless',
                    '--convert-to','xlsx','--outdir',str(recalc),str(case/'candidate.xlsx')])
                shutil.copyfile(recalc/'candidate.xlsx',case/'final.xlsx')
                execute(case,[sys.executable,'-I','/skills/spreadsheet_program.py','check',str(case/'final.xlsx')])
                execute(case,['libreoffice','-env:UserInstallation=file://'+str(case/'lo3'),'--headless',
                    '--convert-to','pdf','--outdir',str(case),str(case/'final.xlsx')])
                execute(case,['pdftoppm','-r','100','-png',str(case/'final.pdf'),str(case/'page')])
            elif key == 'xlsx-dates':
                path=case/'source.csv'; path.write_text('id,date,amount_cents\n0001,2026-01-02,125\n0002,2026-01-03,250\n')
                reject(case,path,'csv')
            elif key in ('xlsx-cross-sheet','xlsx-formula-fill'):
                book=Workbook(); book.active.append(['item','amount_cents']); book.active.append(['A',None])
                book.active.append(['B','=100+25'])
                if key == 'xlsx-cross-sheet':
                    book.create_sheet('Source')['A1']=125; book.active['B2']='=Source!A1'
                book.save(case/'source.xlsx'); reject(case,case/'source.xlsx','xlsx')
            elif key == 'pptx-long':
                shutil.copyfile(output/'pptx-new/final.pptx',case/'source.pptx')
                (case/'request.json').write_text(json.dumps({'field':'body','value':SUMMARY*8}))
                try: pptx.valid_text(SUMMARY*8,'body')
                except ValueError as error: record.update(status='not_supported',reason=str(error),gate='production valid_text')
                else: raise ValueError('long_text_unexpectedly_accepted')
            elif key == 'docx-comments':
                from docx import Document
                document=Document(output/'docx-template/final.docx')
                document.add_comment(document.paragraphs[0].runs[0],text='Comentario sintético',author='Laboratorio')
                document.save(case/'source.docx'); reject(case,case/'source.docx','docx')
            elif key == 'pdf-form':
                writer=PdfWriter(); writer.add_blank_page(595,842)
                from pypdf.generic import TextStringObject
                field=DictionaryObject({NameObject('/FT'):NameObject('/Tx'),NameObject('/T'):TextStringObject('Nombre'),
                    NameObject('/V'):TextStringObject('Private Synthetic Form Marker')})
                writer._root_object[NameObject('/AcroForm')]=DictionaryObject({NameObject('/Fields'):ArrayObject([writer._add_object(field)])})
                writer.write(case/'source.pdf'); reject(case,case/'source.pdf','pdf')
            elif key == 'pdf-ocr':
                image=Image.new('RGB',(827,1170),'white'); ImageDraw.Draw(image).text((50,50),'Private Synthetic OCR Marker',fill='black')
                image.save(case/'source.pdf','PDF'); reject(case,case/'source.pdf','pdf')
            else:
                shutil.copyfile(output/'pdf-spanish/final.pdf',case/'source.pdf')
                (case/'request.json').write_text(json.dumps({'remove':'Piñón','must_not_appear':'Piñón'},ensure_ascii=False))
                reject(case,case/'source.pdf','pdf')
            if record['status'] != 'not_supported':
                (case/'spec.json').write_text(json.dumps(spec,indent=2,ensure_ascii=False))
                execute(case,[sys.executable,__file__,'verify',str(case)])
                record.update(status='passed',visual_review='required')
        except Exception as error:
            record.update(status='failed',reason=str(error))
        record['seconds']=round(time.monotonic()-started,3)
        record['files']={p.name:sha(p) for p in case.iterdir() if p.is_file() and p.suffix in ('.docx','.xlsx','.pptx','.pdf','.png','.csv','.json')}
        record['previews']=[p.name for p in sorted(case.glob('page-*.png'))]
        records.append(record)
        (output/'report.json').write_text(json.dumps({'scope':'Offline production helpers in pinned Linux image, independent verifier process; no agent/model quality claim',
            'cases':records,'passed':sum(r['status']=='passed' for r in records),'failed':sum(r['status']=='failed' for r in records),
            'not_supported':sum(r['status']=='not_supported' for r in records),'paid_calls':0},indent=2,ensure_ascii=False))
    evaluation={'status':'pending_specific_paid_model_authorization','repetitions_per_case':3,'cases':records,
        'model':None,'temperature':0,'seed':None,'provider_limits':'provider_managed, no new generation caps',
        'comparison':'Same authorized model/config and input hashes for skill comparison; Claude product separately',
        'score':['content','structure','recalculated_formulas','editability','visual','unrequested_changes','repairs','latency','cost_per_valid_delivery','privacy'],
        'source_sha256':{str(p.relative_to('/source')):sha(p) for p in Path('/source/regulai/front/agents').glob('*program.py')},
        'runtime_sha256':{str(p.relative_to('/source')):sha(p) for p in Path('/source/regulai/front/agents/runtime').rglob('*.md')}}
    (output/'generative-plan.json').write_text(json.dumps(evaluation,indent=2,ensure_ascii=False))
    print(json.dumps({'cases':len(records),'passed':sum(r['status']=='passed' for r in records),
        'failed':sum(r['status']=='failed' for r in records),'not_supported':sum(r['status']=='not_supported' for r in records)}))


if __name__ == '__main__':
    (verify if sys.argv[1]=='verify' else run)(Path(sys.argv[2]))
