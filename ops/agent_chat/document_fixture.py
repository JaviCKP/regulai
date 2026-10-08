"""Synthetic corporate template for the task-51 engine acceptance experiment."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_template(path):
    doc = Document()
    doc.core_properties.version = '2'
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number'):
        style = doc.styles[name]
        style.font.name = 'Liberation Sans'
        style.font.color.rgb = RGBColor.from_string('193B45' if name != 'Normal' else '263D44')
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(9)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    doc.styles['Title'].font.name = 'Liberation Serif'
    doc.styles['Title'].font.size = Pt(38)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number'):
        fonts = doc.styles[name].element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if key.lower().endswith('theme'):
                del fonts.attrib[key]
    for border in doc.styles['Title'].element.iter(qn('w:pBdr')):
        border.getparent().remove(border)
    doc.styles['Title'].paragraph_format.space_before = Pt(52)
    doc.styles['Title'].paragraph_format.space_after = Pt(24)
    doc.styles['Subtitle'].font.size = Pt(10)
    doc.styles['Subtitle'].font.color.rgb = RGBColor.from_string('5F7479')
    doc.styles['Heading 1'].font.size = Pt(20)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(24)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.85)
    header = doc.sections[0].header.paragraphs[0]
    header.text = 'REGULAY  /  INFORMES'
    header.runs[0].font.size = Pt(9)
    header.runs[0].bold = True
    footer = doc.sections[0].footer.paragraphs[0]
    footer.style = doc.styles['Subtitle']
    footer.add_run('Página ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' · Informe interno')
    doc.add_paragraph('{{TITLE}}', 'Title')
    doc.add_paragraph('INFORME ECONÓMICO  ·  ESPAÑA', 'Subtitle')
    doc.add_heading('Resumen ejecutivo', 1)
    doc.add_paragraph('{{SUMMARY}}')
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph('01  /  DATOS ECONÓMICOS', 'Subtitle')
    doc.add_heading('Detalle de importes', 1)
    doc.add_paragraph('Desglose por concepto, expresado en euros.', 'Subtitle')
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Inches(4.8), Inches(1.75)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Concepto', 'Importe (€)'
    table.cell(1, 0).text, table.cell(1, 1).text = '{{ITEM}}', '{{AMOUNT}}'
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = table.columns[col_index].width
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if col_index else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.runs[0].font.size = Pt(10)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), '193B45')
                cell._tc.get_or_add_tcPr().append(shade)
                paragraph.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                paragraph.runs[0].bold = True
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_heading('Notas de revisión', 1)
    doc.add_paragraph('Cada fila corresponde a un concepto del archivo de origen. El total suma todos los importes en céntimos antes de expresarlos en euros.')
    doc.add_paragraph('Contrastar los conceptos con el archivo de origen.', 'List Number')
    doc.add_paragraph('Revisar el total antes de utilizar este informe.', 'List Number')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == '__main__':
    import sys
    create_template(sys.argv[1])
